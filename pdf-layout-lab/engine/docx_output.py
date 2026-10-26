"""Lab-only fixed-layout DOCX: preserved artwork plus editable text boxes.

The fitted PDF is an intermediate. The delivered DOCX contains Word text, not
page screenshots containing the translated text. Artwork remains a background;
table borders are visual artwork, not reconstructed Word table cells.
"""
from io import BytesIO
import re
import math
from functools import lru_cache

import fitz
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt
from xml.sax.saxutils import escape, quoteattr
from engine.xml_text import clean, INVALID
from engine.typography import aligned_groups, physical_rows, union, lines as logical_lines
from engine.enhanced import cleaned_patch

WPS = 'http://schemas.microsoft.com/office/word/2010/wordprocessingShape'


@lru_cache(maxsize=2)
def _ocr_font(cjk):
    return fitz.Font('korea' if cjk else 'helv')


def _emu(value):
    return round(value * 12700)


def _anchor(shape, identifier, x, y, width, height, behind=False):
    return parse_xml(f'''<w:drawing {nsdecls('w', 'wp', 'a', 'pic', 'r')} xmlns:wps="{WPS}">
      <wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0"
        relativeHeight="{identifier}" behindDoc="{int(behind)}" locked="0" layoutInCell="1" allowOverlap="1">
        <wp:simplePos x="0" y="0"/>
        <wp:positionH relativeFrom="page"><wp:posOffset>{_emu(x)}</wp:posOffset></wp:positionH>
        <wp:positionV relativeFrom="page"><wp:posOffset>{_emu(y)}</wp:posOffset></wp:positionV>
        <wp:extent cx="{_emu(width)}" cy="{_emu(height)}"/>
        <wp:wrapNone/><wp:docPr id="{identifier}" name="Lab object {identifier}"/>
        <wp:cNvGraphicFramePr/><a:graphic>{shape}</a:graphic>
      </wp:anchor></w:drawing>''')


def _font(span):
    name = re.sub(r'^[A-Z]{6}\+', '', span['font'])
    if any('\uac00' <= c <= '\ud7a3' for c in span['text']):
        return 'Noto Sans CJK KR'
    if any('\u3040' <= c <= '\u9fff' for c in span['text']):
        return 'Noto Sans CJK JP'
    if 'times' in name.lower() or name.lower() in {'tiro', 'tibo', 'tiit', 'tibi'}:
        return 'Times New Roman'
    if 'courier' in name.lower():
        return 'Courier New'
    return 'Arial'


def _text_box(line, identifier):
    x0, y0, x1, y1 = line['bbox']
    width, height = max(1, x1 - x0), max(1, y1 - y0)
    direction = line.get('dir', (1, 0))
    vertical = abs(direction[0]) < .01 and abs(abs(direction[1])-1) < .01
    line_height = width if vertical else height
    text_direction = (' vert="vert270"' if direction[1] < 0 else ' vert="vert"') if vertical else ''
    runs = []
    for span_index, span in enumerate(line['spans']):
        text = clean(span['text'])
        if not text:
            continue
        size = max(2, round(span['size'] * 2))
        flags = span['flags']
        style = ('<w:b/>' if flags & 16 else '') + ('<w:i/>' if flags & 2 else '')
        font = quoteattr(_font(span))
        # Fitted PDF spans define the desired physical run width. Word's fitText
        # retains editable characters while compensating for substitute metrics.
        run_width = max(1, round((span['bbox'][3] - span['bbox'][1] if vertical else span['bbox'][2] - span['bbox'][0]) * 20))
        runs.append(f'''<w:r><w:rPr><w:rFonts w:ascii={font} w:hAnsi={font} w:eastAsia={font}/>
          <w:sz w:val="{size}"/><w:szCs w:val="{size}"/><w:color w:val="{span['color']:06X}"/>
          <w:fitText w:val="{run_width}" w:id="{identifier * 1000 + span_index}"/>{style}</w:rPr>
          <w:t xml:space="preserve">{escape(text)}</w:t></w:r>''')
    content = f'''<a:graphicData uri="{WPS}"><wps:wsp><wps:cNvSpPr txBox="1"/>
      <wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{_emu(width)}" cy="{_emu(height)}"/></a:xfrm>
      <a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/><a:ln><a:noFill/></a:ln></wps:spPr>
      <wps:txbx><w:txbxContent><w:p><w:pPr><w:spacing w:before="0" w:after="0"
      w:line="{max(20, round(line_height * 20))}" w:lineRule="exact"/></w:pPr>{''.join(runs)}</w:p></w:txbxContent></wps:txbx>
      <wps:bodyPr wrap="none" lIns="0" tIns="0" rIns="0" bIns="0" anchor="ctr"{text_direction}><a:noAutofit/></wps:bodyPr>
      </wps:wsp></a:graphicData>'''
    return _anchor(content, identifier, x0, y0, width, height)


def _original_lines(region):
    """Use the corrected OCR transcript, not stale draft span text."""
    texts = logical_lines(region['text'])
    if not texts or '[OCR]' in region['text'] or not region.get('spans'):
        raise ValueError('인식된 원문이 없어 텍스트로 복원하지 못했습니다.')
    groups = aligned_groups(region, texts)
    if groups is None and len(texts) == 1 and re.fullmatch(r'[가-힣]+', texts[0]):
        rows = physical_rows(region)
        ink = fitz.Rect(region['source_bbox'])
        if ink.height > ink.width*1.5 and len(rows) == len(texts[0]):
            texts, groups = list(texts[0]), [[r] for r in rows]
    if groups is None:
        # OCR correction can merge or split physical lines. Keep every corrected
        # line editable within the source ink bounds instead of using a stale draft.
        ink = fitz.Rect(region.get('source_bbox', region['bbox']))
        step = ink.height / len(texts)
        boxes = [[ink.x0, ink.y0+i*step, ink.x1, ink.y0+(i+1)*step] for i in range(len(texts))]
    else:
        boxes = [list(union(row['box'] for row in group)) for group in groups]
    result = []
    for text, box in zip(texts, boxes):
        size = min(region['font_size'], (box[3]-box[1])*.85)
        font = _ocr_font(bool(re.search(r'[\u3000-\u9fff\uac00-\ud7a3]', text)))
        # OnlyOffice does not consistently honor w:fitText for OCR/fallback
        # fonts. Fit the actual font size first, with half-point rounding slack.
        unit_width = font.text_length(text, fontsize=1)
        size = min(size, (box[2]-box[0])*.94/max(unit_width,.1))
        if size < 3:
            raise ValueError('인식된 원문이 해당 위치에 들어가지 않아 복원을 보류했습니다.')
        size = math.floor(size*2)/2
        result.append({'bbox':box, 'dir':(1,0), 'spans':[{'text':text, 'bbox':box,
                       'size':size, 'font':'OCR', 'flags':region.get('dominant_flags',0),
                       'color':region.get('dominant_color',0)}]})
    return result


def _ocr_candidates(page, page_model, outcomes, native_lines):
    candidates, details = [], []
    for region in page_model.get('regions', []):
        outcome = outcomes.get(region['id'], {})
        if not region.get('ocr') or outcome.get('status') not in ('UNCHANGED', 'NEEDS_REVIEW'):
            continue
        detail = {'id':region['id'], 'page':page.number, 'editable':False}
        details.append(detail)
        if region.get('keep_original') or region.get('kind') == 'protected_text':
            detail['reason'] = '보호 영역·체크 표시는 원본 이미지로 보존했습니다.'
            continue
        try:
            restored = _original_lines(region)
            masks = list(region.get('source_mask_boxes', region.get('mask_boxes', [region['source_bbox']])))
            if re.search(r'\.{2,}|…', region['text']):
                # Word OCR lines include the dotted leader itself. Clear the
                # original gap between label/value as well, avoiding double dots.
                masks.extend(list(row['box']) for row in physical_rows(region) if len(row['spans']) > 1)
            # A manually applied neighbour may already cover this old source ink.
            # Do not erase that text or put a second OCR layer over it.
            if any((fitz.Rect(mask) & fitz.Rect(line['bbox'])).get_area() > .5
                   for mask in masks for line in native_lines):
                raise ValueError('현재 PDF의 다른 텍스트와 겹쳐 원문 복원을 보류했습니다.')
            bounds = union([*masks, *[line['bbox'] for line in restored]]) + (-1,-1,1,1)
            bounds &= page.rect
            if bounds.is_empty or any(not page.rect.contains(fitz.Rect(line['bbox'])) for line in restored):
                raise ValueError('원문 위치가 페이지 경계를 벗어납니다.')
            candidates.append((region, restored, dict(region, bbox=list(bounds), mask_boxes=masks), detail))
        except ValueError as exc:
            detail['reason'] = str(exc)
    return candidates, details


def create_docx(fitted_pdf, output, layout=None, report=None):
    """Export native text plus retained OCR originals without changing the PDF."""
    document = Document()
    document.core_properties.title = 'Layout Lab editable output'
    warnings, page_reports, ocr_details = [], [], []
    outcomes = {r['id']:r for r in (report or {}).get('regions', [])}
    identifier = 1
    with fitz.open(fitted_pdf) as source:
        for pno, page in enumerate(source):
            section = document.sections[0] if pno == 0 else document.add_section(WD_SECTION_START.NEW_PAGE)
            section.page_width, section.page_height = Pt(page.rect.width), Pt(page.rect.height)
            section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Pt(0)
            section.header_distance = section.footer_distance = Pt(0)
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = Pt(1)
            paragraph.add_run().font.size = Pt(1)
            lines, retained, rotated = [], 0, 0
            for block in page.get_text('dict')['blocks']:
                if block['type'] != 0:
                    continue
                for line in block['lines']:
                    if not any(s['text'].strip() for s in line['spans']):
                        continue
                    horizontal = abs(line['dir'][0]-1) < .01 and abs(line['dir'][1]) < .01
                    vertical = abs(line['dir'][0]) < .01 and abs(abs(line['dir'][1])-1) < .01
                    if page.rotation or not (horizontal or vertical):
                        retained += 1
                    else:
                        lines.append(line)
                        rotated += int(vertical)
            page_model = layout['pages'][pno] if layout and pno < len(layout['pages']) else {}
            candidates, details = _ocr_candidates(page, page_model, outcomes, lines)
            ocr_details.extend(details)
            restored_lines = []
            # Redact only native text; pictures, table rules and decorations stay.
            with fitz.open() as background:
                background.insert_pdf(source, from_page=pno, to_page=pno)
                art = background[0]
                for line in lines:
                    art.add_redact_annot(fitz.Rect(line['bbox']), fill=False, cross_out=False)
                if lines:
                    art.apply_redactions(images=0, graphics=0, text=0)
                overlays = [r for r in page_model.get('regions', []) if r.get('overlay_safe')]
                for region, restored, patch, detail in candidates:
                    try:
                        # Only measured original ink is cleared. A white text-box
                        # fill would cover table rules and leave duplicate letters.
                        png = cleaned_patch(art, patch, overlays)
                        art.insert_image(fitz.Rect(patch['bbox']), stream=png, keep_proportion=False)
                        restored_lines.extend(restored)
                        detail.update(editable=True, lines=len(restored), reason='OCR 원문을 편집 가능한 텍스트로 복원했습니다.')
                    except ValueError as exc:
                        detail['reason'] = str(exc)
                png = art.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False).tobytes('png')
            host = paragraph.add_run()
            picture = host.add_picture(BytesIO(png), width=section.page_width, height=section.page_height)
            graphic = picture._inline.graphic.graphicData.xml
            host._r.remove(host._r.find(qn('w:drawing')))
            host._r.append(_anchor(graphic, identifier, 0, 0, page.rect.width, page.rect.height, behind=True))
            identifier += 1
            lines.extend(restored_lines)
            for line in lines:
                paragraph.add_run()._r.append(_text_box(line, identifier))
                identifier += 1
            invalid = sum(len(INVALID.findall(s['text'])) for line in lines for s in line['spans'])
            if invalid:
                warnings.append(f'{pno + 1}페이지: PDF의 XML 비호환 문자 {invalid}개를 �로 표시했습니다. 원본과 대조하세요.')
            if retained:
                warnings.append(f'{pno + 1}페이지: 대각선·보호 텍스트 {retained}줄은 배경으로 보존됩니다.')
            if not lines:
                warnings.append(f'{pno + 1}페이지: 편집 가능한 텍스트가 없습니다. 이미지/OCR 결과를 확인하세요.')
            page_reports.append({'page': pno, 'editable_lines': len(lines), 'retained_lines': retained,
                                 'editable_ocr_lines':len(restored_lines), 'editable_rotated_lines':rotated})
    document.save(str(output))
    restored_count = sum(r['editable'] for r in ocr_details)
    if restored_count:
        warnings.insert(0, f'미번역·원문 유지 {restored_count}개 영역을 OCR 텍스트로 복원했습니다. 인식 오류와 글꼴·줄 위치를 원본과 대조하세요.')
    return {'ready': True, 'version':2, 'pages': page_reports, 'editable_lines': sum(p['editable_lines'] for p in page_reports),
            'editable_original_regions':restored_count, 'original_regions':ocr_details,
            'editable_rotated_lines':sum(p['editable_rotated_lines'] for p in page_reports),
            'warnings': warnings, 'preview_status': 'pending',
            'scope': '번역문·복원된 OCR 원문·세로 공정명은 텍스트 상자로 편집할 수 있습니다. 스캔 원문은 글꼴·줄 위치가 달라질 수 있습니다. 표 선·그림·보호 영역은 배경입니다.'}
