import json
import threading
import time

from docx import Document
from fastapi.testclient import TestClient

import app as web
from engine import legacy
from engine.inspect import inspect_pdf
from test_lab import make_pdf


class CountingTranslator:
    def __init__(self):
        self.calls = 0

    def translate(self, texts, source_lang, target_lang, glossary, **kwargs):
        self.calls += 1
        return ["번역 " + text for text in texts]


def test_existing_translation_logic_is_cached_and_used_for_docx(tmp_path, monkeypatch):
    legacy.setup()
    from pipeline import translator
    delegate = CountingTranslator()
    monkeypatch.setattr(legacy, "configuration", lambda: {"configured": True, "provider": "test", "model": "test"})
    monkeypatch.setattr(translator, "get_translator", lambda: delegate)
    layout = inspect_pdf(make_pdf(tmp_path / "source.pdf"))
    first = legacy.translate_shared(layout, "en", "ko", tmp_path / "cache")
    second = legacy.translate_shared(layout, "en", "ko", tmp_path / "cache")
    assert delegate.calls == 1
    assert not first["cache_hit"] and second["cache_hit"]
    assert first["replacements"] == second["replacements"]
    output = tmp_path / "comparison.docx"
    legacy.create_docx(layout, first["replacements"], "en", "ko", output)
    actual = [p.text for p in Document(output).paragraphs if p.text]
    assert actual == list(first["replacements"].values())
    legacy.translate_shared(layout, "en", "ja", tmp_path / "cache")
    assert delegate.calls == 2  # Language pair is part of the cache identity.


def test_no_key_never_uses_mock_translation(tmp_path, monkeypatch):
    import pytest
    monkeypatch.setattr(legacy, "configuration", lambda: {"configured": False, "provider": None, "model": None})
    layout = inspect_pdf(make_pdf(tmp_path / "source.pdf"))
    with pytest.raises(ValueError, match="모의 번역"):
        legacy.translate_shared(layout, "en", "ko", tmp_path / "cache")


def wait_job(client, job_id):
    for _ in range(300):
        state = client.get(f"/api/jobs/{job_id}").json()
        if state["status"] in {"completed", "partial", "failed"}:
            return state
        time.sleep(.02)
    raise AssertionError("comparison worker did not finish")


def test_comparison_uses_same_snapshot_and_conversion_token(tmp_path, monkeypatch):
    monkeypatch.setattr(web, "DATA", tmp_path / "data")
    web.DATA.mkdir()
    client = TestClient(web.app)
    monkeypatch.setattr(legacy, "configuration", lambda: {"configured": True, "provider": "test", "model": "test"})
    captured = {}

    def project(source, source_lang, target_lang, cache_dir, output, progress):
        document = Document()
        document.add_paragraph('Independent project output')
        document.save(output)
        return {'origin': 'project_pipeline', 'scope': 'Independent full pipeline test'}
    monkeypatch.setattr(legacy, 'project_output', project)

    def translate(layout, source_lang, target_lang, cache_dir, progress):
        texts = {r["id"]: "Test" for p in layout["pages"] for r in p["regions"]}
        captured["texts"] = texts
        return {"id": "translation-1", "source_lang": source_lang, "target_lang": target_lang,
                "provider": "test", "model": "test", "replacements": texts,
                "unchanged_ids": [], "cache_hit": False}

    def convert(url, key, output):
        token = url.rsplit("/", 1)[1]
        captured["token"] = token
        # A callback while conversion is pending proves no lock is held around network IO.
        response = client.get(f"/api/conversion-input/{token}")
        assert response.status_code == 200
        assert response.content.startswith(b"PK")
        make_pdf(output)  # Test-only conversion double; never used by production.

    monkeypatch.setattr(legacy, "translate_shared", translate)
    monkeypatch.setattr(legacy, "convert_docx", convert)
    source = make_pdf(tmp_path / "sample.pdf")
    item = client.post("/api/documents", files={"file": ("sample.pdf", source.read_bytes())}).json()
    prefix = f"/api/documents/{item['id']}"
    assert client.post(prefix + "/compare", json={"source_lang":"en", "target_lang":"en"}).status_code == 400
    started = client.post(prefix + "/compare", json={"source_lang":"en", "target_lang":"ko"})
    assert started.status_code == 202
    job = wait_job(client, started.json()["id"])
    assert job["status"] == "completed", job
    stages = [step['stage'] for step in job['steps']]
    assert stages.index('new_docx') < stages.index('new_preview') < stages.index('docx') < stages.index('legacy_pdf')
    assert all(step['status'] == 'finished' for step in job['steps'])
    item = client.get(prefix).json()
    assert item["report"]["policy"]["replacements"] == captured["texts"]
    assert item['baseline']['origin'] == 'project_pipeline'
    assert 'text_hash' not in item['baseline']
    assert item['report']['processing']['status'] == 'completed'
    run = item["latest_run"]
    translation = client.get(prefix + f"/download/translation?run={run}").json()
    assert translation["replacements"] == captured["texts"]
    assert client.get(prefix + f"/download/docx?run={run}").content.startswith(b"PK")
    assert client.get(f"/api/conversion-input/{captured['token']}").status_code == 404


def test_conversion_failure_keeps_translation_and_new_output(tmp_path, monkeypatch):
    monkeypatch.setattr(web, "DATA", tmp_path / "data")
    web.DATA.mkdir()
    client = TestClient(web.app)
    monkeypatch.setattr(legacy, "configuration", lambda: {"configured": True, "provider": "test", "model": "test"})
    def translate(layout, source_lang, target_lang, cache_dir, progress):
        return {"id":"cached", "provider":"test", "model":"test", "cache_hit":True,
                "unchanged_ids":[], "replacements":{r["id"]:"Short" for p in layout["pages"] for r in p["regions"]}}
    def fail(*args):
        raise RuntimeError("OnlyOffice unavailable")
    monkeypatch.setattr(legacy, "translate_shared", translate)
    monkeypatch.setattr(legacy, "convert_docx", fail)
    def project(source, source_lang, target_lang, cache_dir, output, progress):
        Document().save(output)
        return {'origin': 'project_pipeline', 'scope': 'test'}
    monkeypatch.setattr(legacy, 'project_output', project)
    source = make_pdf(tmp_path / "sample.pdf")
    item = client.post("/api/documents", files={"file": ("sample.pdf", source.read_bytes())}).json()
    prefix = f"/api/documents/{item['id']}"
    job_id = client.post(prefix + "/compare", json={"source_lang":"en", "target_lang":"ko"}).json()["id"]
    assert wait_job(client, job_id)["status"] == "partial"
    item = client.get(prefix).json()
    assert item["baseline"] is None
    assert item["comparison"]["docx_ready"]
    assert item["comparison"]["legacy_status"] == "unavailable"
    assert client.get(prefix + f"/download/new?run={job_id}").content.startswith(b"%PDF")


def test_busy_document_rejects_other_mutations(tmp_path, monkeypatch):
    monkeypatch.setattr(web, "DATA", tmp_path / "data")
    web.DATA.mkdir()
    client = TestClient(web.app)
    source = make_pdf(tmp_path / "source.pdf")
    item = client.post("/api/documents", files={"file": ("sample.pdf", source.read_bytes())}).json()
    web.ACTIVE[item["id"]] = "test-job"
    try:
        assert client.post(f"/api/documents/{item['id']}/render", json={}).status_code == 409
        assert client.post(f"/api/documents/{item['id']}/baseline", files={"file": ("old.pdf", source.read_bytes())}).status_code == 409
    finally:
        web.ACTIVE.pop(item["id"], None)


def test_progress_remains_readable_while_pdf_lock_is_held(monkeypatch):
    monkeypatch.setitem(web.JOBS, 'progress-lock-test', {
        'stage': 'new_pdf', 'status': 'running', '_started': time.monotonic(),
        '_stage_started': time.monotonic(), 'steps': [],
    })
    done = threading.Event()
    result = {}
    def poll():
        result['response'] = TestClient(web.app).get('/api/jobs/progress-lock-test')
        done.set()
    with web.LOCK:
        thread = threading.Thread(target=poll, daemon=True)
        thread.start()
        responded = done.wait(2)
    thread.join(2)
    assert responded, 'Progress polling must not wait for MuPDF rendering'
    assert result['response'].status_code == 200


def test_stage_timings_reset_counts_and_stop_on_failure(monkeypatch):
    now = [100.0]
    monkeypatch.setattr(web.time, 'monotonic', lambda: now[0])
    monkeypatch.setitem(web.JOBS, 'timing-test', {
        'stage': 'translate', 'status': 'running', '_started': 100.0,
        '_stage_started': 100.0, 'done': 2, 'total': 3,
        'steps': [{'stage': 'translate', 'seconds': 0, 'status': 'running'}],
    })
    now[0] = 105.0
    web.update_job('timing-test', stage='new_pdf')
    snapshot = web.job_snapshot('timing-test')
    assert 'total' not in snapshot and 'done' not in snapshot
    assert snapshot['steps'][0]['seconds'] == 5
    now[0] = 112.0
    web.update_job('timing-test', stage='error', status='failed')
    now[0] = 200.0
    result = web.job_snapshot('timing-test')
    assert result['elapsed_seconds'] == 12
    assert result['steps'][-1] == {'stage': 'new_pdf', 'seconds': 7, 'status': 'failed'}
    result['steps'][-1]['seconds'] = 999
    assert web.job_snapshot('timing-test')['steps'][-1]['seconds'] == 7
