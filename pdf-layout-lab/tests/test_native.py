from zipfile import ZipFile

import fitz
import pytest
from docx import Document
from fastapi.testclient import TestClient

import app as web
import native_api as api
from engine import native_document as native


def sample(path):
    document = Document()
    document.add_paragraph('Original text').runs[0].bold = True
    document.add_table(rows=1, cols=1).cell(0, 0).text = 'Table text'
    document.save(path)
    return path


def test_translation_keeps_package_resources_and_table_structure(tmp_path):
    source = sample(tmp_path / 'source.docx')
    output = tmp_path / 'output.docx'
    report = native.translate_package(source, output, 'docx', lambda xs: ['Translated ' + x for x in xs])
    assert report['segments'] == 2
    with ZipFile(source) as before, ZipFile(output) as after:
        assert before.namelist() == after.namelist()
        for name in before.namelist():
            if name not in report['changed_parts']:
                assert before.read(name) == after.read(name)
    result = Document(output)
    assert result.paragraphs[0].runs[0].bold
    assert result.tables[0].cell(0, 0).text == 'Translated Table text'


@pytest.mark.parametrize('kind', ['pptx', 'hwpx'])
def test_native_xml_preserves_geometry_and_styles(tmp_path, kind):
    source, output = tmp_path / ('source.' + kind), tmp_path / ('output.' + kind)
    part = 'ppt/slides/slide1.xml' if kind == 'pptx' else 'Contents/section0.xml'
    required = 'ppt/presentation.xml' if kind == 'pptx' else 'Contents/content.hpf'
    with ZipFile(source, 'w') as archive:
        archive.writestr(required, '<root/>')
        archive.writestr(part, f'<p xmlns:n="{native.NS[kind]}" width="123"><n:t>Hello</n:t></p>')
        archive.writestr('media/test.bin', b'unchanged-image')
    native.translate_package(source, output, kind, lambda xs: ['Translated' for _ in xs])
    with ZipFile(output) as archive:
        assert b'width="123"' in archive.read(part)
        assert b'Translated' in archive.read(part)
        assert archive.read('media/test.bin') == b'unchanged-image'


def test_invalid_or_encrypted_package_is_rejected(tmp_path):
    path = tmp_path / 'bad.docx'
    path.write_bytes(b'not a docx')
    with pytest.raises(ValueError):
        native.validate(path, 'docx')
    path = sample(path)
    with ZipFile(path, 'a') as archive:
        archive.writestr('word/vbaProject.bin', b'macro')
    with pytest.raises(ValueError):
        native.validate(path, 'docx')


def test_native_routes_preserve_original_and_filter_failed_conversions(tmp_path, monkeypatch):
    monkeypatch.setattr(api, 'DATA', tmp_path / 'native-data')
    monkeypatch.setattr(api, 'convert', lambda *args: (_ for _ in ()).throw(ValueError('unavailable')))
    client = TestClient(web.app)
    source = sample(tmp_path / 'source.docx')
    response = client.post('/api/documents', files={'file': ('source.docx', source.read_bytes())})
    assert response.status_code == 200
    item = response.json()
    assert item['native'] and item['source_format'] == 'docx'
    prefix = '/api/native/' + item['id']
    item = client.post(prefix + '/process', json={}).json()
    assert [x['format'] for x in item['result']['exports']] == ['docx']
    run = item['latest_run']
    assert client.get(prefix + '/export/docx?run=' + run).content == source.read_bytes()
    assert client.get(prefix + '/export/pptx?run=' + run).status_code == 422
    assert client.get(prefix + '/export/pdf?run=' + run).status_code == 422
    assert client.get('/api/documents/' + item['id']).json()['native']


def test_preview_geometry_change_is_reported(tmp_path, monkeypatch):
    monkeypatch.setattr(api, 'DATA', tmp_path / 'native-data')
    def convert(source, output, kind):
        with fitz.open() as pdf:
            pdf.new_page(width=300, height=400)
            pdf.save(output)
    monkeypatch.setattr(api, 'convert', convert)
    client = TestClient(web.app)
    source = sample(tmp_path / 'source.docx')
    item = client.post('/api/documents', files={'file': ('source.docx', source.read_bytes())}).json()
    result = client.post('/api/native/' + item['id'] + '/process', json={}).json()
    assert result['result']['same_page_geometry']
    assert [x['format'] for x in result['result']['exports']] == ['docx', 'pdf']


def test_translation_failure_keeps_previous_result_and_source(tmp_path, monkeypatch):
    monkeypatch.setattr(api, 'DATA', tmp_path / 'native-data')
    monkeypatch.setattr(api, 'convert', lambda *args: (_ for _ in ()).throw(ValueError('unavailable')))
    monkeypatch.setattr(api.legacy, 'configuration', lambda: {'configured': False})
    client = TestClient(web.app)
    source = sample(tmp_path / 'source.docx')
    item = client.post('/api/documents', files={'file': ('source.docx', source.read_bytes())}).json()
    prefix = '/api/native/' + item['id']
    previous = client.post(prefix + '/process', json={}).json()
    assert client.post(prefix + '/process', json={'translate': True}).status_code == 400
    current = client.get(prefix).json()
    assert current['latest_run'] == previous['latest_run']
    assert not current['busy']
    assert (api.DATA / item['id'] / 'source.docx').read_bytes() == source.read_bytes()
