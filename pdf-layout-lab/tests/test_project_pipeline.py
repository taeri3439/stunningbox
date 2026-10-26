import json
import hashlib
from docx import Document
from engine import legacy
from test_lab import make_pdf


def test_project_snapshot_preserves_exact_docx_and_language(tmp_path, monkeypatch):
    storage = tmp_path / 'backend'
    item = storage / 'existing'
    item.mkdir(parents=True)
    source = make_pdf(item / 'source.pdf')
    (item / 'source.json').write_text(json.dumps({'source_lang': 'en'}))
    document = Document()
    document.add_table(rows=2, cols=2).cell(0, 0).text = 'Actual project table'
    document.save(item / 'output.ko.docx')
    monkeypatch.setenv('STORAGE_DIR', str(storage))
    result = legacy.project_output(source, 'en', 'ko', tmp_path / 'cache', tmp_path / 'result.docx')
    assert result['origin'] == 'project_snapshot'
    assert (tmp_path / 'result.docx').read_bytes() == (item / 'output.ko.docx').read_bytes()
    assert result['docx_sha256'] == hashlib.sha256((item / 'output.ko.docx').read_bytes()).hexdigest()
    assert Document(tmp_path / 'result.docx').tables[0].cell(0, 0).text == 'Actual project table'


def test_missing_snapshot_uses_full_pipeline_in_lab_storage(tmp_path, monkeypatch):
    legacy.setup()
    from pipeline import process
    from renderers import docx_renderer
    monkeypatch.setenv('STORAGE_DIR', str(tmp_path / 'empty-backend'))
    monkeypatch.setattr(legacy, 'configuration', lambda: {'provider': 'test'})
    source = make_pdf(tmp_path / 'source.pdf')
    calls = []
    def build(path, identifier, storage, **kwargs):
        assert storage == str(tmp_path / 'lab-cache')
        calls.append('analyze')
    monkeypatch.setattr(process, 'build_source_model', build)
    monkeypatch.setattr(process, 'get_translated_model', lambda *a, **kw: calls.append('translate'))
    monkeypatch.setattr(process, 'get_export_model', lambda *a: 'semantic-model')
    def render(model, path, assets):
        assert model == 'semantic-model'
        assert str(tmp_path / 'lab-cache') in assets
        calls.append('render')
        Document().save(path)
    monkeypatch.setattr(docx_renderer, 'render_docx', render)
    result = legacy.project_output(source, 'en', 'ko', tmp_path / 'lab-cache', tmp_path / 'out.docx')
    assert calls == ['analyze', 'translate', 'render']
    assert result['origin'] == 'project_pipeline'
