"""Render the semantic model to a .docx file.

OOXML is reflow-native and maps almost 1:1 to our model, so the skeleton survives:
headings -> heading styles, list items -> indented paragraphs with LITERAL markers
(never Word auto-numbering, which would re-number and break the original scheme),
tables -> real tables (with cell merges), figures -> embedded images, captions ->
small italic paragraphs.
"""
from __future__ import annotations

import logging
import os
import re
import statistics
from collections import Counter

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt, RGBColor
from docx.exceptions import InvalidSpanError
from PIL import Image

from pipeline.extractor import RENDER_DPI
from models import (
    CaptionBlock,
    DividerBlock,
    FigureBlock,
    HeadingBlock,
    ListItemBlock,
    ParagraphBlock,
    RunningBand,
    SemanticDocument,
    TableBlock,
)
from renderers._layout import scaled_pt as _scaled_pt, widths_would_cram as _widths_would_cram
from renderers import _ooxml

EMU_PER_IN = 914400
INK = RGBColor(0x1B, 0x1A, 0x17)  # near-black; overrides Word's blue heading colour
# Translated text (esp. KR->EN) runs longer than the source, so reproducing a loosely-set
# source's line pitch verbatim balloons the page count and breaks page-for-page alignment.
# We cap the rendered line pitch at a moderate value (a global compression policy) so the
# translation stays close to the source's pagination while remaining readable. The model's
# captured line_spacing stays truthful; this only affects rendering density.
MAX_LINE_PITCH = 1.4  # multiple of font size (source pitch); ~1.4x reads normally
_ALIGN = {
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}

# Target languages where translated text is denser than Latin source text (CJK), so a
# hard break per source page is a net fidelity win instead of the bloat that motivated
# dropping this in e7b3fbc (2026-06-15, see docstring on _page_parity_on below).
PARITY_TARGET_LANGS = {"ko", "ja", "zh"}


def _page_parity_on(doc: SemanticDocument) -> bool:
    """Anchored pagination: ON when translation direction shrinks text (latin->CJK)
    or the model explicitly opts in. Expanding directions (e.g. ko->en) stay OFF —
    see the 2026-06-15 bloat removal (e7b3fbc): breaking at every source-page boundary
    while the translation is EXPANDING makes an overflowing page desync output from
    source pages with no reliable realignment, plus the old mechanism injected a whole
    extra empty paragraph per break. This revival only enables the anchor where it helps
    (shrinking directions) and applies the break to the block's own first paragraph
    instead of a dedicated one — each anchor RESYNCS the two paginations locally rather
    than letting drift accumulate across the whole document."""
    if doc.page_parity is not None:
        return doc.page_parity
    return (doc.target_lang or "") in PARITY_TARGET_LANGS


def _mark_break(p, break_before: bool) -> None:
    """Apply a page-parity anchor to a block's own first paragraph — no extra paragraph
    is inserted (that was the old mechanism's bloat, see _page_parity_on)."""
    if p is not None and break_before:
        p.paragraph_format.page_break_before = True


def _widen_default_tab(d) -> None:
    """Push the document's default tab stop far to the right so it never sits BEFORE our
    explicit right tab stops (running header/footer + TOC dot-leaders). Word suppresses
    default tabs that precede a custom tab, so a single \\t jumps to the right margin — but
    LibreOffice/OnlyOffice honour the default tabs, landing the right element mid-line. A
    large default keeps the right tab the next stop in every renderer."""
    settings = d.settings.element
    dt = settings.find(qn("w:defaultTabStop"))
    if dt is None:
        dt = OxmlElement("w:defaultTabStop")
        settings.append(dt)
    dt.set(qn("w:val"), "20160")  # twips (~14in) — beyond any page width




def _bookmark_block(paragraph, block_id: str, _counter: list) -> None:
    """Wrap a paragraph in a named bookmark ('blk{block_id}') so the editor can scroll to
    this block via editorConfig.actionLink (used to land on a re-translated sentence).
    `block_id` is the stable string id (uuid or short alias). `w:id` must be a small
    integer (Word schema requirement) so we use a per-render counter passed in via a
    single-element list. Skips blocks with no block_id. Inserted after pPr so valid."""
    if not block_id:
        return
    numeric_id = _counter[0]
    _counter[0] += 1
    p = paragraph._p
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(numeric_id))
    start.set(qn("w:name"), f"blk{block_id}")
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(numeric_id))
    pPr = p.find(qn("w:pPr"))
    if pPr is not None:
        pPr.addnext(start)
    else:
        p.insert(0, start)
    p.append(end)


def render_docx(doc: SemanticDocument, out_path: str, asset_dir: str) -> str:
    d = Document()
    _widen_default_tab(d)
    if doc.title:
        d.core_properties.title = doc.title

    # Reproduce the source paper size + margins so the translation paginates like the
    # original instead of defaulting to US Letter (shorter than A4 → balloons the doc).
    sec = d.sections[0]
    _apply_page_geometry(sec, doc)
    # Right edge of the text column — where TOC page numbers and the right-hand running
    # header element are pinned via tab stops.
    content_width = sec.page_width - sec.left_margin - sec.right_margin

    # Reproduce the repeating page header/footer (dropped as chrome upstream) in Word's
    # section header/footer, which repeats automatically. The cover (first page) has no
    # running header/footer in the source, so suppress it there.
    body_font = _body_font(doc)  # so the header/footer match the body's sans/serif class
    if doc.running_header or doc.running_footer:
        sec.different_first_page_header_footer = True
    if doc.running_header:
        _render_band(sec.header, doc.running_header, content_width, font=body_font,
                     target_lang=doc.target_lang)
        if doc.running_header.image:
            _anchor_band_image(sec.header, doc.running_header.image, asset_dir, doc, at_top=True)
    if doc.running_footer:
        _render_band(sec.footer, doc.running_footer, content_width, footer=True, font=body_font,
                     target_lang=doc.target_lang)
        if doc.running_footer.image:
            _anchor_band_image(sec.footer, doc.running_footer.image, asset_dir, doc, at_top=False)

    # Cover/first-page full-bleed background: anchor it behind text in the first-page header.
    if doc.cover_background:
        bg_path = os.path.join(asset_dir, doc.cover_background)
        if os.path.exists(bg_path):
            sec.different_first_page_header_footer = True
            _ooxml.add_page_background(
                sec.first_page_header, bg_path,
                doc.page_width or float(sec.page_width.pt),
                doc.page_height or float(sec.page_height.pt),
            )

    # The document's body font size anchors the relative hierarchy: every block's
    # captured size is rendered proportionally to this, so the original's size tiers
    # survive even though absolute points may differ.
    body_pt = _body_size(doc)
    body_ls = _body_line_spacing(doc)

    cover_break_done = False
    parity = _page_parity_on(doc)
    if parity:
        # The default template's docDefaults add 10pt after every paragraph and a 1.15x
        # line base — measured to inflate dense designed layouts ~1.7x (source TOC pitch
        # 16.9pt rendered at 29.1pt), overflowing their anchored pages. In parity mode
        # underfill is harmless (bottom whitespace, like the source's own margins) while
        # overflow breaks the 1:1 page mapping, so bias density DOWN: zero the inherited
        # paragraph gap and let each block's measured line_spacing carry the pitch.
        normal = d.styles["Normal"].paragraph_format
        normal.space_after = Pt(0)
        normal.line_spacing = 1.0
    cover_page0 = doc.cover_background is not None  # avoid stacking on the cover break
    prev_page = None
    _bm_counter = [0]  # per-render bookmark numeric id counter (Word requires integer w:id)
    for block in doc.blocks:
        page = getattr(block, "page", None)

        # A full-page cover background can't share its page with body text: once we reach the
        # first body block (source page >= 1), break so the cover owns page 1 and the body
        # starts on page 2. Unconditional — independent of the direction-gated anchor below.
        if (doc.cover_background and not cover_break_done
                and page is not None and page >= 1):
            d.add_page_break()
            cover_break_done = True

        # Direction-gated page-parity anchoring: a page_break_before on the block's own
        # first paragraph at every source-page boundary — see _page_parity_on for why this
        # is safe where the old per-source-page break (e7b3fbc) was not. Skip the page
        # 0 -> 1 transition when a cover background is present: the cover-isolation break
        # just above already guarantees a fresh page there, so anchoring it too would
        # stack two page breaks at the same boundary.
        break_before = False
        if parity and page is not None and prev_page is not None and page != prev_page:
            break_before = not (cover_page0 and prev_page == 0)
        if page is not None:
            prev_page = page

        bid = getattr(block, "block_id", None)
        if isinstance(block, HeadingBlock):
            lvl = min(max(block.level, 0), 9)
            mcolor = getattr(block, "marker_color", None)
            mrun = None
            if block.marker and mcolor:
                # Marker in its own colour: two runs in ONE paragraph (an accent 'FIGURE 3'
                # label on a differently-coloured title). docx_sync round-trips on the
                # concatenated text, so keeping both runs in one paragraph is lossless.
                p = d.add_heading("", level=lvl)
                mrun = p.add_run(f"{block.marker} ")
                p.add_run(block.text or "")
            else:
                text = f"{block.marker} {block.text}".strip() if block.marker else block.text
                p = d.add_heading(text, level=lvl)
            _mark_break(p, break_before)
            if parity:
                # Word's Heading styles carry generous before/after spacing (10-12pt each)
                # that dense designed layouts don't have — 3-5 headings inflate a page past
                # its anchor. Bias down to a modest fixed gap; measured space_before /
                # page_space_before still override below in _apply_style.
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(2)
            _apply_style(p, block, body_pt, body_ls, parity, doc.target_lang)  # source hierarchy + spacing
            if mrun is not None:
                try:
                    mrun.font.color.rgb = RGBColor.from_string(mcolor)
                except ValueError:
                    pass  # malformed hex → leave marker in the block colour
            _bookmark_block(p, bid, _bm_counter)
        elif isinstance(block, ListItemBlock):
            p = d.add_paragraph()
            _mark_break(p, break_before)
            # A wrapped TOC entry (leader + page number at the END of its LAST line)
            # reads better as one leader row than as broken preserved lines — join and
            # let _toc_split decide; non-TOC multi-line items keep their line breaks.
            toc = _toc_split(" ".join(
                ln.strip() for ln in block.lines if ln.strip()
            )) if block.lines else _toc_split(block.text)
            # The marker's weight follows the block's captured weight (via _apply_style),
            # not a blanket bold — otherwise every TOC sub-entry number renders bold even
            # where the source is regular weight.
            p.add_run(f"{block.marker} ")
            if toc:
                _render_toc(p, toc[0], toc[1], content_width, _toc_depth(block.marker),
                           getattr(block, "indent", None))
            else:
                p.paragraph_format.left_indent = Inches(0.3 * (block.level + 1))
                _add_text(p, block)
            _apply_style(p, block, body_pt, body_ls, parity, doc.target_lang)
            if toc and parity:
                # TOC entries are single-line blocks, so their line_spacing falls back to
                # the body median — taller than a real contents column. Exact single
                # spacing biases toward underfill (safe: the anchor absorbs it) instead of
                # spilling a ~60-entry contents page past its 1:1 anchor.
                p.paragraph_format.line_spacing = 1.0
            _bookmark_block(p, bid, _bm_counter)
        elif isinstance(block, ParagraphBlock):
            p = d.add_paragraph()
            _mark_break(p, break_before)
            toc = None if block.lines else _toc_split(block.text)
            if toc:
                _render_toc(p, toc[0], toc[1], content_width,
                           source_indent=getattr(block, "indent", None))
            else:
                _add_text(p, block)
            _apply_style(p, block, body_pt, body_ls, parity, doc.target_lang)
            _bookmark_block(p, bid, _bm_counter)
        elif isinstance(block, TableBlock):
            if break_before:
                # A table has no leading paragraph of its own to carry the break, and an
                # ordinary paragraph would inject visible whitespace (the old bloat) — a
                # 1pt, zero-margin spacer paragraph anchors the break without adding space.
                sp = d.add_paragraph()
                sp.paragraph_format.page_break_before = True
                sp.paragraph_format.space_before = Pt(0)
                sp.paragraph_format.space_after = Pt(0)
                r = sp.add_run("")
                r.font.size = Pt(1)
            _add_table(d, block, doc.target_lang)
            if d.tables and d.tables[-1].rows and d.tables[-1].rows[0].cells:
                _bookmark_block(d.tables[-1].rows[0].cells[0].paragraphs[0], bid, _bm_counter)
        elif isinstance(block, FigureBlock):
            p = _add_figure(d, block, asset_dir, content_width)
            _mark_break(p, break_before)
        elif isinstance(block, DividerBlock):
            p = _add_divider(d, block)
            _mark_break(p, break_before)
        elif isinstance(block, CaptionBlock):
            p = d.add_paragraph()
            _mark_break(p, break_before)
            p.alignment = 1  # center
            _add_text(p, block, italic=True)
            _bookmark_block(p, bid, _bm_counter)
            _apply_style(p, block, body_pt, body_ls, parity, doc.target_lang)  # source size wins if known
            # A caption must never separate from the figure/table it describes onto the
            # next page — unconditional, direction-independent (unlike page_space_before).
            p.paragraph_format.keep_with_next = True

    d.save(out_path)
    return out_path


def _apply_page_geometry(sec, doc: SemanticDocument) -> None:
    """Set the section's paper size + margins from the captured source geometry. Each is
    optional — absent values keep python-docx's defaults (e.g. legacy source.json)."""
    if doc.page_width and doc.page_height:
        sec.page_width = Pt(doc.page_width)
        sec.page_height = Pt(doc.page_height)
    for attr, val in (
        ("left_margin", doc.margin_left), ("right_margin", doc.margin_right),
        ("top_margin", doc.margin_top), ("bottom_margin", doc.margin_bottom),
    ):
        if val:
            setattr(sec, attr, Pt(val))


def _add_text(p, block, italic: bool = False) -> None:
    """Add a block's content as runs. Preserved `lines` (semantic breaks) render with a
    line break between each; otherwise the single reflowed `text`."""
    lines = getattr(block, "lines", None)
    segments = lines if lines else [getattr(block, "text", "") or ""]
    for i, seg in enumerate(segments):
        if i > 0:
            p.add_run().add_break(WD_BREAK.LINE)
        run = p.add_run(seg)
        if italic:
            run.italic = True


def _body_line_spacing(doc: SemanticDocument) -> float | None:
    """Median line spacing of body paragraphs/list items, used where a block has none."""
    vals = [
        b.line_spacing
        for b in doc.blocks
        if isinstance(b, (ParagraphBlock, ListItemBlock)) and b.line_spacing
    ]
    return round(statistics.median(vals), 2) if vals else None


def _body_font(doc: SemanticDocument) -> str | None:
    """Most common block font family — applied to the running header/footer so it shares
    the body's sans/serif class instead of the renderer's default serif."""
    fonts = [b.font for b in doc.blocks if getattr(b, "font", None)]
    return Counter(fonts).most_common(1)[0][0] if fonts else None


def _body_size(doc: SemanticDocument) -> float | None:
    """Dominant body-text font size (pt): the most common size among paragraphs and
    list items. Falls back to any sized block, else None (→ no scaling)."""
    sizes = [
        round(b.font_size, 1)
        for b in doc.blocks
        if isinstance(b, (ParagraphBlock, ListItemBlock)) and b.font_size
    ]
    if not sizes:
        sizes = [round(b.font_size, 1) for b in doc.blocks if getattr(b, "font_size", None)]
    return Counter(sizes).most_common(1)[0][0] if sizes else None



_SANS_PAT = re.compile(
    r"gothic|arial|helvetica|franklin|verdana|tahoma|segoe|calibri|noto ?sans|malgun|"
    r"dotum|돋움|고딕",
    re.IGNORECASE,
)
_HANGUL_FONT_PAT = re.compile(
    # A Hangul-capable-name list (NOT a general CJK-capable list): Japanese-only families
    # like "MS PGothic" intentionally do NOT match this — for a ko target those still need
    # substitution, which is correct. ja/zh targets use their own substitute table below
    # regardless of this pattern (see _east_asia_for).
    r"[가-힣]|batang|바탕|malgun|맑은|gulim|굴림|dotum|돋움|gungsuh|궁서|noto.*(kr|cjk)",
    re.IGNORECASE,
)
# Any CJK codepoint present in a run's text: Hangul syllables, CJK Unified Ideographs
# (covers both Japanese kanji and Chinese hanzi), and Hiragana/Katakana. Broadened from
# Hangul-only because ja/zh are shipped target languages (translator.LANG_NAMES, frontend
# LANGS) whose translations also need the eastAsia substitution, not just ko's.
_CJK_TEXT_RE = re.compile(r"[가-힣一-鿿぀-ヿ]")

# Per-target substitute (sans, serif) pair for a latin-only source font that has no glyphs
# for the target script — Word/DS would otherwise silently fall back to a default serif.
# Keyed by target_lang; an other/unknown target with CJK text in the run defensively falls
# back to the ko pair (see _east_asia_for) rather than leaving the latin name unstyled.
_EASTASIA_SUBSTITUTE = {
    "ko": ("Malgun Gothic", "Batang"),
    "ja": ("MS Gothic", "MS Mincho"),
    "zh": ("Microsoft YaHei", "SimSun"),
}


def _east_asia_for(name: str, target_lang: str | None = None) -> str:
    """A latin-only source font (FranklinGothic, TimesNewRomanPSMT, ...) has no CJK
    glyphs — Word/DS would silently fall back to a default SERIF for the CJK text set in
    it. Substitute a class-matched family for the render's target language (sans/serif
    classification via _SANS_PAT on the source name), so a sans-serif report stays
    sans-serif in translation. A name that already identifies a Hangul-capable family
    (Batang, 맑은 고딕, ...) passes through unchanged — even for a ja/zh target, since a
    Hangul-capable font still carries real glyphs, unlike a latin-only one."""
    if _HANGUL_FONT_PAT.search(name):
        return name
    sans, serif = _EASTASIA_SUBSTITUTE.get(target_lang, _EASTASIA_SUBSTITUTE["ko"])
    return sans if _SANS_PAT.search(name) else serif


def _set_run_font(run, name: str, target_lang: str | None = None) -> None:
    """Set a run's font family for Latin AND East-Asian text. `run.font.name` alone does
    not style CJK glyphs in Word — w:rFonts/@w:eastAsia must be set too, else CJK text
    falls back to a default serif. The eastAsia class-matched substitute (_east_asia_for,
    language-aware via `target_lang`) only applies when this run's own text actually
    carries a CJK codepoint: a Latin-only run (an untranslated label, a document that was
    never translated to a CJK language) keeps eastAsia == the source name — there is no
    CJK glyph in that run for the substitute to matter to, and diverging from the source
    font attribute there would be a gratuitous, purely cosmetic no-op change to the XML."""
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    east_asia = _east_asia_for(name, target_lang) if _CJK_TEXT_RE.search(run.text or "") else name
    for attr, val in (
        ("w:ascii", name), ("w:hAnsi", name), ("w:eastAsia", east_asia), ("w:cs", name),
    ):
        rFonts.set(qn(attr), val)


def _apply_style(p, block, body_pt: float | None, body_ls: float | None = None,
                 parity: bool = False, target_lang: str | None = None) -> None:
    """Apply the source's relative size, weight, font family, colour, italic/underline,
    alignment, and line spacing to a paragraph. Weight is honoured BOTH ways (bold where
    the source is bold, un-bold where it isn't — so Word's bold heading styles don't
    over-emphasise). Colour honours a real source colour, else forces ink (kills Word's
    blue headings). `parity` gates `page_space_before` (T3): that field compensates for a
    page-parity hard break eating into a block's natural top margin, so it must only apply
    where such breaks are actually inserted (see _page_parity_on). `target_lang` is threaded
    the same way into `_set_run_font` so the eastAsia substitute (when a run's text carries
    a CJK codepoint) picks the language-matched family instead of always assuming Korean."""
    out = _scaled_pt(getattr(block, "font_size", None), body_pt)
    bold = getattr(block, "bold", None)
    italic = getattr(block, "italic", None)
    underline = getattr(block, "underline", None)
    font = getattr(block, "font", None)
    color = getattr(block, "color", None)
    for run in p.runs:
        if out is not None:
            run.font.size = Pt(out)
        if bold is True:
            run.bold = True
        elif bold is False:
            run.bold = False
        if italic:
            run.italic = True
        if underline:
            run.underline = True
        if font:
            _set_run_font(run, font, target_lang)
        run.font.color.rgb = RGBColor.from_string(color) if color else INK
    align = getattr(block, "align", None)
    if align in _ALIGN:
        p.alignment = _ALIGN[align]
    # Measured source indent (pt) overrides any style/level default so the original
    # left offset survives.
    indent = getattr(block, "indent", None)
    if indent:
        p.paragraph_format.left_indent = Pt(indent)
    # line_spacing is captured as the PDF baseline-to-baseline ÷ font ratio; Word's
    # multiple is of single-line height (~1.15×font), so divide to match visually.
    ls = getattr(block, "line_spacing", None) or body_ls
    if ls:
        ls = min(ls, MAX_LINE_PITCH)  # moderate global compression (translated text expands)
        p.paragraph_format.line_spacing = max(1.0, min(2.0, ls / 1.15))
    # Measured top gap (pt), cover-page-only (see models.HeadingBlock.space_before). Absent
    # on every non-cover block and on old source.json (field defaults to None) → no-op, so
    # ordinary paragraph flow is unchanged everywhere else. block.space_before (cover rhythm)
    # wins when set; else page_space_before (sparse-page rhythm, T3) applies but ONLY in
    # parity mode, since that field compensates for a page-parity break that only exists
    # there — outside parity mode there's no break to compensate for.
    space_before = getattr(block, "space_before", None)
    if space_before is None and parity:
        space_before = getattr(block, "page_space_before", None)
    if space_before:
        p.paragraph_format.space_before = Pt(space_before)


# A TOC / index entry: "<title> ·········· <page>". The leader is literal dots or
# middle-dots (U+00B7); we drop them and re-create the leader with a right tab stop so
# the page number pins to the right margin and the entry never wraps.
_TOC_RE = re.compile(r"^(.+?)\s*[·.]{2,}[·.\s]*(\d+)\s*$")


def _toc_split(text: str | None):
    if not text:
        return None
    m = _TOC_RE.match(text)
    if not m:
        return None
    title = m.group(1).strip()
    return (title, m.group(2)) if title else None


def _toc_depth(marker: str | None) -> int:
    """Nesting depth from a TOC marker's dot-separated numeric parts, e.g. "12." -> 0
    (top level), "12.2" -> 1, "12.2.1" -> 2. A trailing dot is a delimiter, not a part;
    a missing/empty marker (e.g. a TOC line captured as a plain ParagraphBlock, which has
    no marker field) has no depth information -> top level."""
    if not marker:
        return 0
    parts = [p for p in marker.rstrip(".").split(".") if p]
    return max(0, len(parts) - 1)


def _render_toc(p, title: str, page: str, content_width, depth: int = 0,
                source_indent: float | None = None) -> None:
    """Render a TOC entry on one line: title (left) + dotted leader + page (right).
    `depth` (see _toc_depth) indents nested sub-entries so the hierarchy reads visually.
    The leader's right tab stop stays pinned at content_width regardless of indent, so
    every level's page number still lines up flush at the right margin.

    INDENT PRIORITY (explicit, not an accident of call order): `source_indent` — the
    block's source-measured left indent (pt), routine for reconciler/HWP-extracted list
    items — is real fidelity data and WINS when present. The depth-based indent below is
    only the FALLBACK for blocks with no measured indent. We enforce that here by simply
    not setting the depth indent when `source_indent` is present, rather than relying on
    _apply_style's `if indent: ... left_indent = Pt(indent)` to clobber it afterwards —
    the caller passes `getattr(block, "indent", None)` as `source_indent`; _apply_style
    (untouched) still applies the measured indent itself in both cases."""
    p.paragraph_format.tab_stops.add_tab_stop(
        content_width, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
    )
    if not source_indent:
        p.paragraph_format.left_indent = Inches(min(depth, 4) * 0.18)
    p.add_run(title)
    p.add_run("\t" + page)


def _border_sz(thickness_pt: float | None) -> str:
    """Word border size is in eighths of a point. Map a source line weight (pt), default
    0.75pt, clamped to Word's [0.25pt, 12pt] range."""
    pt = thickness_pt if thickness_pt else 0.75
    return str(max(2, min(96, round(pt * 8))))


def _paragraph_border(p, edge: str, thickness=None, color=None) -> None:
    """Draw a single horizontal rule on one edge ('bottom'/'top') of a paragraph, at the
    source line weight and colour."""
    pPr = p._p.get_or_add_pPr()
    pbdr = pPr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        pPr.append(pbdr)
    b = OxmlElement(f"w:{edge}")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), _border_sz(thickness))
    b.set(qn("w:space"), "1")
    b.set(qn("w:color"), color or "1B1A17")
    pbdr.append(b)


def _add_divider(d: "Document", block: DividerBlock):
    """Reproduce a horizontal rule as an empty paragraph with a bottom border at the
    source line weight and colour. Returns the paragraph (so a page-parity anchor can be
    marked on it)."""
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    _paragraph_border(p, "bottom", block.thickness, block.color)
    return p


def _clear_table_borders(table) -> None:
    """Make a table invisible (used for the header/footer layout table)."""
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tblPr.append(borders)


def _clear_cell_borders(tc) -> None:
    """Remove a single cell's borders (overriding the Table Grid style) so an empty
    flowchart-grid cell renders blank instead of as a boxed empty cell."""
    tcPr = tc._tc.get_or_add_tcPr()
    existing = tcPr.find(qn("w:tcBorders"))
    if existing is not None:
        tcPr.remove(existing)
    tb = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        tb.append(el)
    tcPr.append(tb)


def _cell_band_rule(cell, edge: str, thickness, color) -> None:
    """Draw the band's horizontal rule as a single cell border (so it doesn't depend on
    paragraph ordering)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tb = tcPr.find(qn("w:tcBorders"))
    if tb is None:
        tb = OxmlElement("w:tcBorders")
        tcPr.append(tb)
    b = OxmlElement(f"w:{edge}")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), _border_sz(thickness))
    b.set(qn("w:color"), color or "1B1A17")
    tb.append(b)


def _render_band(hf, band: RunningBand, content_width, footer: bool = False,
                 font: str | None = None, target_lang: str | None = None) -> None:
    """Render a running header/footer as a borderless 1-row table: left | (center) | right.
    A table pins each element to its edge in EVERY renderer; tab stops render inconsistently
    across Word/OnlyOffice/LibreOffice (the right element drifts to mid-line). The band's
    horizontal rule is drawn as a cell border (header → bottom edge, footer → top). A
    page_number side gets a Word PAGE field (auto page number) appended in its cell.
    `target_lang` is threaded into `_set_run_font` (same as `_apply_style`/`_fill_cell`) so
    a ja/zh running band gets the language-matched eastAsia substitute instead of always
    assuming Korean."""
    sides = [(band.left, "left"), (band.center, "center"), (band.right, "right")]
    active = [(t, a) for t, a in sides if t or (band.page_number == a)]
    if not active:
        if band.rule:
            _paragraph_border(hf.paragraphs[0], "top" if footer else "bottom",
                              band.rule_size, band.rule_color)
        return
    table = hf.add_table(rows=1, cols=len(active), width=content_width)
    _clear_table_borders(table)
    col_w = int(content_width) // len(active)
    for i, (text, align) in enumerate(active):
        cell = table.cell(0, i)
        cell.width = Emu(col_w)
        cp = cell.paragraphs[0]
        cp.alignment = _ALIGN[align]
        if text:
            run = cp.add_run(text)
            run.font.color.rgb = INK
            if font:
                _set_run_font(run, font, target_lang)
        if band.page_number == align:
            _append_page_number(cp, band.page_number_format)
        if band.rule:
            _cell_band_rule(cell, "top" if footer else "bottom", band.rule_size, band.rule_color)
    # Move the table ahead of the (required) trailing empty paragraph so it sits at the top.
    hf.paragraphs[0]._p.addprevious(table._tbl)


def _append_page_number(paragraph, fmt: str | None) -> None:
    """Emit page-number chrome: literal runs around a PAGE field ({n}) and, when the
    source showed a total ("Page 3 of 32"), a NUMPAGES field ({m})."""
    fmt = fmt or "{n}"
    pos = 0
    for m in re.finditer(r"\{[nm]\}", fmt):
        if m.start() > pos:
            paragraph.add_run(fmt[pos:m.start()])
        (_ooxml.add_page_field if m.group() == "{n}" else _ooxml.add_numpages_field)(paragraph)
        pos = m.end()
    if pos < len(fmt):
        paragraph.add_run(fmt[pos:])


def _anchor_band_image(hf, image_name: str, asset_dir: str, doc: "SemanticDocument",
                       at_top: bool) -> None:
    """Anchor a captured decoration strip full-width behind the band text. The strip was
    rendered across the full page width, so width = page_width and height preserves the PNG
    aspect; a footer strip is pinned to the page bottom."""
    path = os.path.join(asset_dir, image_name)
    if not (image_name and os.path.exists(path) and doc.page_width and doc.page_height):
        return
    try:
        with Image.open(path) as im:
            px_w, px_h = im.size
    except Exception as e:  # noqa: BLE001
        logging.getLogger(__name__).warning(
            "band decoration %s unreadable — dropped from output: %s", path, e)
        return
    if not px_w:
        return
    width_pt = float(doc.page_width)
    height_pt = width_pt * (px_h / px_w)
    y_offset = 0.0 if at_top else max(0.0, float(doc.page_height) - height_pt)
    _ooxml.add_band_image(hf, path, width_pt, height_pt, y_offset)


def _cell_has_text(cell) -> bool:
    return bool((cell.text or "").strip()) or bool(cell.lines)


def _content_pt(d: "Document") -> float:
    """Usable text width (points) of the first section — table widths are sized against it."""
    sec = d.sections[0]
    try:  # Length values are EMU (int subclass); 12700 EMU = 1 pt
        return (int(sec.page_width) - int(sec.left_margin) - int(sec.right_margin)) / 12700
    except (TypeError, ValueError):
        return 455.0  # A4 text width fallback if the section has no geometry



def _set_col_widths(d: "Document", table, fractions: list[float]) -> None:
    """Fix the table layout and size each column to its relative width (recovered from the
    source's grid lines), so the vertical lines fall where the original has them instead of
    being evenly distributed."""
    content_pt = _content_pt(d)
    dxa = [max(1, int(f * content_pt * 20)) for f in fractions]  # twentieths of a point
    table.autofit = False
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:type"), "dxa")
    tblW.set(qn("w:w"), str(sum(dxa)))
    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for col, w in zip(grid.findall(qn("w:gridCol")), dxa):
            col.set(qn("w:w"), str(w))


def _add_table(d: "Document", block: TableBlock, target_lang: str | None = None) -> None:
    if not block.rows:
        return
    ncols = max((sum(max(1, c.colspan) for c in row) for row in block.rows), default=0)
    nrows = len(block.rows)
    if ncols == 0:
        return
    # Flowchart-style tables: HWP draws test-sequence diagrams as a grid of boxes joined
    # by arrows; hwp5html keeps the boxes as mostly-empty cells and the arrows are lost.
    # With full borders such a table reads as a broken empty grid, so when a table is large
    # and predominantly empty we strip the empty cells' borders, leaving just the labelled
    # boxes (closer to the original diagram than a grid of blank cells).
    cells_flat = [c for row in block.rows for c in row]
    sparse = len(cells_flat) >= 12 and sum(
        1 for c in cells_flat if not _cell_has_text(c)) > len(cells_flat) * 0.5
    # Wide reference tables (e.g. 별표2 합격품질수준표, 9 cols) don't fit portrait A4 at the
    # body size and wrap into a cramped mess; shrink the cell text so columns stay legible
    # (mirrors how the source uses a smaller face for these tables).
    cell_pt = 8 if ncols >= 8 else 9 if ncols >= 6 else None
    table = d.add_table(rows=nrows, cols=ncols)
    table.style = "Table Grid"
    # Real column widths recovered from the source's grid lines — but only when they still
    # fit the (now-translated) text. If a source-narrow column can't hold its longest
    # translated word, forcing the width breaks it one-glyph-per-line; autofit instead.
    if (block.col_widths and len(block.col_widths) == ncols
            and not _widths_would_cram(block.rows, block.col_widths, ncols,
                                       _content_pt(d), cell_pt or 10.0)):
        _set_col_widths(d, table, block.col_widths)
    else:
        table.autofit = True  # long (translated) cell text rewraps instead of clipping
    # Track grid cells already covered by a span from an earlier row/cell. Without this a
    # rowspan from above is ignored and the next row's cells land in the merged region,
    # shifting columns and corrupting the table.
    occupied: set[tuple[int, int]] = set()
    for r, row in enumerate(block.rows):
        c = 0
        last = None  # (tc, col_start) of the last placed cell, for trailing-gap extend
        for cell in row:
            while c < ncols and (r, c) in occupied:  # skip cells taken by spans
                c += 1
            if c >= ncols:
                break
            cs = min(max(1, cell.colspan), ncols - c)
            rs = min(max(1, cell.rowspan), nrows - r)
            tc = table.cell(r, c)
            if rs > 1 or cs > 1:
                try:
                    tc = tc.merge(table.cell(r + rs - 1, c + cs - 1))
                except InvalidSpanError:
                    # An irregular detected grid can ask for a non-rectangular merge; leave
                    # the cells unmerged rather than aborting the whole document render.
                    pass
            for rr in range(r, r + rs):
                for cc in range(c, c + cs):
                    occupied.add((rr, cc))
            _fill_cell(tc, cell, font_pt=cell_pt, target_lang=target_lang)
            if sparse and not _cell_has_text(cell):
                _clear_cell_borders(tc)
            last = (tc, c)
            c += cs
        # A label/value form row that stops short of the right edge: extend the last cell
        # to fill the remaining free columns (matches forms where the value spans to edge).
        if last and c < ncols and all((r, cc) not in occupied for cc in range(c, ncols)):
            tc, c0 = last
            try:
                tc = tc.merge(table.cell(r, ncols - 1))
            except InvalidSpanError:
                # The last cell already carries a rowspan, so extending it to the row's edge
                # would be L-shaped (non-rectangular). Skip the fill-extend; never crash.
                pass
            for cc in range(c, ncols):
                occupied.add((r, cc))
    # Never let a row split mid-cell across a page boundary (orphan fragments like
    # "Only)"/"Value)" at the top of the next page) — Word instead pushes the whole row
    # to the next page. Repeat the header row on each page when the first row IS a header
    # (a cosmetic nice-to-have; harmless no-op for headerless tables).
    for r in range(nrows):
        trPr = table.rows[r]._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
    if block.rows and any(c.is_header for c in block.rows[0]):
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    if block.caption:
        cap = d.add_paragraph()
        cap.alignment = 1
        run = cap.add_run(block.caption)
        run.italic = True
        run.font.size = Pt(9)


def _fill_cell(tc, cell, font_pt: float | None = None, target_lang: str | None = None) -> None:
    """Write a cell's content: vertical (rotated) text, preserved multi-line text, or a
    single string. Applies the source cell's shading, alignment, vertical alignment, and
    font/colour/italic. Header cells are bolded. `font_pt` shrinks the text for wide tables
    that would otherwise wrap into a cramped mess in portrait A4. `target_lang` is threaded
    into `_set_run_font` (same as `_apply_style`) so a ja/zh table cell gets the
    language-matched eastAsia substitute instead of always assuming Korean."""
    tcp = tc.paragraphs[0]
    if getattr(cell, "vertical", False):
        _set_cell_vertical(tc)
        tcp.add_run(cell.text)
    elif cell.lines:  # preserved multi-line cell
        for i, ln in enumerate(cell.lines):
            if i > 0:
                tcp.add_run().add_break(WD_BREAK.LINE)
            tcp.add_run(ln)
    else:
        tcp.add_run(cell.text)
    # Per-cell formatting
    if getattr(cell, "shading", None):
        _set_cell_shading(tc, cell.shading)
    if getattr(cell, "valign", None):
        _set_cell_valign(tc, cell.valign)
    if getattr(cell, "align", None) in _ALIGN:
        tcp.alignment = _ALIGN[cell.align]
    font = getattr(cell, "font", None)
    color = getattr(cell, "color", None)
    italic = getattr(cell, "italic", None)
    for run in tcp.runs:
        if font:
            _set_run_font(run, font, target_lang)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        if italic:
            run.italic = True
        if cell.is_header:
            run.bold = True
        if font_pt:
            run.font.size = Pt(font_pt)


_VALIGN = {"top": "top", "center": "center", "bottom": "bottom"}


def _set_cell_shading(tc, hex_color: str) -> None:
    """Set a cell's background fill via w:tcPr/w:shd."""
    tcPr = tc._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)


def _set_cell_valign(tc, val: str) -> None:
    """Set a cell's vertical text alignment via w:tcPr/w:vAlign."""
    if val not in _VALIGN:
        return
    tcPr = tc._tc.get_or_add_tcPr()
    va = tcPr.find(qn("w:vAlign"))
    if va is None:
        va = OxmlElement("w:vAlign")
        tcPr.append(va)
    va.set(qn("w:val"), val)


def _set_cell_vertical(tc) -> None:
    """Render a cell's text vertically (rotated 90°, reading bottom→top) — for source
    labels that were written vertically, e.g. a tall narrow row-spanning header."""
    tcPr = tc._tc.get_or_add_tcPr()
    td = OxmlElement("w:textDirection")
    td.set(qn("w:val"), "btLr")
    tcPr.append(td)


def _embed_figure_group(d: "Document", block: FigureBlock, asset_dir: str, content_width):
    """Try to render `block` as an editable OOXML group (base picture + text boxes).
    Returns the paragraph, or None to signal the caller to fall back to the composite."""
    base_path = os.path.join(asset_dir, block.base_image_ref)
    if not (block.overlays and block.base_image_ref and os.path.exists(base_path)):
        return None
    p = None
    try:
        with Image.open(base_path) as im:
            px_w, px_h = im.size
        w_emu = int(block.width / 72.0 * EMU_PER_IN) if block.width else int(px_w / RENDER_DPI * EMU_PER_IN)
        h_emu = int(w_emu * (px_h / px_w))
        cw_emu = int(content_width)  # Emu
        if w_emu > cw_emu:
            h_emu = int(h_emu * (cw_emu / w_emu)); w_emu = cw_emu
        p = d.add_paragraph()
        run = p.add_run(); host = p.add_run()
        pic = host.add_picture(base_path, width=Emu(w_emu))
        rid = pic._inline.graphic.graphicData.pic.blipFill.blip.get(qn("r:embed"))
        host._r.getparent().remove(host._r)
        from renderers.figure_group import build_figure_group_xml
        xml = build_figure_group_xml(block, rid, w_emu, h_emu, px_w, px_h)
        run._r.append(parse_xml(xml))
        p.alignment = _ALIGN.get(block.align, WD_ALIGN_PARAGRAPH.LEFT)
        return p
    except Exception:  # noqa: BLE001 — any failure → caller falls back to composite
        if p is not None:
            p._element.getparent().remove(p._element)
        return None


def _add_figure(d: "Document", block: FigureBlock, asset_dir: str, content_width):
    """Returns the figure's paragraph (so a page-parity anchor can be marked on it), or
    None if nothing was rendered (missing/unreadable asset)."""
    if not block.image_ref:
        return None
    grouped = _embed_figure_group(d, block, asset_dir, content_width)
    if grouped is not None:
        return grouped
    path = os.path.join(asset_dir, block.image_ref)
    if not os.path.exists(path):
        return None
    # Render at the SOURCE size (pt → in), never stretched to the column. Fall back to the
    # clip's pixel size at the render DPI, which is also 1:1 with the original region.
    if block.width:
        w = Inches(block.width / 72.0)
    else:
        try:
            with Image.open(path) as im:
                w = Inches(im.size[0] / RENDER_DPI)
        except Exception:  # noqa: BLE001
            w = content_width
    if w > content_width:  # only shrink to fit; never enlarge
        w = content_width
    try:
        p = d.add_paragraph()
        run = p.add_run()
        run.add_picture(path, width=w)
    except Exception:  # noqa: BLE001
        return None
    if block.wrap in ("left", "right"):
        try:
            _float_inline(run, block.wrap)  # let neighbouring text wrap beside it
        except Exception:  # noqa: BLE001
            p.alignment = _ALIGN.get(block.align, WD_ALIGN_PARAGRAPH.LEFT)
    else:
        p.alignment = _ALIGN.get(block.align, WD_ALIGN_PARAGRAPH.LEFT)
    return p


def _float_inline(run, side: str) -> None:
    """Convert an inline picture (in `run`) into a floating anchor with square text-wrap,
    aligned to `side` ('left'/'right') of the column, so adjacent text flows beside it."""
    drawing = run._r.find(qn("w:drawing"))
    inline = drawing.find(qn("wp:inline")) if drawing is not None else None
    if inline is None:
        return
    inline.tag = qn("wp:anchor")
    for k, v in {
        "simplePos": "0", "relativeHeight": "251658240", "behindDoc": "0",
        "locked": "0", "layoutInCell": "1", "allowOverlap": "1",
        "distT": "0", "distB": "0", "distL": "114300", "distR": "114300",
    }.items():
        inline.set(k, v)
    extent = inline.find(qn("wp:extent"))
    simple_pos = OxmlElement("wp:simplePos")
    simple_pos.set("x", "0")
    simple_pos.set("y", "0")
    pos_h = OxmlElement("wp:positionH")
    pos_h.set("relativeFrom", "column")
    align_h = OxmlElement("wp:align")
    align_h.text = side
    pos_h.append(align_h)
    pos_v = OxmlElement("wp:positionV")
    pos_v.set("relativeFrom", "paragraph")
    off = OxmlElement("wp:posOffset")
    off.text = "0"
    pos_v.append(off)
    extent.addprevious(simple_pos)
    extent.addprevious(pos_h)
    extent.addprevious(pos_v)
    wrap = OxmlElement("wp:wrapSquare")
    wrap.set("wrapText", "bothSides")
    inline.find(qn("wp:docPr")).addprevious(wrap)
