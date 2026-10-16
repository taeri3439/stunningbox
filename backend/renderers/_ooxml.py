"""Raw-OOXML builders python-docx doesn't expose: behind-text page-anchored images
(full-page covers + full-width header/footer bands) and PAGE fields. Kept separate so
docx_renderer stays focused."""
from __future__ import annotations

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


def _anchor_behind_text(hf, image_path: str, cx_emu: int, cy_emu: int,
                        x_emu: int, y_emu: int) -> None:
    """Anchor `image_path` as a floating, behind-text picture in header/footer `hf`, sized
    cx_emu × cy_emu and positioned at (x_emu, y_emu) relative to the page."""
    p = hf.paragraphs[0] if hf.paragraphs else hf.add_paragraph()
    run = p.add_run()
    run.add_picture(image_path)  # inline first; converted to a page-anchored drawing below
    drawing = run._r.find(qn("w:drawing"))
    inline = drawing.find(qn("wp:inline"))
    inline.tag = qn("wp:anchor")
    for k, v in {
        "simplePos": "0", "relativeHeight": "0", "behindDoc": "1", "locked": "0",
        "layoutInCell": "1", "allowOverlap": "1",
        "distT": "0", "distB": "0", "distL": "0", "distR": "0",
    }.items():
        inline.set(k, v)
    extent = inline.find(qn("wp:extent"))
    extent.set("cx", str(cx_emu))
    extent.set("cy", str(cy_emu))
    for ext in inline.iter(qn("a:ext")):
        ext.set("cx", str(cx_emu))
        ext.set("cy", str(cy_emu))
    simple_pos = OxmlElement("wp:simplePos")
    simple_pos.set("x", "0")
    simple_pos.set("y", "0")
    pos_h = OxmlElement("wp:positionH")
    pos_h.set("relativeFrom", "page")
    off_h = OxmlElement("wp:posOffset")
    off_h.text = str(x_emu)
    pos_h.append(off_h)
    pos_v = OxmlElement("wp:positionV")
    pos_v.set("relativeFrom", "page")
    off_v = OxmlElement("wp:posOffset")
    off_v.text = str(y_emu)
    pos_v.append(off_v)
    extent.addprevious(simple_pos)
    extent.addprevious(pos_h)
    extent.addprevious(pos_v)
    wrap = OxmlElement("wp:wrapNone")
    inline.find(qn("wp:docPr")).addprevious(wrap)


def add_page_background(header, image_path: str, page_width_pt: float, page_height_pt: float) -> None:
    """Anchor `image_path` as a full-page image behind text in `header` (use a section's
    first-page header for a cover). The image fills the whole page and sits behind the text."""
    _anchor_behind_text(header, image_path,
                        int(Pt(page_width_pt)), int(Pt(page_height_pt)), 0, 0)


def add_band_image(hf, image_path: str, width_pt: float, height_pt: float,
                   y_offset_pt: float) -> None:
    """Anchor `image_path` as a full-width strip behind text in header/footer `hf`, sized
    width_pt × height_pt, pinned x=0 and y=y_offset_pt from the page top (0 for a header
    band, page_height−height for a footer band)."""
    _anchor_behind_text(hf, image_path,
                        int(Pt(width_pt)), int(Pt(height_pt)), 0, int(Pt(y_offset_pt)))


def add_page_field(paragraph) -> None:
    """Append a Word PAGE field (auto page number) to `paragraph` as a fldChar run sequence."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def add_numpages_field(paragraph) -> None:
    """Append a Word NUMPAGES field (total page count) to `paragraph` as a fldChar run
    sequence — pairs with add_page_field for "Page N of M" chrome."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "NUMPAGES"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)
