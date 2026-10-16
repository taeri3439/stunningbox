from docx import Document
from docx.oxml.ns import qn
from models import SemanticDocument, HeadingBlock
from renderers.docx_renderer import render_docx
from pipeline.docx_sync import apply_docx_edits_to_model

def _render(tmp_path, block):
    doc = SemanticDocument(blocks=[block], target_lang="ko")
    out = str(tmp_path / "o.docx"); render_docx(doc, out, str(tmp_path))
    return out

def _runs_with_color(docx_path):
    d = Document(docx_path)
    p = d.paragraphs[0]
    res = []
    for r in p.runs:
        c = r.font.color.rgb
        res.append((r.text, str(c) if c else None))
    return res

def test_marker_color_renders_two_runs(tmp_path):
    b = HeadingBlock(block_id="H", level=4, marker="FIGURE 3", marker_color="D4B031",
                     text="1차 인터뷰의 구분", color="0D5775", font_size=13.0)
    runs = _runs_with_color(_render(tmp_path, b))
    # marker run gold, title run teal
    joined = "".join(t for t, _ in runs)
    assert "FIGURE 3" in joined and "1차 인터뷰의 구분" in joined
    assert any(t.strip() == "FIGURE 3" and (c or "").upper() == "D4B031" for t, c in runs)
    assert any("인터뷰" in t and (c or "").upper() == "0D5775" for t, c in runs)

def test_no_marker_color_single_run(tmp_path):
    b = HeadingBlock(block_id="H", level=3, marker="2.1.2.2", marker_color=None,
                     text="1차 자료의 구분", color="007C9E", font_size=12.0)
    d = Document(_render(tmp_path, b))
    assert "2.1.2.2 1차 자료의 구분" in d.paragraphs[0].text  # combined, one visual colour

def test_marker_render_roundtrips_through_sync(tmp_path):
    b = HeadingBlock(block_id="H", level=4, marker="FIGURE 3", marker_color="D4B031",
                     text="타이틀", color="0D5775", font_size=13.0)
    out = _render(tmp_path, b)
    model = SemanticDocument(blocks=[b], target_lang="ko")
    synced = apply_docx_edits_to_model(model, out)
    h = [x for x in synced.blocks if isinstance(x, HeadingBlock)][0]
    assert h.text == "타이틀"  # marker stripped, title intact (round-trip lossless)
