from docx import Document

from pipeline.docx_edit import apply_replacements, count_replacements
from pipeline.text_replace import ReplaceRule


def _rule(find, replace, match="smart"):
    return ReplaceRule(find=find, replace=replace, match=match)


def test_replace_in_single_run_paragraph():
    d = Document()
    d.add_paragraph("Rated Load is high")
    counts = apply_replacements(d, [_rule("Rated Load", "Load Rating")], "en")
    assert d.paragraphs[0].text == "Load Rating is high"
    assert counts == {"Rated Load": 1}


def test_replace_across_run_boundary():
    d = Document()
    p = d.add_paragraph()
    p.add_run("Ra")
    p.add_run("ted Load")  # "Rated" spans two runs → caught by the cross-run pass
    counts = apply_replacements(d, [_rule("Rated", "Nominal")], "en")
    assert p.text == "Nominal Load"
    assert counts == {"Rated": 1}


def test_replace_in_table_cell():
    d = Document()
    t = d.add_table(rows=1, cols=1)
    t.rows[0].cells[0].paragraphs[0].add_run("정격하중")
    counts = apply_replacements(d, [_rule("정격하중", "정격")], "ko")
    assert t.rows[0].cells[0].text == "정격"
    assert counts == {"정격하중": 1}


def test_count_does_not_mutate():
    d = Document()
    d.add_paragraph("load load")
    counts = count_replacements(d, [_rule("load", "x")], "en")
    assert counts == {"load": 2}
    assert d.paragraphs[0].text == "load load"


def test_no_double_replace_expanding_cjk_smart():
    # find is a substring of replace; smart+CJK has no word boundary → must still replace once.
    d = Document()
    d.add_paragraph("하중")
    counts = apply_replacements(d, [_rule("하중", "하중량", match="smart")], "ko")
    assert d.paragraphs[0].text == "하중량"
    assert counts == {"하중": 1}


def test_no_double_replace_expanding_exact():
    d = Document()
    d.add_paragraph("load")
    counts = apply_replacements(d, [_rule("load", "loaded", match="exact")], "en")
    assert d.paragraphs[0].text == "loaded"
    assert counts == {"load": 1}
