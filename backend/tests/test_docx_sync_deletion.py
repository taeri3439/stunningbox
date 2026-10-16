"""Deletion tests for apply_docx_edits_to_model (T9 rewrite).

Deletion is now structural: a bookmark-less paragraph ⇒ that block is DROPPED from
the rebuilt model (absent id ⇒ not in the new blocks list). The old tombstone flag
is gone; these tests verify the drop-path and the survivors-intact path.
"""
import os
import models
from pipeline.docx_sync import apply_docx_edits_to_model
from renderers.docx_renderer import render_docx
from docx import Document
from docx.oxml.ns import qn


def _delete_paragraph_by_bookmark(path, name):
    """Simulate an editor deletion: remove the w:p that contains bookmark `name`."""
    d = Document(path)
    for bm in list(d.element.body.iter(qn("w:bookmarkStart"))):
        if bm.get(qn("w:name")) == name:
            p = bm
            while p is not None and p.tag != qn("w:p"):
                p = p.getparent()
            if p is not None:
                p.getparent().remove(p)
            break
    d.save(path)


def test_missing_bookmark_drops_block(tmp_path):
    """A text block whose bookmark was deleted is dropped from synced.blocks."""
    m = models.SemanticDocument(source_lang="ko", blocks=[
        models.ParagraphBlock(text="first",  block_id="id0", source_span_ids=["s0"]),
        models.ParagraphBlock(text="second", block_id="id1", source_span_ids=["s1"]),
        models.ParagraphBlock(text="third",  block_id="id2", source_span_ids=["s2"]),
    ])
    out = os.path.join(str(tmp_path), "doc.docx")
    render_docx(m, out, str(tmp_path))
    _delete_paragraph_by_bookmark(out, "blkid1")  # user deletes the 2nd paragraph (block_id="id1")

    synced = apply_docx_edits_to_model(m, out)

    ids = [b.block_id for b in synced.blocks]
    assert "id1" not in ids             # dropped — absent from rebuilt list
    assert "id0" in ids
    assert "id2" in ids
    assert len(synced.blocks) == 2      # exactly two blocks remain

    # Survivor text is still synced
    by_id = {b.block_id: b for b in synced.blocks}
    assert by_id["id2"].text == "third"


def test_all_present_leaves_all_blocks(tmp_path):
    """When all bookmarks are present, all blocks are included in synced.blocks."""
    m = models.SemanticDocument(source_lang="ko", blocks=[
        models.ParagraphBlock(text="alive", block_id="id0", source_span_ids=["s0"]),
    ])
    out = os.path.join(str(tmp_path), "doc.docx")
    render_docx(m, out, str(tmp_path))

    synced = apply_docx_edits_to_model(m, out)

    assert len(synced.blocks) == 1
    assert synced.blocks[0].block_id == "id0"
