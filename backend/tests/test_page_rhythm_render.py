"""page_space_before applies only in parity mode; captions keep with next.

T3 of the template-fidelity work: the reconciler (T2) measures a per-page-boundary
compensating gap (`page_space_before`) for the first block on a source page; this
renderer wiring applies it ONLY when page-parity anchoring is on (see
docx_renderer._page_parity_on) — parity is the only mode that inserts the hard page
breaks this gap is meant to compensate for. `space_before` (the cover-rhythm field,
T-earlier) always takes priority when set. Captions get an unconditional
keep_with_next so they never separate from the figure/table they describe, regardless
of translation direction.
"""
import base64

import docx as docxlib

import models
from renderers.docx_renderer import render_docx

# Mirrors backend/tests/test_docx_extractor.py's _PNG_1x1 fixture.
_PNG_1PX = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAIAAACQd1PeAAAADElEQVR4nGP4//8/AAX+Av4N70a4AAAAAElFTkSuQmCC"
)


def test_page_space_before_applied_in_parity_mode(tmp_path):
    d = models.SemanticDocument(source_lang="en", blocks=[
        models.ParagraphBlock(text="a", page=0, source_span_ids=["s1"]),
        models.ParagraphBlock(text="b", page=1, page_space_before=120.0, source_span_ids=["s2"]),
    ])
    d.target_lang = "ko"
    out = str(tmp_path / "o.docx")
    render_docx(d, out, str(tmp_path))
    doc = docxlib.Document(out)
    p = next(p for p in doc.paragraphs if p.text == "b")
    assert p.paragraph_format.space_before.pt == 120.0


def test_page_space_before_ignored_when_parity_off(tmp_path):
    d = models.SemanticDocument(source_lang="ko", blocks=[
        models.ParagraphBlock(text="a", page=0, source_span_ids=["s1"]),
        models.ParagraphBlock(text="b", page=1, page_space_before=120.0, source_span_ids=["s2"]),
    ])
    d.target_lang = "en"
    out = str(tmp_path / "o.docx")
    render_docx(d, out, str(tmp_path))
    doc = docxlib.Document(out)
    p = next(p for p in doc.paragraphs if p.text == "b")
    assert p.paragraph_format.space_before is None or p.paragraph_format.space_before.pt != 120.0


def test_cover_space_before_still_wins(tmp_path):
    # If a block has both space_before (cover rhythm) and page_space_before, space_before wins.
    d = models.SemanticDocument(source_lang="en", blocks=[
        models.ParagraphBlock(text="t", page=0, space_before=50.0, page_space_before=99.0,
                              source_span_ids=["s1"]),
    ])
    d.target_lang = "ko"
    out = str(tmp_path / "o.docx")
    render_docx(d, out, str(tmp_path))
    doc = docxlib.Document(out)
    p = next(p for p in doc.paragraphs if p.text == "t")
    assert p.paragraph_format.space_before.pt == 50.0


def test_caption_keeps_with_next(tmp_path):
    d = models.SemanticDocument(source_lang="ko", blocks=[
        models.CaptionBlock(text="그림 1 제목", page=0, source_span_ids=["s1"]),
        models.FigureBlock(page=0, image_ref="img.png", source_span_ids=["s2"]),
    ])
    d.target_lang = "en"  # direction-independent — must apply regardless
    (tmp_path / "img.png").write_bytes(_PNG_1PX)
    out = str(tmp_path / "o.docx")
    render_docx(d, out, str(tmp_path))
    doc = docxlib.Document(out)
    cap = next(p for p in doc.paragraphs if "그림 1" in p.text)
    assert cap.paragraph_format.keep_with_next is True
