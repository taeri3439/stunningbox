from models import RunningBand, SemanticDocument


def test_running_band_image_and_page_number_default_none():
    b = RunningBand()
    assert b.image is None
    assert b.page_number is None


def test_running_band_round_trips_image_and_page_number():
    b = RunningBand(left="Market Report", image="hdr.png", page_number="right",
                    rule=True, rule_size=0.75, rule_color="123456")
    loaded = RunningBand.model_validate(b.model_dump())
    assert loaded.image == "hdr.png"
    assert loaded.page_number == "right"
    assert loaded.left == "Market Report"


def test_old_source_json_without_new_fields_loads():
    # A source.json written before Phase 2 has no image/page_number keys.
    doc = SemanticDocument.model_validate(
        {"running_header": {"left": "X", "rule": False}})
    assert doc.running_header.image is None
    assert doc.running_header.page_number is None


import os
from docx import Document
from docx.oxml.ns import qn
from PIL import Image
from renderers import _ooxml


def _png(path, w=600, h=80, color=(40, 80, 160)):
    Image.new("RGB", (w, h), color).save(path)
    return path


def test_add_band_image_anchors_full_width_behind_text_at_offset(tmp_path):
    img = _png(str(tmp_path / "band.png"))
    d = Document()
    sec = d.sections[0]
    # width 500pt, the band image is 500pt wide x 60pt tall, pinned 740pt down (a footer).
    _ooxml.add_band_image(sec.footer, img, width_pt=500.0, height_pt=60.0, y_offset_pt=740.0)
    xml = sec.footer._element.xml
    assert "anchor" in xml and 'behindDoc="1"' in xml
    extent = sec.footer._element.findall(".//" + qn("wp:extent"))[0]
    assert extent.get("cx") == str(int(500.0 * 12700))  # full-width EMU
    assert extent.get("cy") == str(int(60.0 * 12700))
    # The vertical posOffset places the strip at the bottom.
    offsets = [e.text for e in sec.footer._element.iter(qn("wp:posOffset"))]
    assert str(int(740.0 * 12700)) in offsets


def test_add_page_field_inserts_PAGE_fldchar(tmp_path):
    d = Document()
    p = d.add_paragraph()
    _ooxml.add_page_field(p)
    xml = p._p.xml
    assert "fldChar" in xml and 'w:fldCharType="begin"' in xml
    assert "PAGE" in xml
    assert 'w:fldCharType="end"' in xml


def test_add_page_background_still_full_page(tmp_path):
    # Regression: Phase-1 cover anchor unchanged after the refactor.
    img = _png(str(tmp_path / "cover.png"), w=595, h=842)
    d = Document()
    _ooxml.add_page_background(d.sections[0].first_page_header, img, 595.0, 842.0)
    xml = d.sections[0].first_page_header._element.xml
    assert 'behindDoc="1"' in xml
    extent = d.sections[0].first_page_header._element.findall(".//" + qn("wp:extent"))[0]
    assert extent.get("cx") == str(int(595.0 * 12700))
    assert extent.get("cy") == str(int(842.0 * 12700))


import fitz
from pipeline import extractor, reconciler
from pipeline.reconciler import reconcile


def _chrome_pdf(path, pages=3):
    """A doc whose body pages carry a colored top band + header title + a footer color bar
    + a page number — the decoration Phase 2 must capture."""
    doc = fitz.open()
    for i in range(pages):
        page = doc.new_page(width=300, height=400)
        # decorative header band (full-width colored strip across the top)
        page.draw_rect(fitz.Rect(0, 0, 300, 36), fill=(0.15, 0.3, 0.55))
        page.insert_text((20, 24), "Market Report", fontsize=11, color=(1, 1, 1))
        # body text
        page.insert_text((20, 200), f"body line on page {i}", fontsize=10)
        # decorative footer color bar + page number
        page.draw_rect(fitz.Rect(0, 372, 300, 400), fill=(0.15, 0.3, 0.55))
        page.insert_text((270, 392), str(i + 1), fontsize=10, color=(1, 1, 1))
    doc.save(path)
    doc.close()
    return path


def _plain_pdf(path, pages=3):
    """Body pages with header/footer TEXT but no decoration (white background)."""
    doc = fitz.open()
    for i in range(pages):
        page = doc.new_page(width=300, height=400)
        page.insert_text((20, 16), "Plain Header", fontsize=10)
        page.insert_text((20, 200), f"body {i}", fontsize=10)
        page.insert_text((270, 392), str(i + 1), fontsize=10)
    doc.save(path)
    doc.close()
    return path


def _blocks_for(raw, header_text="Market Report"):
    # Minimal analyzer output: mark the header title as chrome on each page so the band TEXT
    # is collected (page-number digits are chrome too); body line as a paragraph.
    page_blocks = []
    for pg in raw.pages:
        blocks = []
        for s in pg.spans:
            t = s.text.strip()
            if t == header_text or t.isdigit():
                blocks.append({"type": "chrome", "text": t, "source_span_ids": [s.id]})
            else:
                blocks.append({"type": "paragraph", "text": t, "source_span_ids": [s.id]})
        page_blocks.append(blocks)
    return page_blocks


def test_capture_header_and_footer_band_images(tmp_path):
    p = _chrome_pdf(str(tmp_path / "chrome.pdf"))
    ad = str(tmp_path / "assets")
    raw = extractor.extract_document(p, ad)
    out, _ = reconcile(raw, _blocks_for(raw), p, ad, source_lang="en")
    assert out.running_header and out.running_header.image, "header decoration captured"
    assert out.running_footer and out.running_footer.image, "footer decoration captured"
    assert os.path.exists(os.path.join(ad, out.running_header.image))
    assert os.path.exists(os.path.join(ad, out.running_footer.image))
    # The running title stays translatable text (not baked only into the image).
    assert out.running_header.left == "Market Report"


def test_footer_page_number_side_detected(tmp_path):
    p = _chrome_pdf(str(tmp_path / "chrome.pdf"))
    ad = str(tmp_path / "assets")
    raw = extractor.extract_document(p, ad)
    out, _ = reconcile(raw, _blocks_for(raw), p, ad, source_lang="en")
    assert out.running_footer.page_number == "right"  # digit sat at the right edge


def test_plain_header_emits_no_band_image(tmp_path):
    p = _plain_pdf(str(tmp_path / "plain.pdf"))
    ad = str(tmp_path / "assets")
    raw = extractor.extract_document(p, ad)
    out, _ = reconcile(raw, _blocks_for(raw, header_text="Plain Header"), p, ad, source_lang="en")
    # No decoration → no blank image anchored; the header text band still exists.
    assert out.running_header is None or out.running_header.image is None


from renderers.docx_renderer import render_docx
from models import RunningBand, SemanticDocument, ParagraphBlock


def test_render_band_image_anchored_behind_header_text(tmp_path):
    ad = str(tmp_path / "assets")
    os.makedirs(ad, exist_ok=True)
    Image.new("RGB", (600, 80), (40, 80, 160)).save(os.path.join(ad, "hdr_band.png"))
    doc = SemanticDocument(
        running_header=RunningBand(left="Market Report", image="hdr_band.png"),
        page_width=595.0, page_height=842.0,
        margin_left=72.0, margin_right=72.0, margin_top=72.0, margin_bottom=72.0,
        blocks=[ParagraphBlock(text="body", page=1)])
    out = str(tmp_path / "o.docx")
    render_docx(doc, out, ad)
    d = Document(out)
    xml = d.sections[0].header._element.xml
    assert 'behindDoc="1"' in xml  # band image anchored behind text
    assert "Market Report" in xml   # translatable title still rendered on top


def test_render_footer_page_field(tmp_path):
    ad = str(tmp_path / "assets")
    os.makedirs(ad, exist_ok=True)
    doc = SemanticDocument(
        running_footer=RunningBand(center="", page_number="right"),
        page_width=595.0, page_height=842.0,
        margin_left=72.0, margin_right=72.0, margin_top=72.0, margin_bottom=72.0,
        blocks=[ParagraphBlock(text="body", page=1)])
    out = str(tmp_path / "o.docx")
    render_docx(doc, out, ad)
    d = Document(out)
    xml = d.sections[0].footer._element.xml
    assert "fldChar" in xml and "PAGE" in xml  # auto page number field present


def test_render_chromeless_doc_unaffected(tmp_path):
    ad = str(tmp_path / "assets")
    os.makedirs(ad, exist_ok=True)
    doc = SemanticDocument(blocks=[ParagraphBlock(text="hello", page=0)])
    out = str(tmp_path / "o.docx")
    render_docx(doc, out, ad)
    d = Document(out)
    assert "hello" in d.paragraphs[0].text or any("hello" in p.text for p in d.paragraphs)


def test_end_to_end_chrome_pdf_to_docx(tmp_path):
    p = _chrome_pdf(str(tmp_path / "chrome.pdf"))
    ad = str(tmp_path / "assets")
    raw = extractor.extract_document(p, ad)
    doc, _ = reconcile(raw, _blocks_for(raw), p, ad, source_lang="en")
    out = str(tmp_path / "o.docx")
    render_docx(doc, out, ad)
    d = Document(out)
    # different-first-page so page-1 chrome differs from body pages
    assert d.sections[0].different_first_page_header_footer
    header_xml = d.sections[0].header._element.xml
    footer_xml = d.sections[0].footer._element.xml
    assert 'behindDoc="1"' in header_xml  # header band image
    assert 'behindDoc="1"' in footer_xml  # footer bar image
    assert "fldChar" in footer_xml and "PAGE" in footer_xml  # auto page number


def test_page_number_side_ignores_empty_text():
    # Regression: an empty/whitespace chrome item (e.g. a decorative rule the analyzer marked
    # chrome) must NOT be treated as a page number, or it injects a spurious PAGE field.
    # (_page_number_side was replaced by _page_number_spec, which returns (side, format).)
    from pipeline import reconciler
    assert reconciler._page_number_spec([("", "left", 1), ("  ", "left", 2)]) is None
    # A real recurring digit still wins.
    assert reconciler._page_number_spec(
        [("3", "right", 1), ("4", "right", 2), ("Title", "left", 1)]) == ("right", "{n}")


def test_thin_colored_footer_bar_captured_via_saturation(tmp_path):
    # A footer with ONLY a thin (low-area) but vivid color bar — below the non-white AREA
    # threshold, captured via the saturation signal so a brand bar isn't lost (dogfood case).
    doc = fitz.open()
    for i in range(2):
        page = doc.new_page(width=300, height=400)
        page.insert_text((20, 200), f"body {i}", fontsize=10)
        page.draw_rect(fitz.Rect(0, 396, 300, 398), fill=(1, 0, 0))  # 2pt red bar (thin)
    p = str(tmp_path / "thinbar.pdf")
    doc.save(p)
    doc.close()
    ad = str(tmp_path / "assets")
    raw = extractor.extract_document(p, ad)
    out, _ = reconcile(raw, _blocks_for(raw, header_text="__none__"), p, ad, source_lang="en")
    assert out.running_footer and out.running_footer.image, "thin color bar captured via saturation"


def test_plain_black_rule_footer_not_captured(tmp_path):
    # A footer with a thin BLACK rule + text only (no colour) → no image; the rule path handles it.
    doc = fitz.open()
    for i in range(2):
        page = doc.new_page(width=300, height=400)
        page.insert_text((20, 200), f"body {i}", fontsize=10)
        page.draw_line(fitz.Point(0, 396), fitz.Point(300, 396), color=(0, 0, 0), width=1)
        page.insert_text((150, 392), str(i + 1), fontsize=9)
    p = str(tmp_path / "plainrule.pdf")
    doc.save(p)
    doc.close()
    ad = str(tmp_path / "assets")
    raw = extractor.extract_document(p, ad)
    out, _ = reconcile(raw, _blocks_for(raw, header_text="__none__"), p, ad, source_lang="en")
    assert out.running_footer is None or out.running_footer.image is None
