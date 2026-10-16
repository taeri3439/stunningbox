"""Tests for apply_docx_edits_to_model: reorder + drop.

After T9, apply_docx_edits_to_model rebuilds translated.blocks from the docx's
bookmark order. Reordering paragraphs in the editor ⇒ translated.blocks reorders;
deleting a paragraph ⇒ its block is DROPPED from translated.blocks (not tombstoned).
"""
import os
import models
from pipeline.docx_sync import apply_docx_edits_to_model
from renderers.docx_renderer import render_docx
from docx import Document
from docx.oxml.ns import qn


def _reorder_paragraphs(path):
    """Swap the first two w:p that carry blk bookmarks (simulates editor drag/cut-paste)."""
    d = Document(path)
    body = d.element.body
    ps = [p for p in body.iter(qn("w:p")) if any(
        (bm.get(qn("w:name")) or "").startswith("blk")
        for bm in p.iter(qn("w:bookmarkStart")))]
    if len(ps) >= 2:
        a, b = ps[0], ps[1]
        a.addprevious(b)  # move b before a
    d.save(path)


def _delete_paragraph_by_bookmark(path, name):
    """Simulate an editor deletion: remove the w:p that contains bookmark `name`."""
    d = Document(path)
    for bm in list(d.element.body.iter(qn("w:bookmarkStart"))):
        if bm.get(qn("w:name")) == name:
            p = bm
            while p is not None and p.tag != qn("w:p"):
                p = p.getparent()
            if p is not None:
                p.getparent().remove(p)
            break
    d.save(path)


def test_apply_reorders_blocks_to_match_docx(tmp_path):
    """After swapping two bookmark paragraphs, synced.blocks order matches docx order."""
    m = models.SemanticDocument(source_lang="ko", blocks=[
        models.ParagraphBlock(text="zero", block_id="B0", source_span_ids=["s"]),
        models.ParagraphBlock(text="one",  block_id="B1", source_span_ids=["s"]),
        models.ParagraphBlock(text="two",  block_id="B2", source_span_ids=["s"]),
    ])
    out = os.path.join(str(tmp_path), "d.docx")
    render_docx(m, out, str(tmp_path))
    _reorder_paragraphs(out)  # editor moves B1 before B0

    synced = apply_docx_edits_to_model(m, out)

    assert [b.block_id for b in synced.blocks][:3] == ["B1", "B0", "B2"]


def test_apply_drops_deleted_block(tmp_path):
    """A paragraph whose bookmark was deleted is dropped from synced.blocks entirely."""
    m = models.SemanticDocument(source_lang="ko", blocks=[
        models.ParagraphBlock(text="first",  block_id="id0", source_span_ids=["s0"]),
        models.ParagraphBlock(text="second", block_id="id1", source_span_ids=["s1"]),
        models.ParagraphBlock(text="third",  block_id="id2", source_span_ids=["s2"]),
    ])
    out = os.path.join(str(tmp_path), "doc.docx")
    render_docx(m, out, str(tmp_path))
    _delete_paragraph_by_bookmark(out, "blkid1")  # user deletes the 2nd paragraph

    synced = apply_docx_edits_to_model(m, out)

    ids = [b.block_id for b in synced.blocks]
    assert "id1" not in ids             # dropped — not in rebuilt list
    assert "id0" in ids
    assert "id2" in ids
    assert len(synced.blocks) == 2      # exactly two blocks remain


def test_apply_non_bookmarked_blocks_ride_along(tmp_path):
    """Non-bookmarked blocks (figures/dividers) ride along with the preceding surviving block."""
    m = models.SemanticDocument(source_lang="ko", blocks=[
        models.ParagraphBlock(text="text0", block_id="P0", source_span_ids=["s0"]),
        models.FigureBlock(image_ref="img.png"),   # no block_id → rides with P0
        models.ParagraphBlock(text="text1", block_id="P1", source_span_ids=["s1"]),
    ])
    out = os.path.join(str(tmp_path), "d.docx")
    render_docx(m, out, str(tmp_path))

    synced = apply_docx_edits_to_model(m, out)

    # Figure should still be present
    types = [type(b).__name__ for b in synced.blocks]
    assert "FigureBlock" in types
    assert len(synced.blocks) == 3


def test_apply_deleted_text_block_not_ridden_along(tmp_path):
    """A DELETED text block (its id absent) is dropped, NOT ridden along as a figure."""
    m = models.SemanticDocument(source_lang="ko", blocks=[
        models.ParagraphBlock(text="keep",  block_id="P0", source_span_ids=["s0"]),
        models.ParagraphBlock(text="drop",  block_id="P1", source_span_ids=["s1"]),
        models.ParagraphBlock(text="keep2", block_id="P2", source_span_ids=["s2"]),
    ])
    out = os.path.join(str(tmp_path), "d.docx")
    render_docx(m, out, str(tmp_path))
    _delete_paragraph_by_bookmark(out, "blkP1")  # delete P1's bookmark

    synced = apply_docx_edits_to_model(m, out)

    ids = [b.block_id for b in synced.blocks]
    assert "P1" not in ids
    # Only text blocks with surviving ids remain (no phantom ride-along)
    assert ids == ["P0", "P2"]
