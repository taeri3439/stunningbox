# backend/tests/test_docx_sync_apply.py
import os
from models import SemanticDocument, HeadingBlock, ParagraphBlock, ListItemBlock, iter_segments
from renderers.docx_renderer import render_docx, _toc_split
from pipeline.docx_sync import apply_docx_edits_to_model
from docx import Document
from docx.oxml.ns import qn


def _render(tmp_path, doc, name="out.docx"):
    out = os.path.join(str(tmp_path), name)
    render_docx(doc, out, str(tmp_path))
    return out


def _model():
    return SemanticDocument(blocks=[
        HeadingBlock(level=1, marker="1.", text="제목", block_id="h0", source_span_ids=["s0"]),
        ParagraphBlock(text="원래 문장.", block_id="p1", source_span_ids=["s1"]),
        ParagraphBlock(lines=["줄 A", "줄 B"], block_id="p2", source_span_ids=["s2"]),
        ListItemBlock(marker="-", text="항목", block_id="l3", source_span_ids=["s3"]),
    ])


def test_unedited_roundtrip_is_identity(tmp_path):
    m = _model()
    out = _render(tmp_path, m)
    result = apply_docx_edits_to_model(m, out)
    assert result.model_dump() == m.model_dump()   # no spurious change


def test_text_edit_updates_only_that_block_preserving_structure(tmp_path):
    m = _model()
    out = _render(tmp_path, m)
    # Simulate an OnlyOffice text edit: change blkp1's run text in the saved docx.
    d = Document(out)
    for bm in d.element.body.iter(qn("w:bookmarkStart")):
        if bm.get(qn("w:name")) == "blkp1":
            p = bm
            while p.tag != qn("w:p"):
                p = p.getparent()
            for t in p.iter(qn("w:t")):
                t.text = "고친 문장."
                break
    d.save(out)

    result = apply_docx_edits_to_model(m, out)
    assert result.blocks[1].text == "고친 문장."
    # structure + other blocks + source linkage preserved:
    assert result.blocks[0].marker == "1." and result.blocks[0].text == "제목"
    assert result.blocks[2].lines == ["줄 A", "줄 B"]
    assert result.blocks[3].marker == "-" and result.blocks[3].text == "항목"
    assert result.blocks[1].source_span_ids == ["s1"]


def test_lines_block_segment_count_preserved_when_edited_line_count_differs(tmp_path):
    """Fix 1: when an editor changes a .lines block's line count, the block's segment
    count must stay the same (N lines → N segments) so positional pairing is not shifted.
    The following paragraph must still be independently correct."""
    m = SemanticDocument(blocks=[
        ParagraphBlock(lines=["줄 A", "줄 B"], block_id="p0", source_span_ids=["s0"]),  # 2 lines → 2 segments
        ParagraphBlock(text="다음 문단", block_id="p1", source_span_ids=["s1"]),         # 1 segment
    ])
    out = _render(tmp_path, m)

    # Simulate editor adding a third line-break run inside blkp0's paragraph.
    d = Document(out)
    for bm in d.element.body.iter(qn("w:bookmarkStart")):
        if bm.get(qn("w:name")) == "blkp0":
            p = bm
            while p.tag != qn("w:p"):
                p = p.getparent()
            # Replace every w:t inside this paragraph to simulate 3-line content.
            # We inject a new run with two line-breaks to turn it into 3 visual lines.
            from docx.oxml import OxmlElement
            # Clear existing runs and add fresh content with two <w:br/> breaks.
            for r in list(p.findall(qn("w:r"))):
                p.remove(r)
            for i, text in enumerate(["줄 A", "줄 B", "줄 C"]):
                if i > 0:
                    br_run = OxmlElement("w:r")
                    br = OxmlElement("w:br")
                    br.set(qn("w:type"), "textWrapping")
                    br_run.append(br)
                    p.append(br_run)
                run = OxmlElement("w:r")
                t = OxmlElement("w:t")
                t.text = text
                run.append(t)
                p.append(run)
            break
    d.save(out)

    result = apply_docx_edits_to_model(m, out)

    # Segment count must be unchanged (2 lines → 2 segments still).
    assert len(list(iter_segments(result))) == len(list(iter_segments(m)))
    # The .lines list must still have length 2 (original segment count preserved).
    assert result.blocks[0].lines is not None
    assert len(result.blocks[0].lines) == 2
    # The following paragraph must still be its own block, not merged in.
    assert result.blocks[1].text == "다음 문단"


def test_toc_block_not_overwritten_on_roundtrip(tmp_path):
    """Fix 2: a TOC-pattern block (text with dot-leader like '제목 ........... 5') must
    survive a render→apply round-trip with its text unchanged. The renderer drops the
    leader and emits 'title\\tpage'; reading that back must NOT overwrite the original."""
    toc_text = "제목 ........... 5"
    # Verify _toc_split recognises this pattern (sanity check for the regex).
    assert _toc_split(toc_text) is not None

    m = SemanticDocument(blocks=[
        ParagraphBlock(text=toc_text, block_id="t0", source_span_ids=["s0"]),
        ParagraphBlock(text="일반 문단", block_id="t1", source_span_ids=["s1"]),
    ])
    out = _render(tmp_path, m)

    result = apply_docx_edits_to_model(m, out)

    # TOC block's text must be unchanged — the dot-leader is preserved.
    assert result.blocks[0].text == toc_text
    # Non-TOC block is still updated normally on an unedited round-trip.
    assert result.blocks[1].text == "일반 문단"


def test_wrapped_toc_list_item_not_overwritten_on_roundtrip(tmp_path):
    """Fix 3: a wrapped TOC entry is a ListItemBlock whose `.lines` wrap across multiple
    preserved lines, with the LAST line ending in a dot-leader + page number. render_docx
    joins those lines and renders them as a single dot-leader row (_toc_split/_render_toc),
    dropping the literal line breaks. Reading that rendered row back must NOT glue the page
    number onto the title and overwrite the block's original wrapped `.lines` — even on an
    UNEDITED round-trip."""
    toc_lines = [
        "세부 항목 설명 텍스트가 길어서 줄바꿈된 목차 제목(MC11M-0047)",
        "................. 8",
    ]
    # Sanity check: the joined lines do form a TOC leader row per the renderer's own logic.
    assert _toc_split(" ".join(l for l in toc_lines if l.strip())) is not None

    m = SemanticDocument(blocks=[
        ListItemBlock(marker="2.2.1.2.2.", lines=list(toc_lines), block_id="t0", source_span_ids=["s0"]),
        ParagraphBlock(text="일반 문단", block_id="t1", source_span_ids=["s1"]),
    ])
    out = _render(tmp_path, m)

    result = apply_docx_edits_to_model(m, out)

    # Wrapped TOC block's lines (and text) must be unchanged — round-trip identity.
    assert result.blocks[0].lines == toc_lines
    assert result.blocks[0].text == m.blocks[0].text
    # Non-TOC block is still updated normally on an unedited round-trip.
    assert result.blocks[1].text == "일반 문단"
