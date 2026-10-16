import os
import models
from pipeline import process
from renderers.docx_renderer import render_docx
from docx import Document
from docx.oxml.ns import qn


def test_lazy_sync_reflects_editor_edit_when_docx_newer(tmp_path):
    """get_translated_model lazy-syncs: if output.{lang}.docx is newer than the cached
    translation (editor saved after), reading the model reflects + persists the edit. This
    fixes the race where the HTML views fetch before the save-callback sync completes."""
    doc_id, lang = "lz", "en"
    d = os.path.join(str(tmp_path), doc_id)
    os.makedirs(os.path.join(d, "assets"), exist_ok=True)
    m = models.SemanticDocument(source_lang="ko", blocks=[
        models.ParagraphBlock(text="orig text", block_id="p0", source_span_ids=["s0"]),
    ])
    process.save_translated_model(str(tmp_path), doc_id, lang, m)
    out = process.output_docx_path(str(tmp_path), doc_id, lang)
    render_docx(m, out, os.path.join(d, "assets"))

    # simulate an OnlyOffice text edit in the saved docx
    dd = Document(out)
    for t in dd.element.body.iter(qn("w:t")):
        t.text = "editor changed"
        break
    dd.save(out)
    # guarantee output.docx is newer than translated.json (avoid fs mtime-resolution flakiness)
    newer = os.path.getmtime(process.translated_path(str(tmp_path), doc_id, lang)) + 10
    os.utime(out, (newer, newer))

    got = process.get_translated_model(str(tmp_path), doc_id, lang)
    assert got.blocks[0].text == "editor changed"                      # lazy-synced on read
    with open(process.translated_path(str(tmp_path), doc_id, lang), encoding="utf-8") as fh:
        assert "editor changed" in fh.read()                            # and persisted


def test_lazy_sync_no_op_when_cache_newer(tmp_path):
    """When translated.json is newer than the docx (no editor save since), the read is a
    plain cache read — no sync, model returned as-is."""
    doc_id, lang = "lz2", "en"
    d = os.path.join(str(tmp_path), doc_id)
    os.makedirs(os.path.join(d, "assets"), exist_ok=True)
    m = models.SemanticDocument(source_lang="ko", blocks=[
        models.ParagraphBlock(text="orig text", block_id="p0", source_span_ids=["s0"]),
    ])
    out = process.output_docx_path(str(tmp_path), doc_id, lang)
    render_docx(m, out, os.path.join(d, "assets"))
    # edit the docx but make the CACHE the newer file → no lazy-sync should occur
    dd = Document(out)
    for t in dd.element.body.iter(qn("w:t")):
        t.text = "should be ignored"
        break
    dd.save(out)
    process.save_translated_model(str(tmp_path), doc_id, lang, m)  # written last → newest
    newer = os.path.getmtime(out) + 10
    os.utime(process.translated_path(str(tmp_path), doc_id, lang), (newer, newer))

    got = process.get_translated_model(str(tmp_path), doc_id, lang)
    assert got.blocks[0].text == "orig text"  # cache wins; docx edit not pulled in
