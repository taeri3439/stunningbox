"""Latin-only fonts get a class-matched eastAsia substitute; TOC entries indent by depth."""
import docx as docxlib
from docx.oxml.ns import qn

import models
from renderers.docx_renderer import render_docx


def _rfonts(path, needle):
    doc = docxlib.Document(path)
    for p in doc.paragraphs:
        if needle in p.text:
            for r in p.runs:
                rPr = r._element.find(qn("w:rPr"))
                if rPr is None:
                    continue
                rf = rPr.find(qn("w:rFonts"))
                if rf is not None:
                    return {k.split('}')[-1]: v for k, v in rf.attrib.items()}
    return None


def _render_one(tmp_path, block):
    d = models.SemanticDocument(source_lang="en", blocks=[block])
    d.target_lang = "ko"
    out = str(tmp_path / "o.docx")
    render_docx(d, out, str(tmp_path))
    return out


def test_latin_sans_font_swaps_eastasia_to_malgun(tmp_path):
    out = _render_one(tmp_path, models.ParagraphBlock(
        text="시장 개요", page=0, font="FranklinGothic-Book", source_span_ids=["s1"]))
    rf = _rfonts(out, "시장 개요")
    assert rf["eastAsia"] == "Malgun Gothic"
    assert rf["ascii"] == "FranklinGothic-Book"  # 라틴은 소스 이름 유지


def test_latin_serif_font_swaps_eastasia_to_batang(tmp_path):
    out = _render_one(tmp_path, models.ParagraphBlock(
        text="본문", page=0, font="TimesNewRomanPSMT", source_span_ids=["s1"]))
    rf = _rfonts(out, "본문")
    assert rf["eastAsia"] == "Batang"


def test_korean_font_name_unchanged(tmp_path):
    out = _render_one(tmp_path, models.ParagraphBlock(
        text="본문", page=0, font="맑은 고딕", source_span_ids=["s1"]))
    rf = _rfonts(out, "본문")
    assert rf["eastAsia"] == "맑은 고딕" and rf["ascii"] == "맑은 고딕"


def test_toc_entries_indent_by_marker_depth(tmp_path):
    d = models.SemanticDocument(source_lang="en", blocks=[
        models.ListItemBlock(marker="12.", text="ZOETIS ......... 149", page=0, source_span_ids=["a"]),
        models.ListItemBlock(marker="12.2", text="개요 ......... 149", page=0, source_span_ids=["b"]),
        models.ListItemBlock(marker="12.2.1", text="사업 ......... 150", page=0, source_span_ids=["c"]),
    ])
    d.target_lang = "ko"
    out = str(tmp_path / "o.docx")
    render_docx(d, out, str(tmp_path))
    doc = docxlib.Document(out)
    tocs = [p for p in doc.paragraphs if "\t" in p.text or "149" in p.text or "150" in p.text]
    ind = [(p.paragraph_format.left_indent.inches if p.paragraph_format.left_indent else 0.0)
           for p in tocs[:3]]
    assert ind[0] < ind[1] < ind[2]
    assert abs(ind[1] - 0.18) < 0.02 and abs(ind[2] - 0.36) < 0.02


# ---- FINDING 1: TOC indent priority (source-measured indent WINS over depth fallback) ----

def test_toc_entry_with_measured_indent_source_wins(tmp_path):
    # A reconciler/HWP-extracted TOC list item routinely carries a source-measured
    # `indent` (pt). That real fidelity data must win over the depth-based fallback,
    # even though the marker's depth ("12.2.1" -> depth 2) would otherwise imply 0.36in.
    d = models.SemanticDocument(source_lang="en", blocks=[
        models.ListItemBlock(marker="12.2.1", text="사업 ......... 150", page=0,
                             indent=50.0, source_span_ids=["a"]),
    ])
    d.target_lang = "ko"
    out = str(tmp_path / "o.docx")
    render_docx(d, out, str(tmp_path))
    doc = docxlib.Document(out)
    tocs = [p for p in doc.paragraphs if "\t" in p.text or "150" in p.text]
    assert tocs and tocs[0].paragraph_format.left_indent == docxlib.shared.Pt(50)


def test_toc_entry_without_measured_indent_depth_fallback_applies(tmp_path):
    # No source-measured indent -> the depth-based fallback (min(depth, 4) * 0.18in) applies.
    d = models.SemanticDocument(source_lang="en", blocks=[
        models.ListItemBlock(marker="12.2.1", text="사업 ......... 150", page=0,
                             source_span_ids=["a"]),
    ])
    d.target_lang = "ko"
    out = str(tmp_path / "o.docx")
    render_docx(d, out, str(tmp_path))
    doc = docxlib.Document(out)
    tocs = [p for p in doc.paragraphs if "\t" in p.text or "150" in p.text]
    assert tocs and abs(tocs[0].paragraph_format.left_indent.inches - 0.36) < 0.02


# ---- FINDING 2: ja/zh eastAsia substitution gate + language-aware substitute table ----

def test_ja_target_kanji_text_franklingothic_swaps_to_ms_gothic(tmp_path):
    d = models.SemanticDocument(source_lang="en", blocks=[
        models.ParagraphBlock(text="市場概要", page=0, font="FranklinGothic-Book",
                              source_span_ids=["s1"]),
    ])
    d.target_lang = "ja"
    out = str(tmp_path / "o.docx")
    render_docx(d, out, str(tmp_path))
    rf = _rfonts(out, "市場概要")
    assert rf["eastAsia"] == "MS Gothic"
    assert rf["ascii"] == "FranklinGothic-Book"  # latin ascii/hAnsi stay the source name


def test_zh_target_hanzi_text_serif_latin_font_swaps_to_simsun(tmp_path):
    d = models.SemanticDocument(source_lang="en", blocks=[
        models.ParagraphBlock(text="市场概述", page=0, font="TimesNewRomanPSMT",
                              source_span_ids=["s1"]),
    ])
    d.target_lang = "zh"
    out = str(tmp_path / "o.docx")
    render_docx(d, out, str(tmp_path))
    rf = _rfonts(out, "市场概述")
    assert rf["eastAsia"] == "SimSun"


def test_unknown_target_with_cjk_text_falls_back_to_ko_pair(tmp_path):
    # Defensive fallback: an other/unknown target_lang with CJK text in the run still
    # gets a substitute (never left as the latin source name) — falls back to the ko pair.
    d = models.SemanticDocument(source_lang="en", blocks=[
        models.ParagraphBlock(text="시장 개요", page=0, font="FranklinGothic-Book",
                              source_span_ids=["s1"]),
    ])
    d.target_lang = "fr"
    out = str(tmp_path / "o.docx")
    render_docx(d, out, str(tmp_path))
    rf = _rfonts(out, "시장 개요")
    assert rf["eastAsia"] == "Malgun Gothic"


# ---- FINDING 3: ko->en direction, Korean-named + latin fonts, byte-identical rFonts ----

def test_ko_to_en_direction_korean_font_and_latin_font_unchanged(tmp_path):
    d = models.SemanticDocument(source_lang="ko", blocks=[
        models.ParagraphBlock(text="hello", page=0, font="맑은 고딕", source_span_ids=["s1"]),
        models.ParagraphBlock(text="world", page=1, font="Arial", source_span_ids=["s2"]),
    ])
    d.target_lang = "en"
    out = str(tmp_path / "o.docx")
    render_docx(d, out, str(tmp_path))
    rf1 = _rfonts(out, "hello")
    rf2 = _rfonts(out, "world")
    assert rf1 == {"ascii": "맑은 고딕", "hAnsi": "맑은 고딕", "eastAsia": "맑은 고딕", "cs": "맑은 고딕"}
    assert rf2 == {"ascii": "Arial", "hAnsi": "Arial", "eastAsia": "Arial", "cs": "Arial"}


# ---- FINDING 4: ja/zh target_lang must flow into table cells and running bands too, not
# just body paragraphs — otherwise those pick up the ko substitute pair while body text
# gets the correct ja/zh pair (inconsistent CJK families in one document). ----

def _table_cell_rfonts(path, row=0, col=0):
    doc = docxlib.Document(path)
    cell = doc.tables[0].rows[row].cells[col]
    for run in cell.paragraphs[0].runs:
        rPr = run._element.find(qn("w:rPr"))
        if rPr is None:
            continue
        rf = rPr.find(qn("w:rFonts"))
        if rf is not None:
            return {k.split('}')[-1]: v for k, v in rf.attrib.items()}
    return None


def test_ja_target_table_cell_kanji_latin_font_swaps_to_ms_gothic(tmp_path):
    d = models.SemanticDocument(source_lang="en", blocks=[
        models.TableBlock(page=0, source_span_ids=["t1"], rows=[
            [models.TableCell(text="市場概要", font="FranklinGothic-Book")],
        ]),
    ])
    d.target_lang = "ja"
    out = str(tmp_path / "o.docx")
    render_docx(d, out, str(tmp_path))
    rf = _table_cell_rfonts(out)
    assert rf is not None
    assert rf["eastAsia"] == "MS Gothic"
    assert rf["ascii"] == "FranklinGothic-Book"
