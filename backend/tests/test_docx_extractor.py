from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from models import HeadingBlock, ParagraphBlock
from pipeline import docx_extractor


def _save(d, tmp_path, name="in.docx"):
    p = tmp_path / name
    d.save(str(p))
    return str(p)


def test_heading_and_paragraph_mapping(tmp_path):
    d = Document()
    d.add_heading("Title One", level=1)
    d.add_heading("Sub", level=2)
    p = d.add_paragraph("Body text here.")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    path = _save(d, tmp_path)

    doc = docx_extractor.extract_docx(path, str(tmp_path / "assets"))
    kinds = [(type(b).__name__, getattr(b, "text", None)) for b in doc.blocks]
    assert ("HeadingBlock", "Title One") in kinds
    h1 = next(b for b in doc.blocks if isinstance(b, HeadingBlock) and b.text == "Title One")
    assert h1.level == 1
    h2 = next(b for b in doc.blocks if isinstance(b, HeadingBlock) and b.text == "Sub")
    assert h2.level == 2
    body = next(b for b in doc.blocks if isinstance(b, ParagraphBlock))
    assert body.text == "Body text here."
    assert body.align == "center"


from models import ListItemBlock


def test_list_items_get_level(tmp_path):
    d = Document()
    d.add_paragraph("First bullet", style="List Bullet")
    d.add_paragraph("Second bullet", style="List Bullet")
    path = _save(d, tmp_path, "lists.docx")
    doc = docx_extractor.extract_docx(path, str(tmp_path / "assets2"))
    items = [b for b in doc.blocks if isinstance(b, ListItemBlock)]
    assert len(items) == 2
    assert items[0].text == "First bullet"
    assert items[0].level == 0


from models import TableBlock


def test_table_basic_and_colspan(tmp_path):
    d = Document()
    t = d.add_table(rows=2, cols=2)
    t.cell(1, 0).text = "C"
    t.cell(1, 1).text = "D"
    # Merge first, THEN set text — python-docx concatenates the source cells' text on
    # merge, so setting it afterwards gives the merged cell a single clean value.
    merged = t.cell(0, 0).merge(t.cell(0, 1))
    merged.text = "A"
    path = _save(d, tmp_path, "table.docx")
    doc = docx_extractor.extract_docx(path, str(tmp_path / "assets3"))
    tables = [b for b in doc.blocks if isinstance(b, TableBlock)]
    assert len(tables) == 1
    rows = tables[0].rows
    assert rows[0][0].colspan == 2
    assert rows[0][0].text == "A"
    assert [c.text for c in rows[1]] == ["C", "D"]


import base64, os as _os
from models import FigureBlock

_PNG_1x1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAIAAACQd1PeAAAADElEQVR4nGP4//8/AAX+Av4N70a4AAAAAElFTkSuQmCC"
)


def test_inline_image_becomes_figure(tmp_path):
    img = tmp_path / "pic.png"
    img.write_bytes(_PNG_1x1)
    d = Document()
    d.add_paragraph("Before image.")
    d.add_picture(str(img))
    path = _save(d, tmp_path, "img.docx")
    assets = tmp_path / "assets_img"
    doc = docx_extractor.extract_docx(path, str(assets))
    figs = [b for b in doc.blocks if isinstance(b, FigureBlock)]
    assert len(figs) == 1
    assert figs[0].image_ref
    assert _os.path.exists(_os.path.join(str(assets), figs[0].image_ref))


def test_header_footer_become_bands(tmp_path):
    d = Document()
    d.sections[0].header.paragraphs[0].text = "My Running Header"
    d.add_paragraph("Body.")
    path = _save(d, tmp_path, "hf.docx")
    doc = docx_extractor.extract_docx(path, str(tmp_path / "assets_hf"))
    assert doc.running_header is not None
    assert "My Running Header" in (doc.running_header.center or "")
