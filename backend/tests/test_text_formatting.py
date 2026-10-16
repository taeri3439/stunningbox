"""Phase 1A — block-level text formatting (font family, colour, italic/underline,
justified alignment) and the over-bold fix, across the PDF reconciler, the DOCX
extractor, and the renderer."""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

from models import (
    HeadingBlock, ListItemBlock, ParagraphBlock, RawDocument, RawPage, RawSpan,
    RunningBand, SemanticDocument,
)
from pipeline import reconciler as R
from pipeline.reconciler import reconcile
from renderers.docx_renderer import render_docx


# ---- PDF path: reconciler capture --------------------------------------------------

def test_norm_font_maps_to_universal_class():
    # subset prefix dropped, then classified to a universally-available family
    assert R._norm_font("ABCDEF+Batang") == "Times New Roman"   # serif
    assert R._norm_font("MalgunGothicBold") == "Arial"          # gothic -> sans
    assert R._norm_font("ABCDEF+MalgunGothic") == "Arial"
    assert R._norm_font("Arial-BoldItalic") == "Arial"
    assert R._norm_font("Times New Roman,Bold") == "Times New Roman"
    assert R._norm_font("") is None


def test_toc_marker_not_duplicated_as_line(tmp_path):
    # A TOC entry whose marker ("4.") sits on its own line must NOT keep the marker as a
    # preserved line — else it renders twice ("4. 4.") on two lines instead of one.
    spans = [
        RawSpan(id="p0b0l0s0", page=0, text="4.", bbox=(72, 100, 90, 112), size=11),
        RawSpan(id="p0b0l1s0", page=0, text="Nonclinical Overview .......... 2",
                bbox=(108, 100, 500, 112), size=11),
    ]
    raw = RawDocument(pages=[RawPage(width=595, height=842, spans=spans)])
    pb = [[{"type": "list_item", "marker": "4.", "level": 0,
            "text": "Nonclinical Overview .......... 2",
            "source_span_ids": ["p0b0l0s0", "p0b0l1s0"]}]]
    doc, _ = reconcile(raw, pb, str(tmp_path / "x.pdf"), str(tmp_path), source_lang="en")
    li = [b for b in doc.blocks if b.type == "list_item"][0]
    assert li.marker == "4."
    assert li.lines is None              # marker line dropped -> single line -> reflow/TOC
    assert li.text.startswith("Nonclinical Overview")


def test_universal_font_classifier():
    from models import universal_font
    assert universal_font("맑은 고딕") == "Arial"        # Hangul sans
    assert universal_font("바탕체") == "Times New Roman"  # Hangul serif
    assert universal_font("Dotum") == "Arial"
    assert universal_font("SomeUnknownFace") is None     # unknown -> renderer default
    assert universal_font(None) is None


def _spans(font="Batang", color=0, italic_font=False):
    name = "Batang-Italic" if italic_font else font
    return {
        "a": RawSpan(id="a", page=0, text="x", bbox=(0, 0, 10, 10), font=name, color=color),
        "b": RawSpan(id="b", page=0, text="y", bbox=(0, 0, 10, 10), font=name, color=color),
    }


def test_dominant_inline_font_color_italic():
    font, color, italic = R._dominant_inline(["a", "b"], _spans(font="Batang", color=0xCC0000))
    assert font == "Times New Roman"  # Batang (serif) -> universal serif
    assert color == "CC0000"          # int -> hex
    assert italic is None
    # near-black colour is treated as ink (None)
    _, color2, _ = R._dominant_inline(["a", "b"], _spans(color=0x101010))
    assert color2 is None
    # italic inferred from font name
    _, _, italic2 = R._dominant_inline(["a", "b"], _spans(italic_font=True))
    assert italic2 is True


def test_justified_detection():
    # Three lines, the first two reach the same right edge (520) -> justified.
    by_id = {
        "p0b0l0s0": RawSpan(id="p0b0l0s0", page=0, text="a", bbox=(72, 100, 520, 112)),
        "p0b0l1s0": RawSpan(id="p0b0l1s0", page=0, text="b", bbox=(72, 120, 519, 132)),
        "p0b0l2s0": RawSpan(id="p0b0l2s0", page=0, text="c", bbox=(72, 140, 300, 152)),  # last, ragged
    }
    assert R._justified(list(by_id), by_id) is True
    # Ragged right (short inner line) -> not justified
    by_id["p0b0l1s0"] = RawSpan(id="p0b0l1s0", page=0, text="b", bbox=(72, 120, 300, 132))
    assert R._justified(list(by_id), by_id) is False


# ---- Renderer ----------------------------------------------------------------------

def _render(block, tmp_path):
    out = str(tmp_path / "t.docx")
    render_docx(SemanticDocument(blocks=[block]), out, str(tmp_path))
    return Document(out)


def test_renderer_applies_font_color_italic(tmp_path):
    d = _render(ParagraphBlock(text="hello", font="Arial", color="CC0000", italic=True), tmp_path)
    run = d.paragraphs[-1].runs[0]
    assert run.font.name == "Arial"
    rFonts = run._element.rPr.find(qn("w:rFonts"))
    assert rFonts.get(qn("w:eastAsia")) == "Arial"   # CJK font set too
    assert str(run.font.color.rgb) == "CC0000"
    assert run.italic is True


def test_renderer_no_color_forces_ink(tmp_path):
    d = _render(ParagraphBlock(text="plain"), tmp_path)
    # regression guard: without a captured colour, runs render near-black (not Word blue)
    assert str(d.paragraphs[-1].runs[0].font.color.rgb) == "1B1A17"


def test_renderer_justify_alignment(tmp_path):
    d = _render(ParagraphBlock(text="body", align="justify"), tmp_path)
    assert d.paragraphs[-1].alignment == WD_ALIGN_PARAGRAPH.JUSTIFY


def test_heading_bold_false_unbolds(tmp_path):
    # A heading the source captured as NOT bold must not render bold despite Word's style.
    d = _render(HeadingBlock(level=1, text="Title", bold=False), tmp_path)
    assert all(r.bold is False for r in d.paragraphs[-1].runs)


def test_list_marker_not_force_bolded(tmp_path):
    # Over-bold fix: a non-bold list item's marker run is not forced bold.
    d = _render(ListItemBlock(marker="4.1", text="entry", bold=False, level=0), tmp_path)
    assert all(not r.bold for r in d.paragraphs[-1].runs)


def test_running_header_left_right_uses_borderless_table(tmp_path):
    # A left|right header is rendered as a borderless 2-cell table so each element pins to
    # its edge in every renderer (tab stops render inconsistently across Word/OO/LO).
    doc = SemanticDocument(
        blocks=[ParagraphBlock(text="body")],
        running_header=RunningBand(left="Korean Fire Equipment Standards",
                                   right="Aerial Ladder Vehicle", rule=True),
    )
    out = str(tmp_path / "hdr.docx")
    render_docx(doc, out, str(tmp_path))
    hdr = Document(out).sections[0].header
    assert len(hdr.tables) == 1
    t = hdr.tables[0]
    assert len(t.rows[0].cells) == 2
    assert t.cell(0, 0).text == "Korean Fire Equipment Standards"
    assert t.cell(0, 1).text == "Aerial Ladder Vehicle"
    assert t.cell(0, 1).paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT  # pinned right


# ---- DOCX path: extractor capture --------------------------------------------------

def test_docx_extractor_captures_style(tmp_path):
    from pipeline.docx_extractor import extract_docx
    src = Document()
    p = src.add_paragraph()
    r = p.add_run("colored italic")
    r.italic = True
    r.font.name = "Arial"
    r.font.color.rgb = __import__("docx").shared.RGBColor(0xCC, 0x00, 0x00)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    sp = str(tmp_path / "in.docx")
    src.save(sp)

    doc = extract_docx(sp, str(tmp_path / "assets"))
    blk = doc.blocks[0]
    assert blk.italic is True
    assert blk.color == "CC0000"
    assert blk.align == "justify"
    assert blk.font == "Arial"
