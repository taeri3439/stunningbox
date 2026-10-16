"""Page-parity anchoring: source-page boundaries become page_break_before in docx.

Direction-gated: ON when the translation direction shrinks text (latin -> CJK target,
e.g. en->ko) or the model explicitly opts in via `page_parity`; OFF for expanding
directions (e.g. ko->en) — see the 2026-06-15 bloat removal (e7b3fbc) this task revives
in a non-bloating form (paragraph_format.page_break_before on the block's own first
paragraph, not an extra empty paragraph)."""
import docx as docxlib

import models
from renderers.docx_renderer import render_docx


def _doc(blocks, target_lang="ko", page_parity=None):
    d = models.SemanticDocument(source_lang="en", blocks=blocks)
    d.target_lang = target_lang
    d.page_parity = page_parity
    return d


def _para_texts_with_break(path):
    out = []
    doc = docxlib.Document(path)
    for p in doc.paragraphs:
        pf = p.paragraph_format
        out.append((p.text, bool(pf.page_break_before)))
    return out


def test_break_before_first_block_of_each_source_page(tmp_path):
    doc = _doc([
        models.ParagraphBlock(text="p0 a", page=0, source_span_ids=["s1"]),
        models.ParagraphBlock(text="p0 b", page=0, source_span_ids=["s2"]),
        models.ParagraphBlock(text="p1 first", page=1, source_span_ids=["s3"]),
        models.ParagraphBlock(text="p2 first", page=2, source_span_ids=["s4"]),
    ])
    out = str(tmp_path / "o.docx")
    render_docx(doc, out, str(tmp_path))
    marks = dict(_para_texts_with_break(out))
    assert marks["p0 a"] is False and marks["p0 b"] is False
    assert marks["p1 first"] is True and marks["p2 first"] is True


def test_no_empty_break_paragraphs(tmp_path):
    doc = _doc([
        models.ParagraphBlock(text="p0", page=0, source_span_ids=["s1"]),
        models.ParagraphBlock(text="p1", page=1, source_span_ids=["s2"]),
    ])
    out = str(tmp_path / "o.docx")
    render_docx(doc, out, str(tmp_path))
    # 과거 bloat의 원인이던 브레이크 전용 빈 문단이 없어야 한다: every paragraph carrying a
    # break also carries real text.
    assert all(t.strip() for t, br in _para_texts_with_break(out) if br)


def test_zero_block_source_page_gets_no_extra_break(tmp_path):
    # page 1에 블록이 없어도(전량 소실) page 2 첫 블록은 브레이크 1개만 갖는다
    doc = _doc([
        models.ParagraphBlock(text="p0", page=0, source_span_ids=["s1"]),
        models.ParagraphBlock(text="p2", page=2, source_span_ids=["s2"]),
    ])
    out = str(tmp_path / "o.docx")
    render_docx(doc, out, str(tmp_path))
    marks = dict(_para_texts_with_break(out))
    assert marks["p2"] is True  # 하나의 break_before일 뿐 blank page 2장 아님


def test_table_first_page_uses_1pt_spacer(tmp_path):
    doc = _doc([
        models.ParagraphBlock(text="p0", page=0, source_span_ids=["s1"]),
        models.TableBlock(page=1, rows=[[models.TableCell(text="c")]], source_span_ids=["s2"]),
    ])
    out = str(tmp_path / "o.docx")
    render_docx(doc, out, str(tmp_path))
    d = docxlib.Document(out)
    spacers = [p for p in d.paragraphs
               if p.paragraph_format.page_break_before and not p.text.strip()]
    assert len(spacers) == 1
    sp = spacers[0]
    assert sp.runs == [] or all(r.font.size.pt <= 2 for r in sp.runs if r.font.size)
    # 스페이서는 여백을 만들지 않아야 한다
    assert (sp.paragraph_format.space_after or 0) == 0


def test_parity_off_for_expanding_direction(tmp_path):
    doc = _doc([
        models.ParagraphBlock(text="p0", page=0, source_span_ids=["s1"]),
        models.ParagraphBlock(text="p1", page=1, source_span_ids=["s2"]),
    ], target_lang="en")
    doc.source_lang = "ko"
    out = str(tmp_path / "o.docx")
    render_docx(doc, out, str(tmp_path))
    assert not any(br for _, br in _para_texts_with_break(out))


def test_explicit_page_parity_overrides_direction(tmp_path):
    doc = _doc([
        models.ParagraphBlock(text="p0", page=0, source_span_ids=["s1"]),
        models.ParagraphBlock(text="p1", page=1, source_span_ids=["s2"]),
    ], target_lang="en", page_parity=True)
    doc.source_lang = "ko"
    out = str(tmp_path / "o.docx")
    render_docx(doc, out, str(tmp_path))
    marks = dict(_para_texts_with_break(out))
    assert marks["p1"] is True


def test_parity_skips_first_break_when_cover_already_isolated_it(tmp_path):
    """A cover background already forces the page-0 -> page-1 break (existing mechanism);
    the parity anchor must not double it into two page breaks."""
    from PIL import Image
    import os

    Image.new("RGB", (300, 400), (20, 80, 130)).save(os.path.join(str(tmp_path), "cover_bg.png"))
    doc = _doc([
        models.HeadingBlock(text="cover title", page=0),
        models.ParagraphBlock(text="body intro", page=1, source_span_ids=["s1"]),
        models.ParagraphBlock(text="page two", page=2, source_span_ids=["s2"]),
    ])
    doc.cover_background = "cover_bg.png"
    doc.page_width = 300.0
    doc.page_height = 400.0
    out = str(tmp_path / "o.docx")
    render_docx(doc, out, str(tmp_path))
    body_xml = docxlib.Document(out).element.body.xml
    # Exactly one hard w:br page break (the cover-isolation break) plus the anchor's
    # page_break_before on "page two" — not two hard breaks stacked at the same boundary.
    assert body_xml.count('w:type="page"') == 1
    marks = dict(_para_texts_with_break(out))
    assert marks["body intro"] is False  # cover break already isolated this page
    assert marks["page two"] is True


def test_parity_zeroes_default_paragraph_spacing(tmp_path):
    """Parity mode kills the template's inherited 10pt-after/1.15x docDefaults (dense
    designed layouts overflow their anchored pages under that inflation); non-parity
    documents keep the template defaults untouched."""
    import docx as docxlib2

    on = _doc([models.ParagraphBlock(text="p0", page=0, source_span_ids=["s1"])])
    out_on = str(tmp_path / "on.docx")
    render_docx(on, out_on, str(tmp_path))
    normal_on = docxlib2.Document(out_on).styles["Normal"].paragraph_format
    assert normal_on.space_after is not None and normal_on.space_after.pt == 0
    assert normal_on.line_spacing == 1.0

    off = _doc([models.ParagraphBlock(text="p0", page=0, source_span_ids=["s1"])],
               target_lang="en")
    off.source_lang = "ko"
    out_off = str(tmp_path / "off.docx")
    render_docx(off, out_off, str(tmp_path))
    normal_off = docxlib2.Document(out_off).styles["Normal"].paragraph_format
    assert normal_off.space_after is None
    assert normal_off.line_spacing is None


def test_parity_biases_heading_and_toc_density(tmp_path):
    """Parity mode: headings get a modest fixed gap (not Word's airy defaults) and
    TOC entries use exact single spacing — underfill is safe under anchors, spill
    is not. Non-parity documents keep today's spacing."""
    import docx as docxlib3

    blocks = [
        models.HeadingBlock(text="개요", level=1, page=0, source_span_ids=["s1"]),
        models.ListItemBlock(marker="1.", text="적용범위 ......... 3", page=0,
                             source_span_ids=["s2"]),
    ]
    on = _doc(list(blocks))
    out_on = str(tmp_path / "on.docx")
    render_docx(on, out_on, str(tmp_path))
    d_on = docxlib3.Document(out_on)
    head = next(p for p in d_on.paragraphs if "개요" in p.text)
    assert head.paragraph_format.space_before.pt == 4
    assert head.paragraph_format.space_after.pt == 2
    toc_p = next(p for p in d_on.paragraphs if "적용범위" in p.text)
    assert toc_p.paragraph_format.line_spacing == 1.0

    off = _doc(list(blocks), target_lang="en")
    off.source_lang = "ko"
    out_off = str(tmp_path / "off.docx")
    render_docx(off, out_off, str(tmp_path))
    d_off = docxlib3.Document(out_off)
    head_off = next(p for p in d_off.paragraphs if "개요" in p.text)
    assert head_off.paragraph_format.space_before is None
