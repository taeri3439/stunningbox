"""Phase 1B — table cell shading / per-cell alignment / vertical alignment + cell font
and colour, across the renderer, the analyzer-output validators, and the DOCX path."""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor

from models import RawSpan, SemanticDocument, TableBlock, TableCell
from pipeline import reconciler as R
from pipeline.docx_extractor import _tc_align, _tc_shading, _tc_valign
from renderers.docx_renderer import render_docx


def _render_table(cell: TableCell, tmp_path):
    doc = SemanticDocument(blocks=[TableBlock(rows=[[cell]])])
    out = str(tmp_path / "t.docx")
    render_docx(doc, out, str(tmp_path))
    return Document(out)


def test_render_cell_shading_align_valign_font(tmp_path):
    d = _render_table(
        TableCell(text="hdr", shading="EFEFEF", align="center", valign="center",
                  font="Arial", color="CC0000"),
        tmp_path,
    )
    tc = d.tables[0].cell(0, 0)._tc
    tcPr = tc.tcPr
    assert tcPr.find(qn("w:shd")).get(qn("w:fill")) == "EFEFEF"
    assert tcPr.find(qn("w:vAlign")).get(qn("w:val")) == "center"
    cellp = d.tables[0].cell(0, 0).paragraphs[0]
    assert cellp.alignment == WD_ALIGN_PARAGRAPH.CENTER
    run = cellp.runs[0]
    assert run.font.name == "Arial"
    assert str(run.font.color.rgb) == "CC0000"


def test_render_cell_no_shading_emits_no_shd(tmp_path):
    d = _render_table(TableCell(text="plain"), tmp_path)
    tcPr = d.tables[0].cell(0, 0)._tc.tcPr
    assert tcPr is None or tcPr.find(qn("w:shd")) is None  # regression: no phantom fill


def test_hex6_validation():
    assert R._hex6("EFEFEF") == "EFEFEF"
    assert R._hex6("ababab") == "ABABAB"
    assert R._hex6("xyz") is None
    assert R._hex6("12345") is None      # too short
    assert R._hex6(None) is None
    assert R._cell_align("center") == "center"
    assert R._cell_align("bogus") is None
    assert R._cell_valign("bottom") == "bottom"
    assert R._cell_valign("middle") is None  # not a valid Word value


def _vstack(chars, x=70.0, y0=100.0, dy=15.0):
    """Spans for a vertically-stacked label: one single char per line at the same x."""
    return {f"p0b0l{i}s0": RawSpan(id=f"p0b0l{i}s0", page=0, text=ch,
                                   bbox=(x, y0 + i * dy, x + 10, y0 + i * dy + 12))
            for i, ch in enumerate(chars)}


def test_absorb_vertical_label_fills_empty_rowspan_cell():
    by_id = _vstack(["개", "정", "내", "역"])
    cell = TableCell(text="", rowspan=4, is_header=True)
    metas = [{"cell": cell, "yr": (95.0, 200.0), "left_x": 70.0}]
    referenced = set()
    R._absorb_vertical_labels(list(by_id.values()), referenced, by_id, metas)
    assert cell.text == "개정내역" and cell.vertical is True
    assert set(cell.source_span_ids) == set(by_id)
    assert referenced == set(by_id)  # absorbed spans are marked, so they won't leak


def test_absorb_skips_non_vertical_text():
    # 2+ chars per line is real prose, not a stacked label -> not absorbed
    spans = [RawSpan(id=f"p0b0l{i}s0", page=0, text=f"word{i}",
                     bbox=(70, 100 + i * 15, 140, 112 + i * 15)) for i in range(4)]
    by_id = {s.id: s for s in spans}
    cell = TableCell(text="", rowspan=4, is_header=True)
    referenced = set()
    R._absorb_vertical_labels(spans, referenced, by_id, [{"cell": cell, "yr": (95.0, 200.0), "left_x": 70.0}])
    assert cell.text == "" and referenced == set()


def test_absorb_skips_outside_table_yrange():
    by_id = _vstack(["개", "정", "내", "역"])
    cell = TableCell(text="", rowspan=4, is_header=True)
    R._absorb_vertical_labels(list(by_id.values()), set(), by_id,
                              [{"cell": cell, "yr": (300.0, 400.0), "left_x": 70.0}])
    assert cell.text == ""


def test_absorb_skips_when_not_left_edge():
    by_id = _vstack(["개", "정", "내", "역"], x=300.0)  # far from the table's left edge (70)
    cell = TableCell(text="", rowspan=4, is_header=True)
    R._absorb_vertical_labels(list(by_id.values()), set(), by_id,
                              [{"cell": cell, "yr": (95.0, 200.0), "left_x": 70.0}])
    assert cell.text == ""


def test_absorb_noop_without_empty_cells():
    by_id = _vstack(["개", "정", "내", "역"])
    R._absorb_vertical_labels(list(by_id.values()), set(), by_id, [])  # no cells -> no error


def test_docx_extractor_reads_cell_shading_valign_align(tmp_path):
    src = Document()
    t = src.add_table(rows=1, cols=1)
    tc = t.cell(0, 0)._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), "DDEEFF"); tcPr.append(shd)
    va = OxmlElement("w:vAlign"); va.set(qn("w:val"), "center"); tcPr.append(va)
    p = t.cell(0, 0).paragraphs[0]
    p.add_run("x")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    assert _tc_shading(tc) == "DDEEFF"
    assert _tc_valign(tc) == "center"
    assert _tc_align(tc) == "right"
    # auto/white fill is treated as no shading
    shd.set(qn("w:fill"), "auto")
    assert _tc_shading(tc) is None


# ---- Sparse "flowchart" tables: empty cells render borderless --------------------

def _render_rows(rows, tmp_path):
    doc = SemanticDocument(blocks=[TableBlock(rows=rows)])
    out = str(tmp_path / "tb.docx")
    render_docx(doc, out, str(tmp_path))
    return Document(out)


def _top_is_nil(tc_el):
    tb = tc_el.find(".//" + qn("w:tcBorders"))
    top = tb.find(qn("w:top")) if tb is not None else None
    return top is not None and top.get(qn("w:val")) == "nil"


def test_sparse_flowchart_table_clears_empty_cell_borders(tmp_path):
    # HWP test-sequence flowcharts arrive as mostly-empty grids; the empty cells should
    # render borderless (leaving the labelled boxes) instead of as a broken empty grid.
    rows = [[TableCell(text="박스" if (r, c) == (1, 1) else "") for c in range(4)]
            for r in range(4)]
    t = _render_rows(rows, tmp_path).tables[0]
    cells = [tc for row in t.rows for tc in row._tr.findall(qn("w:tc"))]
    assert sum(_top_is_nil(tc) for tc in cells) == 15      # 15 empty of 16 -> borderless
    # the single labelled box keeps its border
    labelled = t.rows[1]._tr.findall(qn("w:tc"))[1]
    assert not _top_is_nil(labelled)


def test_dense_table_keeps_all_borders(tmp_path):
    # A normal, mostly-filled table is left fully gridded (not mistaken for a flowchart).
    rows = [[TableCell(text=f"r{r}c{c}") for c in range(4)] for r in range(4)]
    t = _render_rows(rows, tmp_path).tables[0]
    assert all(not _top_is_nil(tc)
               for row in t.rows for tc in row._tr.findall(qn("w:tc")))


# ---- Wide tables shrink cell text so they stay legible in portrait A4 -------------

def test_wide_table_shrinks_cell_font(tmp_path):
    from docx.shared import Pt
    rows = [[TableCell(text=f"c{c}") for c in range(9)] for _ in range(2)]
    t = _render_rows(rows, tmp_path).tables[0]
    sizes = [run.font.size for row in t.rows for cell in row.cells
             for run in cell.paragraphs[0].runs]
    assert sizes and all(s == Pt(8) for s in sizes)        # 9 cols -> 8pt


def test_narrow_table_keeps_default_font(tmp_path):
    rows = [[TableCell(text=f"c{c}") for c in range(4)] for _ in range(2)]
    t = _render_rows(rows, tmp_path).tables[0]
    sizes = [run.font.size for row in t.rows for cell in row.cells
             for run in cell.paragraphs[0].runs]
    assert all(s is None for s in sizes)                   # inherits style default


# ---- Recovered column widths size the grid (vertical lines match the source) -----

def test_col_widths_size_grid_columns(tmp_path):
    rows = [[TableCell(text="a"), TableCell(text="b"), TableCell(text="c")]]
    doc = SemanticDocument(blocks=[TableBlock(rows=rows, col_widths=[0.6, 0.25, 0.15])])
    out = str(tmp_path / "cw.docx")
    render_docx(doc, out, str(tmp_path))
    t = Document(out).tables[0]
    grid = t._tbl.find(qn("w:tblGrid"))
    ws = [int(c.get(qn("w:w"))) for c in grid.findall(qn("w:gridCol"))]
    assert len(ws) == 3
    fr = [w / sum(ws) for w in ws]
    assert abs(fr[0] - 0.6) < 0.03 and abs(fr[1] - 0.25) < 0.03 and abs(fr[2] - 0.15) < 0.03
    layout = t._tbl.tblPr.find(qn("w:tblLayout"))
    assert layout is not None and layout.get(qn("w:type")) == "fixed"


def test_no_col_widths_keeps_autofit(tmp_path):
    doc = SemanticDocument(blocks=[TableBlock(rows=[[TableCell(text="a"), TableCell(text="b")]])])
    out = str(tmp_path / "nf.docx")
    render_docx(doc, out, str(tmp_path))
    t = Document(out).tables[0]
    layout = t._tbl.tblPr.find(qn("w:tblLayout"))  # autofit, not fixed-width
    assert layout is None or layout.get(qn("w:type")) != "fixed"


# ---- Irregular tables rebuilt from the real PDF grid (find_tables) ----------------

def test_grid_span_counts_cells_between_lines():
    from pipeline.extractor import _grid_span
    lines = [58, 104, 149, 225, 293, 369, 478, 538]
    assert _grid_span(58, 149, lines) == 2     # spans two columns
    assert _grid_span(225, 478, lines) == 3    # a wide merged "개정내용" cell
    assert _grid_span(104, 149, lines) == 1


def test_fractions_and_rebuild_helpers():
    from pipeline.reconciler import _fractions, _rebuild_table_rows, _should_rebuild
    assert _fractions([0, 60, 100]) == [0.6, 0.4]
    assert _fractions([10, 10]) is None        # zero-width column → None
    rows = _rebuild_table_rows(
        [[{"text": "a", "colspan": 2, "rowspan": 1},
          {"text": "b", "colspan": 5, "rowspan": 3}]])
    assert rows[0][0].text == "a" and rows[0][0].colspan == 2
    assert rows[0][1].colspan == 5 and rows[0][1].rowspan == 3
    solid = [[{"text": "x", "colspan": 1, "rowspan": 1}]] * 3
    assert _should_rebuild(solid, [[1], [1]]) is True
    assert _should_rebuild([], [[1], [1]]) is False
    assert _should_rebuild([[{"text": "", "colspan": 1, "rowspan": 1}]] * 3, [[1]]) is False


def test_best_grid_picks_overlapping_table():
    from pipeline.reconciler import _best_grid
    grids = [((0, 0, 100, 50), [0, 50, 100], [[{"text": "a", "colspan": 1, "rowspan": 1}]]),
             ((0, 200, 100, 250), [0, 40, 100], [[{"text": "b", "colspan": 1, "rowspan": 1}]])]
    gbbox, bounds, struct = _best_grid(grids, (5, 5, 90, 45))
    assert gbbox == (0, 0, 100, 50) and bounds == [0, 50, 100]
    assert _best_grid(grids, None) == (None, None, None)
    assert _best_grid([], (0, 0, 10, 10)) == (None, None, None)


# ---- Table border mask: union the detected grid bbox, else pad the text range -----

def test_table_occupied_range_unions_detected_grid():
    # With a detected grid, the mask spans the union of the text y-range and the grid bbox
    # (which covers the real border lines the text stops short of).
    from pipeline.reconciler import _table_occupied_range
    assert _table_occupied_range((166.0, 730.8), (57.8, 157.1, 537.9, 740.2)) == (157.1, 740.2)


def test_table_occupied_range_pads_text_when_grid_missing():
    # No detected grid → pad the text range by a constant so the table's own borders (≈ one
    # cell-pad outside the text) stay masked, while a genuine separator a blank line away
    # still becomes a real divider. This is the S1 generalization fix (gbbox-None fallback).
    from pipeline.reconciler import _table_occupied_range, _suppressed, _TABLE_BORDER_PAD
    yr = (166.0, 730.8)
    rng = _table_occupied_range(yr, None)
    assert rng == (166.0 - _TABLE_BORDER_PAD, 730.8 + _TABLE_BORDER_PAD)
    occ = [rng]
    assert _suppressed(166.0 - 9, occ) and _suppressed(730.8 + 9, occ)  # borders masked
    assert not _suppressed(730.8 + 30, occ)                             # real separator kept


# ---- Grid rebuild is gated on the detected cells matching the vision table --------

def test_rebuild_text_overlap_high_for_same_table():
    # Same source text (whitespace differences ignored) → high overlap → trust the grid.
    from pipeline.reconciler import _rebuild_text_overlap
    vision = [[TableCell(text="기본규격번호"), TableCell(text="(KFS)-1-0011")],
              [TableCell(text="장비명"), TableCell(text="소방사다리차")]]
    rebuilt = [[TableCell(text="기본규격번호"), TableCell(text=""), TableCell(text="(KFS)-1-0011")],
               [TableCell(text="장 비 명"), TableCell(text="소방사다리차"), TableCell(text="")]]
    assert _rebuild_text_overlap(rebuilt, vision) >= 0.5


def test_rebuild_text_overlap_low_for_wrong_region():
    # find_tables locked onto a different region → disjoint texts → keep the vision table.
    from pipeline.reconciler import _rebuild_text_overlap
    vision = [[TableCell(text="기본규격번호"), TableCell(text="소방사다리차")]]
    rebuilt = [[TableCell(text="전혀"), TableCell(text="다른"), TableCell(text="표")]]
    assert _rebuild_text_overlap(rebuilt, vision) < 0.5


def test_rebuild_text_overlap_empty_rebuild_is_zero():
    from pipeline.reconciler import _rebuild_text_overlap
    assert _rebuild_text_overlap([[TableCell(text="")]], [[TableCell(text="x")]]) == 0.0
