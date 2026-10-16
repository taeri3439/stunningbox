import json, os
from docx import Document
from pipeline import process


def test_office_ingest_writes_source_json(tmp_path):
    # A .docx input means convert.to_docx is a no-op (no LibreOffice needed).
    d = Document()
    d.add_heading("Hello", level=1)
    d.add_paragraph("World")
    storage = str(tmp_path / "storage")
    doc_id = "doc123"
    base = process.doc_dir(storage, doc_id)
    os.makedirs(base, exist_ok=True)
    src = os.path.join(base, "source.docx")
    d.save(src)

    stages = []
    model, report = process.build_source_model_from_office(
        src, doc_id, storage, source_lang="ko",
        on_progress=lambda s, k, t: stages.append(s),
    )
    assert any(b.text == "Hello" for b in model.blocks)
    assert os.path.exists(os.path.join(base, "source.json"))
    saved = json.load(open(os.path.join(base, "source.json"), encoding="utf-8"))
    assert saved["blocks"]
    assert "convert" in stages and "extract" in stages
