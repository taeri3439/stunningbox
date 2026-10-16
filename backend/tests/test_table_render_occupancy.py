"""Row-grouped table rendering (fidelity wave 2, workstream A).

Diagnosis on the real watergun fixture (backend/storage/ahTZyKrBSX7bj_D4KLER6iPl, p21
0-indexed "마개 게이지" plug-gauge table) traced the visible zig-zag/column-shift defect
to its root cause: `_add_table`'s grid-occupancy placement (the `occupied` set below) is
ALREADY correct — it places each row's cells into the first free column and skips columns
covered by an active rowspan, exactly as this workstream's spec calls for. The synthetic
reproduction below (mirroring the spec's own example: a rowspan=2 cell in column 0, next
row with one fewer cell) confirms this passes on unmodified code; it is kept here as a
regression lock, not as a fix target.

The REAL defect is a wrong `rowspan` VALUE from the vision analyzer: page_blocks.json
(cached, no API) declares rowspan=4 for each nominal-size cell (40/65/75/90/100/125/150)
when the source table (verified against the rendered PDF page) groups each size into only
2 rows (통과/정지). Because the resulting column overshoot is CONSISTENT (every affected
row credits the same wrong extra column via carry, rather than fluctuating), reconciler's
`_table_structure_unreliable` ragged-fraction check (>=30% of rows disagree with the modal
width) doesn't trip — only 3/15 rows differ from the modal, 20% < 30% — so the table was
never rerouted to the (already-existing, already-tested) image-fallback path that
`test_table_fallback_watergun.py` documents as the intended defense for exactly this
"self-contradictory rowspan" shape. That gap is fixed in `_table_structure_unreliable`
(see `test_table_fallback.py::test_oversized_rowspan_double_claim_is_unreliable`): a modal
credited width that EXCEEDS the table's own real column count (any row's own colspan sum,
ignoring carry) is now also treated as unreliable, regardless of the ragged fraction. A
targeted regression sweep over every cached fixture's every table (154 tables across 22
docs in backend/storage) confirmed this new check flips exactly one verdict — this table —
and zero others, so no existing table was disturbed.

This file therefore verifies the two things that ARE renderer-side and in scope for this
workstream: (1) occupancy-aware placement was already correct (regression lock), and
(2) rows now never split mid-cell across a page break (`w:cantSplit`), which was missing
and is a genuine fix.
"""
from docx import Document
from docx.oxml.ns import qn

from models import TableBlock, TableCell
from renderers.docx_renderer import _add_table


def _tr_pr_children(tr, tag):
    trPr = tr.find(qn("w:trPr"))
    if trPr is None:
        return []
    return trPr.findall(qn(tag))


def test_rowspan_carry_places_next_row_in_free_columns():
    """Mirrors the spec's own example: col-0 cell has rowspan=2 in row 0; row 1 has ONE
    FEWER cell (its own column-0 slot is covered by row 0's span). Row 1's OWN cells must
    land in columns 1..N, not shift back into column 0 (which read-back reports as the
    rowspan-owner's text via the merge, not a duplicate/misplaced value)."""
    block = TableBlock(
        rows=[
            [TableCell(text="A"), TableCell(text="B"), TableCell(text="C")],
            [TableCell(text="D"), TableCell(text="E")],
        ]
    )
    block.rows[0][0].rowspan = 2

    d = Document()
    _add_table(d, block)
    t = d.tables[0]

    assert len(t.rows) == 2
    assert len(t.columns) == 3
    # Row 0: the declared cells, unmerged text as authored.
    assert [c.text for c in t.rows[0].cells] == ["A", "B", "C"]
    # Row 1: column 0 reads back as "A" (python-docx reports a merged cell's text on every
    # cell it spans) — the important assertion is that row 1's OWN cells ("D", "E") sit at
    # columns 1 and 2, not columns 0 and 1.
    row1_texts = [c.text for c in t.rows[1].cells]
    assert row1_texts[0] == "A"          # covered by the row-0 rowspan, not overwritten
    assert row1_texts[1:] == ["D", "E"]  # row 1's own cells occupy columns 1..N


def test_every_row_gets_cant_split():
    """No `<w:tr>` may split mid-cell across a page boundary — every row's trPr must carry
    `<w:cantSplit/>` so Word pushes the whole row to the next page instead."""
    block = TableBlock(
        rows=[
            [TableCell(text="h1", is_header=True), TableCell(text="h2", is_header=True)],
            [TableCell(text="a1"), TableCell(text="a2")],
            [TableCell(text="b1"), TableCell(text="b2")],
        ]
    )

    d = Document()
    _add_table(d, block)
    t = d.tables[0]

    for row in t.rows:
        cant_split = _tr_pr_children(row._tr, "w:cantSplit")
        assert len(cant_split) == 1, "every row must carry exactly one w:cantSplit"


def test_header_row_gets_tbl_header_when_first_row_is_a_header():
    """Nice-to-have: a first row whose cells are marked is_header repeats on each page."""
    block = TableBlock(
        rows=[
            [TableCell(text="h1", is_header=True), TableCell(text="h2", is_header=True)],
            [TableCell(text="a1"), TableCell(text="a2")],
        ]
    )

    d = Document()
    _add_table(d, block)
    t = d.tables[0]

    assert len(_tr_pr_children(t.rows[0]._tr, "w:tblHeader")) == 1
    assert len(_tr_pr_children(t.rows[1]._tr, "w:tblHeader")) == 0


def test_no_header_row_skips_tbl_header():
    """A table with no is_header cells at all never gets a (meaningless) w:tblHeader."""
    block = TableBlock(
        rows=[
            [TableCell(text="a1"), TableCell(text="a2")],
            [TableCell(text="b1"), TableCell(text="b2")],
        ]
    )

    d = Document()
    _add_table(d, block)
    t = d.tables[0]

    for row in t.rows:
        assert len(_tr_pr_children(row._tr, "w:tblHeader")) == 0
