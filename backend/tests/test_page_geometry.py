"""Source page geometry (A4 size + margins) survives model -> renderer, and the PDF
reconciler captures it. Guards the page-ballooning fix (A4 source must not render as
US Letter)."""
import os

from docx import Document

from models import ParagraphBlock, RawDocument, RawPage, RawSpan, SemanticDocument
from pipeline.reconciler import _page_geometry
from renderers.docx_renderer import render_docx

A4_W, A4_H = 595.0, 841.0


def _raw_a4_page():
    spans = [
        RawSpan(id=f"p0b0l{i}s0", page=0, text=f"line {i}",
                bbox=(70.0, 80.0 + i * 22, 520.0, 92.0 + i * 22), size=11.0)
        for i in range(20)
    ]
    return RawDocument(pages=[RawPage(width=A4_W, height=A4_H, spans=spans)])


def test_reconciler_captures_a4_and_margins():
    geo = _page_geometry(_raw_a4_page())
    assert geo["page_width"] == A4_W and geo["page_height"] == A4_H
    # left margin ~70pt (text starts at x0=70); right margin ~ 595-520=75pt; both clamped sane
    assert 18.0 <= geo["margin_left"] <= 108.0
    assert 60.0 <= geo["margin_left"] <= 80.0
    assert 60.0 <= geo["margin_right"] <= 90.0


def test_empty_raw_returns_no_geometry():
    assert _page_geometry(RawDocument(pages=[])) == {}


def test_renderer_applies_a4_size_and_margins(tmp_path):
    doc = SemanticDocument(
        blocks=[ParagraphBlock(text="hello")],
        page_width=A4_W, page_height=A4_H,
        margin_left=72.0, margin_right=72.0, margin_top=54.0, margin_bottom=54.0,
    )
    out = str(tmp_path / "a4.docx")
    render_docx(doc, out, str(tmp_path))
    sec = Document(out).sections[0]
    assert round(sec.page_width.pt) == round(A4_W)   # not US Letter (612)
    assert round(sec.page_height.pt) == round(A4_H)  # not US Letter (792)
    assert round(sec.left_margin.pt) == 72
    assert round(sec.top_margin.pt) == 54


def test_renderer_without_geometry_keeps_defaults(tmp_path):
    # Legacy model (no page_width/height) must still render, using python-docx defaults.
    doc = SemanticDocument(blocks=[ParagraphBlock(text="hi")])
    out = str(tmp_path / "legacy.docx")
    render_docx(doc, out, str(tmp_path))
    assert os.path.exists(out)
    sec = Document(out).sections[0]
    assert sec.page_width.pt > 0  # whatever the default is, it's valid
