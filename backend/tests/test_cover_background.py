"""Phase 1 — cover/first-page background."""
import os

from PIL import Image
from docx import Document
from docx.shared import Pt
from models import SemanticDocument, HeadingBlock, ParagraphBlock
import fitz

from pipeline import extractor
from pipeline.reconciler import reconcile
from renderers.docx_renderer import render_docx


def test_cover_background_defaults_none():
    doc = SemanticDocument()
    assert doc.cover_background is None


def test_cover_background_roundtrips():
    doc = SemanticDocument(cover_background="cover_bg.png")
    loaded = SemanticDocument.model_validate(doc.model_dump())
    assert loaded.cover_background == "cover_bg.png"


def _cover_pdf(tmp_path):
    doc = fitz.open()
    page = doc.new_page(width=300, height=400)
    page.draw_rect(fitz.Rect(0, 0, 300, 400), fill=(0.1, 0.3, 0.5))  # full-page design
    page.insert_text((40, 120), "VETERINARY RAPID TEST MARKET", fontsize=18, color=(1, 1, 1))
    p = str(tmp_path / "cover.pdf")
    doc.save(p); doc.close()
    return p


def test_reconcile_sets_cover_background_and_suppresses_inline_figure(tmp_path):
    pdf = _cover_pdf(tmp_path)
    ad = str(tmp_path / "assets")
    raw = extractor.extract_document(pdf, ad)
    title_ids = [s.id for s in raw.pages[0].spans]  # the title span(s)
    page_blocks = [[
        {"type": "figure", "figure_bbox": [0, 0, 300, 400], "source_span_ids": []},
        {"type": "heading", "text": "VETERINARY RAPID TEST MARKET", "source_span_ids": title_ids},
    ]]
    doc, _report = reconcile(raw, page_blocks, pdf, ad, source_lang="en")

    assert doc.cover_background, "cover_background should be set"
    assert (tmp_path / "assets" / doc.cover_background).exists()
    assert not [b for b in doc.blocks if b.type == "figure"], "inline cover figure suppressed"
    assert any(b.type == "heading" for b in doc.blocks), "title stays a translatable heading"


def test_reconcile_small_figure_on_page0_stays_inline(tmp_path):
    doc = fitz.open()
    page = doc.new_page(width=300, height=400)
    page.draw_rect(fitz.Rect(20, 20, 120, 90), fill=(0, 0, 0))  # small figure (<50% of page)
    p = str(tmp_path / "small.pdf")
    doc.save(p); doc.close()
    ad = str(tmp_path / "assets")
    raw = extractor.extract_document(p, ad)
    page_blocks = [[{"type": "figure", "figure_bbox": [20, 20, 120, 90], "source_span_ids": []}]]
    out, _ = reconcile(raw, page_blocks, p, ad, source_lang="en")
    assert out.cover_background is None
    assert [b for b in out.blocks if b.type == "figure"], "small figure stays inline"


def test_render_docx_anchors_cover_background(tmp_path):
    ad = str(tmp_path)
    Image.new("RGB", (300, 400), (20, 80, 130)).save(os.path.join(ad, "cover_bg.png"))
    doc = SemanticDocument(source_lang="en", cover_background="cover_bg.png",
                           page_width=300.0, page_height=400.0)
    out = str(tmp_path / "out.docx")
    render_docx(doc, out, ad)

    d2 = Document(out)
    sec = d2.sections[0]
    assert sec.different_first_page_header_footer
    xml = sec.first_page_header._element.xml
    assert 'behindDoc="1"' in xml           # background sits behind the text
    assert str(int(Pt(400.0))) in xml       # full-page height extent (EMU)


def test_reconcile_cover_with_split_subhalf_figures(tmp_path):
    # Real covers: a colored full-bleed page the analyzer splits into multiple sub-50%
    # figures. The page is mostly non-white → must be detected as a cover anyway.
    doc = fitz.open()
    page = doc.new_page(width=300, height=400)
    page.draw_rect(fitz.Rect(0, 0, 300, 400), fill=(0.15, 0.2, 0.3))  # full colored design
    page.insert_text((40, 100), "VETERINARY RAPID TEST MARKET", fontsize=16, color=(1, 1, 1))
    p = str(tmp_path / "split.pdf")
    doc.save(p); doc.close()
    ad = str(tmp_path / "assets")
    raw = extractor.extract_document(p, ad)
    title_ids = [s.id for s in raw.pages[0].spans]
    page_blocks = [[
        {"type": "figure", "figure_bbox": [0, 0, 300, 150], "source_span_ids": []},   # 37%
        {"type": "figure", "figure_bbox": [0, 250, 300, 400], "source_span_ids": []},  # 37%
        {"type": "heading", "text": "VETERINARY RAPID TEST MARKET", "source_span_ids": title_ids},
    ]]
    out, _ = reconcile(raw, page_blocks, p, ad, source_lang="en")
    assert out.cover_background, "colored cover detected even with split sub-50% figures"
    assert not [b for b in out.blocks if b.type == "figure"], "all cover figures suppressed"
    assert any(b.type == "heading" for b in out.blocks), "title stays a translatable heading"


def test_reconcile_textless_cover_drops_vision_blocks(tmp_path):
    # A fully-designed cover (colored page, title is non-extractable outlined art → 0 spans).
    # The vision analyzer still "reads" the title as a heading; those blocks must be dropped so
    # the translated text doesn't double the title baked into the cover image.
    doc = fitz.open()
    page = doc.new_page(width=300, height=400)
    page.draw_rect(fitz.Rect(0, 0, 300, 400), fill=(0.15, 0.2, 0.3))  # colored cover, NO text
    p = str(tmp_path / "textless.pdf")
    doc.save(p); doc.close()
    ad = str(tmp_path / "assets")
    raw = extractor.extract_document(p, ad)
    assert raw.pages[0].spans == []  # fixture sanity: cover has no extractable text
    page_blocks = [[
        {"type": "figure", "figure_bbox": [0, 0, 300, 400], "source_span_ids": []},
        {"type": "heading", "text": "VETERINARY RAPID TEST MARKET", "source_span_ids": []},
    ]]
    out, report = reconcile(raw, page_blocks, p, ad, source_lang="en")
    assert out.cover_background, "textless cover still detected as a background"
    assert not out.blocks, "vision-read blocks on a textless cover are dropped (kept as image)"
    assert any("cover page 0: dropped" in w for w in report.warnings)  # drop is surfaced, not silent


def test_render_docx_breaks_after_cover(tmp_path):
    ad = str(tmp_path)
    Image.new("RGB", (300, 400), (20, 80, 130)).save(os.path.join(ad, "cover_bg.png"))
    doc = SemanticDocument(
        source_lang="en", cover_background="cover_bg.png",
        page_width=300.0, page_height=400.0,
        blocks=[HeadingBlock(text="cover title", page=0),
                ParagraphBlock(text="body intro", page=1)],
    )
    out = str(tmp_path / "out.docx")
    render_docx(doc, out, ad)
    body_xml = Document(out).element.body.xml
    assert 'w:type="page"' in body_xml, "a page break isolates the cover on page 1"


def test_render_docx_no_break_without_cover(tmp_path):
    doc = SemanticDocument(
        source_lang="en",
        blocks=[HeadingBlock(text="t", page=0), ParagraphBlock(text="b", page=1)],
    )
    out = str(tmp_path / "out.docx")
    render_docx(doc, out, str(tmp_path))
    assert 'w:type="page"' not in Document(out).element.body.xml, \
        "no source-page breaks when there's no cover (unchanged behaviour)"


def test_reconcile_image_heavy_content_page0_not_cover(tmp_path):
    # page 0 is mostly colored (passes the non-white check) BUT has many body blocks → a real
    # CONTENT page, not a cover: must NOT be treated as a cover (figure kept, nothing dropped).
    doc = fitz.open()
    page = doc.new_page(width=300, height=400)
    page.draw_rect(fitz.Rect(0, 0, 300, 220), fill=(0.2, 0.3, 0.5))  # big banner (>40% non-white)
    page.insert_text((20, 300), "real body content here", fontsize=10)
    p = str(tmp_path / "content.pdf")
    doc.save(p); doc.close()
    ad = str(tmp_path / "assets")
    raw = extractor.extract_document(p, ad)
    page_blocks = [[
        {"type": "figure", "figure_bbox": [0, 0, 300, 220], "source_span_ids": []},
        {"type": "paragraph", "text": "p1", "source_span_ids": []},
        {"type": "paragraph", "text": "p2", "source_span_ids": []},
        {"type": "paragraph", "text": "p3", "source_span_ids": []},
        {"type": "paragraph", "text": "p4", "source_span_ids": []},
    ]]
    out, _ = reconcile(raw, page_blocks, p, ad, source_lang="en")
    assert out.cover_background is None, "image-heavy content page (many body blocks) is not a cover"
    assert [b for b in out.blocks if b.type == "figure"], "its figure is kept, not suppressed"
