"""Task 5: end-to-end wiring check for editable-figure-text.

Builds a tiny SemanticDocument with one overlay FigureBlock, renders it to a real
.docx via the full render_docx path, and reopens it with python-docx to confirm the
editable OOXML group + overlay text are present. Then mutates overlays[0].target_text
(exactly what re-translating an `overlay:{oi}` segment does via models.iter_segments'
_overlay_accessor setter) and re-renders, proving the new text reflects and the old
text is gone — the re-translation round-trip the whole feature exists for.
"""
from docx import Document
from PIL import Image

from models import FigureBlock, FigureOverlay, SemanticDocument
from renderers.docx_renderer import render_docx


def test_overlay_edit_reflects_on_rerender(tmp_path):
    Image.new("RGB", (300, 200), "white").save(tmp_path / "b.png")
    fig = FigureBlock(
        block_id="F", image_ref="b.png", base_image_ref="b.png",
        width=216.0, height=144.0,
        overlays=[
            FigureOverlay(id="o0", target_text="처음", bbox=(20, 20, 120, 24), font_size=16),
            FigureOverlay(id="o1", target_text="둘째", bbox=(20, 60, 120, 24), font_size=16),
        ],
    )
    doc = SemanticDocument(blocks=[fig], target_lang="ko")
    out = str(tmp_path / "o.docx")

    render_docx(doc, out, str(tmp_path))
    xml = Document(out).element.body.xml
    assert "<wpg:wgp" in xml  # editable group, not a flattened composite image
    assert "처음" in xml and "둘째" in xml

    # Simulate re-translation: iter_segments' overlay setter writes overlays[oi].target_text.
    fig.overlays[0].target_text = "다시"
    render_docx(doc, out, str(tmp_path))
    xml = Document(out).element.body.xml
    assert "다시" in xml and "둘째" in xml
    assert "처음" not in xml
