"""Workstream B (2026-07-02 fidelity wave 2): cover vertical rhythm.

A designed/text cover's blocks (title, subtitle, doc-number line, ...) carry no y geometry
once reconciled into the semantic model, so a large deliberate gap in the source (e.g. the
title sitting ~1/3 down the page) collapses into ordinary top-flush paragraph flow on
render. The reconciler now measures each page-0-cover block's top gap and records it as
`space_before` (pt); the docx renderer reproduces it as `paragraph_format.space_before`.
"""
import os

import fitz
import pytest
from docx import Document

import pipeline.reconciler as reconciler_mod
from models import ParagraphBlock, SemanticDocument
from pipeline import extractor, figure_overlay
from pipeline.reconciler import reconcile
from renderers.docx_renderer import render_docx

_TITLE = "COVER TITLE"


def _two_block_cover_pdf(tmp_path):
    """A text cover with a title far below the page top, then a subtitle further down
    still — mirrors the real ladder cover's deliberate title placement + spaced-out
    subtitle/doc-number block. The subtitle ends with a period so the (pre-existing,
    out-of-scope) cross-page paragraph stitcher doesn't fold the page-1 body paragraph
    into it (it merges a paragraph that doesn't end a sentence with the next page's)."""
    doc = fitz.open()
    p0 = doc.new_page(width=600, height=800)
    p0.insert_text((60, 250), _TITLE, fontsize=42)  # title starts well below the top
    p0.insert_text((60, 500), "Subtitle line.", fontsize=14)  # large gap below the title
    p1 = doc.new_page(width=600, height=800)
    p1.insert_text((60, 60), "Body paragraph text on page two.", fontsize=11)
    path = str(tmp_path / "rhythmcover.pdf")
    doc.save(path)
    doc.close()
    return path


def _reconcile_two_block_cover(tmp_path, monkeypatch):
    pdf = _two_block_cover_pdf(tmp_path)
    ad = str(tmp_path / "assets")
    raw = extractor.extract_document(pdf, ad)
    monkeypatch.setattr(reconciler_mod, "_page_nonwhite_frac", lambda *a, **k: 0.05)
    monkeypatch.setattr(figure_overlay, "render_redacted_clip", lambda *a, **k: "cover_bg.png")

    heading_ids = [s.id for s in raw.pages[0].spans if _TITLE in s.text]
    sub_ids = [s.id for s in raw.pages[0].spans if s.text.strip() == "Subtitle line."]
    body_ids = [s.id for s in raw.pages[1].spans]
    assert heading_ids and sub_ids and body_ids  # fixture sanity

    page_blocks = [
        [
            {"type": "heading", "text": _TITLE, "source_span_ids": heading_ids},
            {"type": "paragraph", "text": "Subtitle line.", "source_span_ids": sub_ids},
        ],
        [
            {"type": "paragraph", "text": "Body paragraph text on page two.",
             "source_span_ids": body_ids},
        ],
    ]
    return reconcile(raw, page_blocks, pdf, ad, source_lang="ko")


def test_cover_second_block_gets_measured_space_before(tmp_path, monkeypatch):
    doc, _report = _reconcile_two_block_cover(tmp_path, monkeypatch)
    assert doc.cover_background == "cover_bg.png"  # fixture sanity: cover was detected

    page0 = [b for b in doc.blocks if getattr(b, "page", None) == 0]
    title = next(b for b in page0 if getattr(b, "text", "") == _TITLE)
    subtitle = next(b for b in page0 if getattr(b, "text", "") == "Subtitle line.")

    # The subtitle's gap = its own top y minus the title's measured bottom y (span bbox
    # y1) ~ 484.95 - 262.56 ~= 222.4pt — a real, large, unambiguous design gap.
    assert subtitle.space_before is not None
    assert 190.0 < subtitle.space_before < 260.0

    # The title itself sits far enough below the estimated document top margin (~48pt for
    # this fixture) that it also picks up a (smaller but still clear) measured gap.
    assert title.space_before is not None
    assert title.space_before > 100.0


def test_cover_space_before_clamped_and_noise_skipped(tmp_path, monkeypatch):
    """Values are clamped to [0, 300] and gaps <= 6pt are left unset (noise, not a
    deliberate design gap)."""
    pdf = _two_block_cover_pdf(tmp_path)
    ad = str(tmp_path / "assets")
    raw = extractor.extract_document(pdf, ad)
    monkeypatch.setattr(reconciler_mod, "_page_nonwhite_frac", lambda *a, **k: 0.05)
    monkeypatch.setattr(figure_overlay, "render_redacted_clip", lambda *a, **k: "cover_bg.png")

    heading_ids = [s.id for s in raw.pages[0].spans if _TITLE in s.text]
    sub_ids = [s.id for s in raw.pages[0].spans if s.text.strip() == "Subtitle line."]
    body_ids = [s.id for s in raw.pages[1].spans]

    page_blocks = [
        [
            {"type": "heading", "text": _TITLE, "source_span_ids": heading_ids},
            {"type": "paragraph", "text": "Subtitle line.", "source_span_ids": sub_ids},
        ],
        [
            {"type": "paragraph", "text": "Body paragraph text on page two.",
             "source_span_ids": body_ids},
        ],
    ]
    doc, _report = reconcile(raw, page_blocks, pdf, ad, source_lang="ko")
    page0 = [b for b in doc.blocks if getattr(b, "page", None) == 0]
    for b in page0:
        sb = getattr(b, "space_before", None)
        if sb is not None:
            assert 0.0 <= sb <= 300.0


def _content_page0_pdf(tmp_path):
    """A content-heavy (non-cover) page 0, same shape as test_text_cover.py's guard
    fixture, but with a large vertical gap between two paragraphs so a false-positive
    space_before would be caught."""
    doc = fitz.open()
    p0 = doc.new_page(width=300, height=400)
    p0.insert_text((230, 25), "KFS 0011", fontsize=10)
    p0.insert_text((40, 100), "Paragraph one.", fontsize=11)
    p0.insert_text((40, 300), "Paragraph two, far below.", fontsize=11)
    p0.insert_text((40, 320), "Paragraph three.", fontsize=11)
    p0.insert_text((40, 340), "Paragraph four.", fontsize=11)
    p1 = doc.new_page(width=300, height=400)
    p1.insert_text((40, 60), "Body paragraph text on page two.", fontsize=11)
    path = str(tmp_path / "content.pdf")
    doc.save(path)
    doc.close()
    return path


def test_non_cover_page_never_gets_space_before(tmp_path, monkeypatch):
    """A content-heavy (non-cover) page 0 never sets space_before anywhere in the
    document, even though it has a large vertical gap between two of its paragraphs —
    the rhythm capture is gated on `cover_bg is not None`."""
    pdf = _content_page0_pdf(tmp_path)
    ad = str(tmp_path / "assets")
    raw = extractor.extract_document(pdf, ad)
    monkeypatch.setattr(reconciler_mod, "_page_nonwhite_frac", lambda *a, **k: 0.05)

    chrome_ids = [s.id for s in raw.pages[0].spans if s.text.strip() == "KFS 0011"]
    para_texts = ["Paragraph one.", "Paragraph two, far below.", "Paragraph three.",
                  "Paragraph four."]
    para_blocks = [
        {"type": "paragraph", "text": t,
         "source_span_ids": [s.id for s in raw.pages[0].spans if s.text.strip() == t]}
        for t in para_texts
    ]
    body_ids = [s.id for s in raw.pages[1].spans]
    assert chrome_ids and all(b["source_span_ids"] for b in para_blocks) and body_ids

    page_blocks = [
        [{"type": "chrome", "text": "KFS 0011", "source_span_ids": chrome_ids}] + para_blocks,
        [{"type": "paragraph", "text": "Body paragraph text on page two.",
          "source_span_ids": body_ids}],
    ]
    doc, _report = reconcile(raw, page_blocks, pdf, ad, source_lang="ko")

    assert doc.cover_background is None, "content-heavy page 0 is not a cover"
    assert all(getattr(b, "space_before", None) is None for b in doc.blocks), (
        "no block anywhere gets space_before when page 0 isn't a cover"
    )


# --------------------------------------------------------------------------- #
# Renderer: space_before -> w:spacing w:before (docx paragraph_format.space_before)
# --------------------------------------------------------------------------- #


def test_renderer_applies_space_before(tmp_path):
    m = SemanticDocument(source_lang="ko", blocks=[
        ParagraphBlock(text="Gapped paragraph", block_id="B0", source_span_ids=["s"],
                        space_before=120.0),
        ParagraphBlock(text="Normal paragraph", block_id="B1", source_span_ids=["s2"]),
    ])
    out = os.path.join(str(tmp_path), "d.docx")
    render_docx(m, out, str(tmp_path))

    d = Document(out)
    paras = [p for p in d.paragraphs if p.text in ("Gapped paragraph", "Normal paragraph")]
    gapped = next(p for p in paras if p.text == "Gapped paragraph")
    normal = next(p for p in paras if p.text == "Normal paragraph")

    assert gapped.paragraph_format.space_before is not None
    assert gapped.paragraph_format.space_before.pt == pytest.approx(120.0)
    assert gapped.paragraph_format.space_before.twips == pytest.approx(2400, abs=1)
    # A block with no space_before renders with Word's normal (unset) spacing.
    assert normal.paragraph_format.space_before is None


# --------------------------------------------------------------------------- #
# Real-data acceptance (dev-local; skips when the storage copy is absent). Never mutates
# the original fixture at backend/storage/fsOdj49iXYJWO1L96wrXY1HP — operates on a
# separate copy, `_rhythm-verify-ladder`.
# --------------------------------------------------------------------------- #

STORAGE = os.path.join(os.path.dirname(__file__), "..", "storage")
DOC = "_rhythm-verify-ladder"


@pytest.mark.skipif(
    not os.path.exists(os.path.join(STORAGE, DOC, "page_blocks.json")),
    reason="dev-local cached fixture copy not present",
)
def test_ladder_cover_title_gets_a_large_measured_space_before():
    """The real ladder cover's title sits well below the page top (source design: title
    span y0=149.1pt on an 841pt page, doc top margin 54.6pt → 94.5pt of deliberate space
    above it). The title is NOT the page's first block — the 'KFS 0011' chrome-derived
    paragraph sits above it (cluster y≈105.5) and legitimately claims the first chunk of
    that gap. So the design intent ("title well below the top") is carried by the
    CUMULATIVE space_before of the blocks up to and including the title (measured 50.9 +
    43.6 = 94.5pt — exactly the title's true offset below the margin), not by the title's
    own gap alone. Values printed for the report."""
    from pipeline import process

    base = os.path.join(STORAGE, DOC)
    pdf = os.path.join(base, "source.pdf")
    doc, _report = process.rebuild_source_model(
        pdf, DOC, os.path.abspath(STORAGE), source_lang="ko"
    )

    assert doc.cover_background, "the ladder cover must still be recognized"
    page0 = [b for b in doc.blocks if getattr(b, "page", None) == 0]
    print("ladder page-0 space_before values:")
    for b in page0:
        print(f"  {getattr(b, 'type', '?'):10s} {(getattr(b, 'text', '') or '')[:40]!r:44s} "
              f"space_before={getattr(b, 'space_before', None)}")

    heading = next((b for b in page0 if getattr(b, "type", None) == "heading"), None)
    assert heading is not None, "the cover title must survive as a heading block"
    assert heading.space_before is not None and heading.space_before > 30.0, (
        "the title carries its own measured gap from the chrome line above it"
    )
    cumulative = 0.0
    for b in page0:
        cumulative += getattr(b, "space_before", None) or 0.0
        if b is heading:
            break
    assert cumulative > 50.0, (
        "the accumulated space above the title reproduces its position well below the top"
    )
