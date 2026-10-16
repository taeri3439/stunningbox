import os
from docx import Document
from docx.oxml.ns import qn
from models import FigureBlock, FigureOverlay
from renderers.docx_renderer import _add_figure
from PIL import Image


def _fig(tmp_path, overlays=True):
    base = tmp_path / "b.png"
    Image.new("RGB", (300, 200), "white").save(base)
    return FigureBlock(
        block_id="F1", image_ref="b.png", base_image_ref="b.png" if overlays else "",
        width=216.0, height=144.0,
        overlays=[FigureOverlay(id="o0", target_text="테스트 라벨", bbox=(20, 20, 120, 24), font_size=16)] if overlays else [],
    )


def test_overlay_figure_renders_group_with_textbox(tmp_path):
    d = Document()
    _add_figure(d, _fig(tmp_path), str(tmp_path), d.sections[0].page_width)
    xml = d.element.body.xml
    # NOTE: the raw substring "wordprocessingGroup" is a false-positive signal here —
    # python-docx's default template already declares xmlns:wpg="...wordprocessingGroup"
    # on every document, even a totally blank one (verified: Document().element.body.xml
    # contains it with zero figures). The actual, unambiguous signal that a group was
    # inserted is the "<wpg:wgp" element tag itself.
    assert "<wpg:wgp" in xml and "테스트 라벨" in xml
    assert 'name="figF1"' in xml


def test_missing_base_falls_back_to_composite(tmp_path):
    fig = _fig(tmp_path)
    fig.base_image_ref = "nonexistent.png"
    d = Document()
    p = _add_figure(d, fig, str(tmp_path), d.sections[0].page_width)
    xml = d.element.body.xml
    assert "<wpg:wgp" not in xml     # fell back
    # NOTE: raw xml.count("blip") over-counts — python-docx always wraps a blip in an
    # open+close <pic:blipFill> tag pair, so a single embedded picture yields 3 "blip"
    # substring hits, not 1. Count the actual <a:blip element instead.
    assert p is not None and xml.count("<a:blip ") == 1  # composite image embedded


def test_opaque_figure_unchanged(tmp_path):
    d = Document()
    _add_figure(d, _fig(tmp_path, overlays=False), str(tmp_path), d.sections[0].page_width)
    assert "<wpg:wgp" not in d.element.body.xml
