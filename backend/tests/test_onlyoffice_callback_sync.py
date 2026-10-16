import os, json
import models, main
from renderers.docx_renderer import render_docx
from pipeline import process
from docx import Document
from docx.oxml.ns import qn


def test_callback_sync_reflects_in_all_views(tmp_path, monkeypatch):
    """An editor text edit, synced on save, must reflect in EVERY model-based view: the
    preview/render (get_export_model), AND the 재번역 segment list + 용어/고유명사 audits
    (build_segments → get_translated_model). The sync therefore writes the single live model
    translated.{lang}.json (NOT a separate edited.json, which would leave the list stale)."""
    monkeypatch.setattr(main, "STORAGE", str(tmp_path))
    doc_id, lang = "docX", "en"
    d = os.path.join(str(tmp_path), doc_id)
    os.makedirs(os.path.join(d, "assets"), exist_ok=True)
    # source (Korean) + translated (the auto-translation / current live model)
    src = models.SemanticDocument(source_lang="ko", blocks=[
        models.ParagraphBlock(text="안녕 세상", block_id="p0", source_span_ids=["s0"]),
    ])
    tr = models.SemanticDocument(source_lang="ko", blocks=[
        models.ParagraphBlock(text="hello world", block_id="p0", source_span_ids=["s0"]),
    ])
    with open(os.path.join(d, "source.json"), "w", encoding="utf-8") as fh:
        json.dump(src.model_dump(), fh)
    with open(os.path.join(d, f"translated.{lang}.json"), "w", encoding="utf-8") as fh:
        json.dump(tr.model_dump(), fh)
    out = process.output_docx_path(str(tmp_path), doc_id, lang)
    render_docx(tr, out, os.path.join(d, "assets"))

    # simulate an OnlyOffice text edit in the saved docx
    dd = Document(out)
    for t in dd.element.body.iter(qn("w:t")):
        t.text = "edited world"; break
    dd.save(out)

    main._sync_model_from_saved_docx(doc_id, lang)

    # the single live model (translated.json) carries the edit → every consumer reflects it:
    assert process.get_translated_model(str(tmp_path), doc_id, lang).blocks[0].text == "edited world"
    assert process.get_export_model(str(tmp_path), doc_id, lang).blocks[0].text == "edited world"  # preview/render
    segs = process.build_segments(str(tmp_path), doc_id, lang)                                       # 재번역 list + audits
    assert segs and segs[0]["translated"] == "edited world"
    # Phase 1 does NOT create a separate edited.json (it would shadow translated.json for render
    # while build_segments stayed stale — the gap this test guards against).
    assert not os.path.exists(process.edited_path(str(tmp_path), doc_id, lang))
