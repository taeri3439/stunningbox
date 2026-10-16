import os

from docx import Document

from models import ParagraphBlock, SemanticDocument
from pipeline import process
from pipeline.text_replace import ReplaceRule
from renderers.docx_renderer import render_docx


def _rule(find, replace, save=True, match="smart"):
    return ReplaceRule(find=find, replace=replace, save=save, match=match)


def _seed(tmp_path, text="Rated Load test"):
    storage, doc_id, lang = str(tmp_path), "doc1", "en"
    base = process.doc_dir(storage, doc_id)
    os.makedirs(process.asset_dir(storage, doc_id), exist_ok=True)
    process._write_json(
        os.path.join(base, "source.json"),
        SemanticDocument(source_lang="ko",
                         blocks=[ParagraphBlock(text="원문", source_span_ids=[])]).model_dump(),
    )
    translated = SemanticDocument(
        source_lang="ko", target_lang=lang,
        blocks=[ParagraphBlock(text=text, source_span_ids=[])],
    )
    process.save_translated_model(storage, doc_id, lang, translated)
    render_docx(translated, process.output_docx_path(storage, doc_id, lang),
                process.asset_dir(storage, doc_id))
    return storage, doc_id, lang


def _docx_text(storage, doc_id, lang):
    d = Document(process.output_docx_path(storage, doc_id, lang))
    return "\n".join(p.text for p in d.paragraphs)


def test_bulk_replace_edits_docx_and_mirrors_model(tmp_path):
    storage, doc_id, lang = _seed(tmp_path)
    res = process.bulk_replace(storage, doc_id, lang, [_rule("Rated Load", "Load Rating")])
    assert res["total"] == 1
    text = _docx_text(storage, doc_id, lang)
    assert "Load Rating" in text and "Rated Load" not in text
    assert process.get_translated_model(storage, doc_id, lang).blocks[0].text == "Load Rating test"


def test_bulk_replace_persists_saved_rules_only(tmp_path):
    storage, doc_id, lang = _seed(tmp_path, "alpha beta")
    process.bulk_replace(storage, doc_id, lang, [
        _rule("alpha", "A", save=True),
        _rule("beta", "B", save=False),
    ])
    saved = process.load_replacements(storage, doc_id, lang)
    assert {r["find"] for r in saved} == {"alpha"}


def test_save_replacements_dedup_later_wins(tmp_path):
    storage, doc_id, lang = str(tmp_path), "doc1", "en"
    os.makedirs(process.doc_dir(storage, doc_id), exist_ok=True)
    process.save_replacements(storage, doc_id, lang, [_rule("x", "1")])
    process.save_replacements(storage, doc_id, lang, [_rule("x", "2")])
    assert process.load_replacements(storage, doc_id, lang) == [
        {"find": "x", "replace": "2", "match": "smart"}
    ]


def test_load_replacements_corrupt_is_empty(tmp_path):
    storage, doc_id, lang = str(tmp_path), "doc1", "en"
    os.makedirs(process.doc_dir(storage, doc_id), exist_ok=True)
    with open(process.replacements_path(storage, doc_id, lang), "w", encoding="utf-8") as fh:
        fh.write("{ not json")
    assert process.load_replacements(storage, doc_id, lang) == []


def test_preview_replace_counts_without_writing(tmp_path):
    storage, doc_id, lang = _seed(tmp_path, "load the load")
    res = process.preview_replace(storage, doc_id, lang, [_rule("load", "cargo")])
    assert res["total"] == 2
    assert "cargo" not in _docx_text(storage, doc_id, lang)


def test_bulk_replace_mirrors_edited_model(tmp_path):
    import json
    storage, doc_id, lang = _seed(tmp_path, "Rated Load")
    # Seed an edited model too (export prefers edited.{lang}.json); it must also be mirrored.
    process.save_edited_model(
        storage, doc_id, lang,
        SemanticDocument(source_lang="ko", target_lang=lang,
                         blocks=[ParagraphBlock(text="Rated Load", source_span_ids=[])]),
    )
    process.bulk_replace(storage, doc_id, lang, [_rule("Rated Load", "Load Rating")])
    with open(process.edited_path(storage, doc_id, lang), encoding="utf-8") as fh:
        data = json.load(fh)
    assert data["blocks"][0]["text"] == "Load Rating"
