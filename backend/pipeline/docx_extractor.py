"""DOCX -> SemanticDocument. The office-ingest counterpart to extractor+analyzer+
reconciler: instead of recovering structure from PDF geometry with a vision model, it
reads the structure DOCX already carries (heading styles, lists, tables, images).

PDF-only fields (page, source_span_ids, font geometry) are left at their defaults —
they are Optional in the model and unused by the translator/renderers for this path.
"""
from __future__ import annotations

import os
import re
from collections import Counter

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from models import (
    CaptionBlock,
    FigureBlock,
    HeadingBlock,
    ListItemBlock,
    ParagraphBlock,
    RunningBand,
    SemanticDocument,
    TableBlock,
    TableCell,
    universal_font,
)


_HEADING_RE = re.compile(r"^Heading (\d+)$")
_ALIGN = {
    WD_ALIGN_PARAGRAPH.CENTER: "center",
    WD_ALIGN_PARAGRAPH.RIGHT: "right",
    WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
}


def _align(p: Paragraph):
    return _ALIGN.get(p.alignment)


def _font_size_pt(p: Paragraph):
    for run in p.runs:
        if run.font.size is not None:
            return round(run.font.size.pt, 1)
    return None


def _bold(p: Paragraph):
    flags = [r.bold for r in p.runs if r.text.strip()]
    return bool(flags) and all(flags)


def _run_font(r) -> str | None:
    """A run's family — prefer the East-Asian font (Korean docs set it via w:rFonts
    w:eastAsia) over the Latin `name`, since the body script here is Korean."""
    rpr = r._element.rPr
    if rpr is not None:
        rf = rpr.find(qn("w:rFonts"))
        if rf is not None:
            ea = rf.get(qn("w:eastAsia"))
            if ea:
                return ea
    return r.font.name


def _font(p: Paragraph):
    names = [f for r in p.runs if r.text.strip() and (f := _run_font(r))]
    # Map to a universal family by class (sans→Arial / serif→Times New Roman) so the
    # translated output renders consistently regardless of installed fonts.
    return universal_font(Counter(names).most_common(1)[0][0]) if names else None


def _color(p: Paragraph):
    cols = [str(r.font.color.rgb) for r in p.runs
            if r.text.strip() and r.font.color and r.font.color.rgb is not None]
    return Counter(cols).most_common(1)[0][0].upper() if cols else None


def _italic(p: Paragraph):
    flags = [r.italic for r in p.runs if r.text.strip()]
    return True if (flags and all(flags)) else None


def _underline(p: Paragraph):
    flags = [bool(r.underline) for r in p.runs if r.text.strip()]
    return True if (flags and all(flags)) else None


def _heading_level(style_name: str):
    m = _HEADING_RE.match(style_name or "")
    return int(m.group(1)) if m else None


def _list_level(p: Paragraph):
    """0-based list level if the paragraph is numbered/bulleted, else None."""
    pPr = p._p.pPr
    if pPr is not None and pPr.numPr is not None:
        ilvl = pPr.numPr.ilvl
        return int(ilvl.val) if ilvl is not None and ilvl.val is not None else 0
    if (p.style.name or "").startswith("List"):
        return 0
    return None


def _style(p: Paragraph) -> dict:
    """Block-level visual style shared by every paragraph kind."""
    return {
        "bold": _bold(p), "align": _align(p), "font_size": _font_size_pt(p),
        "font": _font(p), "color": _color(p),
        "italic": _italic(p), "underline": _underline(p),
    }


def _paragraph_block(p: Paragraph):
    """Map a paragraph to a single Block, or None for an empty paragraph."""
    text = p.text.strip()
    if not text:
        return None
    style = p.style.name if p.style else ""
    st = _style(p)
    level = _heading_level(style)
    if level is not None:
        return HeadingBlock(level=level, text=text, **st)
    lvl = _list_level(p)
    if lvl is not None:
        return ListItemBlock(level=lvl, text=text, **st)
    if (style or "").lower().startswith("caption"):
        # captions don't carry bold in the model the same way; keep the full style set
        return CaptionBlock(text=text, **st)
    return ParagraphBlock(text=text, **st)


def _tc_text(tc) -> str:
    return "".join(n.text or "" for n in tc.findall(".//" + qn("w:t"))).strip()


_JC_MAP = {"center": "center", "right": "right", "end": "right", "both": "justify"}


def _tc_shading(tc):
    """Cell background fill (w:shd/@w:fill), ignoring 'auto'/white/no-fill."""
    tcPr = tc.tcPr
    if tcPr is None:
        return None
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        return None
    fill = shd.get(qn("w:fill"))
    if not fill or fill.lower() in ("auto", "ffffff"):
        return None
    return fill.upper() if re.match(r"^[0-9A-Fa-f]{6}$", fill) else None


def _tc_valign(tc):
    tcPr = tc.tcPr
    if tcPr is None:
        return None
    va = tcPr.find(qn("w:vAlign"))
    return va.get(qn("w:val")) if va is not None else None


def _tc_align(tc):
    """Horizontal alignment of the cell's first paragraph (w:pPr/w:jc)."""
    jc = tc.find(".//" + qn("w:pPr") + "/" + qn("w:jc"))
    return _JC_MAP.get(jc.get(qn("w:val"))) if jc is not None else None


def _table_block(table: Table) -> TableBlock:
    """Map a python-docx Table to a TableBlock, deriving colspan from w:gridSpan and
    rowspan from w:vMerge (restart starts a span; continue extends the cell above)."""
    rows = []
    grid_open: dict[int, "TableCell"] = {}  # grid column -> cell spanning down
    for tr in table._tbl.tr_lst:
        cells = []
        grid_col = 0
        trPr = tr.find(qn("w:trPr"))
        header_row = trPr is not None and trPr.find(qn("w:tblHeader")) is not None
        for tc in tr.tc_lst:
            tcPr = tc.tcPr
            colspan = 1
            vmerge = None
            if tcPr is not None:
                gs = tcPr.find(qn("w:gridSpan"))
                if gs is not None:
                    colspan = int(gs.get(qn("w:val")))
                vm = tcPr.find(qn("w:vMerge"))
                if vm is not None:
                    vmerge = vm.get(qn("w:val")) or "continue"
            if vmerge == "continue":
                origin = grid_open.get(grid_col)
                if origin is not None:
                    origin.rowspan += 1
                grid_col += colspan
                continue
            cell = TableCell(
                text=_tc_text(tc), colspan=colspan, rowspan=1, is_header=header_row,
                shading=_tc_shading(tc), valign=_tc_valign(tc), align=_tc_align(tc),
            )
            if vmerge == "restart":
                grid_open[grid_col] = cell
            else:
                grid_open.pop(grid_col, None)
            cells.append(cell)
            grid_col += colspan
        rows.append(cells)
    return TableBlock(rows=rows)


def _band(hf):
    """A section header/footer -> RunningBand (all non-empty lines joined into center)."""
    if hf is None:
        return None
    texts = [p.text.strip() for p in hf.paragraphs if p.text.strip()]
    if not texts:
        return None
    return RunningBand(center=" ".join(texts))


def _paragraph_images(p: Paragraph, d, asset_dir: str, counter: list) -> list:
    """Extract embedded images in `p` to asset_dir; return FigureBlocks (in order)."""
    figs = []
    for blip in p._p.findall(".//" + qn("a:blip")):
        rid = blip.get(qn("r:embed"))
        if not rid:
            continue
        part = d.part.related_parts.get(rid)
        if part is None:
            continue
        counter[0] += 1
        ext = os.path.splitext(part.partname)[1] or ".png"
        fname = f"docximg{counter[0]}{ext}"
        with open(os.path.join(asset_dir, fname), "wb") as fh:
            fh.write(part.blob)
        figs.append(FigureBlock(image_ref=fname))
    return figs


def extract_docx(docx_path: str, asset_dir: str, source_lang: str = "ko") -> SemanticDocument:
    os.makedirs(asset_dir, exist_ok=True)
    d = Document(docx_path)
    blocks = []
    img_counter = [0]
    for item in d.iter_inner_content():
        if isinstance(item, Paragraph):
            for fig in _paragraph_images(item, d, asset_dir, img_counter):
                blocks.append(fig)
            blk = _paragraph_block(item)
            if blk is not None:
                blocks.append(blk)
        elif isinstance(item, Table):
            blocks.append(_table_block(item))
    sec = d.sections[0]
    header = _band(sec.header)
    footer = _band(sec.footer)
    return SemanticDocument(
        source_lang=source_lang, running_header=header, running_footer=footer,
        blocks=blocks, **_section_geometry(sec),
    )


def _pt(length) -> float | None:
    """python-docx Length (EMU) -> points, or None."""
    return round(length.pt, 1) if length is not None else None


def _section_geometry(sec) -> dict:
    """Carry the docx's own paper size + margins through so the renderer reproduces them
    (matches the PDF path's page-geometry capture)."""
    return {
        "page_width": _pt(sec.page_width),
        "page_height": _pt(sec.page_height),
        "margin_left": _pt(sec.left_margin),
        "margin_right": _pt(sec.right_margin),
        "margin_top": _pt(sec.top_margin),
        "margin_bottom": _pt(sec.bottom_margin),
    }
